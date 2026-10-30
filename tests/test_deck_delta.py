"""Tests du mecanisme de generation PAR DELTA (docs/cadrage-ppt/deck_delta/) : version, refus d'ecraser, titre exact, statut, bout en bout sur un mini-deck."""
import os
import sys

import pytest
from pptx import Presentation

DELTA_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "docs", "cadrage-ppt", "deck_delta")
if DELTA_DIR not in sys.path:
    sys.path.insert(0, DELTA_DIR)

import deck  # noqa: E402
import deck_base  # noqa: E402
import deck_draw  # noqa: E402
import deck_garde  # noqa: E402
import deck_rapport  # noqa: E402

NOM = "offre-infra-as-a-product-V{}.pptx"


def _dossier(tmp_path, numeros, extras=()):
    for n in numeros:
        (tmp_path / NOM.format(n)).write_bytes(b"")
    for e in extras:
        (tmp_path / e).write_bytes(b"")
    return str(tmp_path)


def _mini_deck(chemin, titres=("Prochain pas", "Prochain pas : la suite", "Autre slide")):
    prs = Presentation()
    for t in titres:
        s = prs.slides.add_slide(prs.slide_layouts[5])  # titre seul
        s.shapes.title.text_frame.text = t
    prs.save(chemin)
    return str(chemin)


# --- resolution de version -------------------------------------------------------------------------------------------------------------------------------

def test_version_numerique_v10_depasse_v9(tmp_path):
    d = _dossier(tmp_path, [1, 9, 10, 2])
    assert [n for n, _ in deck_base.versions(d)] == [1, 2, 9, 10]  # un tri lexical donnerait 1, 10, 2, 9
    chemin, n = deck_base.dernier_livrable(d)
    assert n == 10 and chemin.endswith(NOM.format(10))
    assert deck_base.chemin_sortie(d).endswith(NOM.format(11))


def test_fichiers_hors_convention_ignores(tmp_path):
    d = _dossier(tmp_path, [3], extras=["~$offre-infra-as-a-product-V9.pptx", "bmad-iap-cadrage-synthese.pptx", "offre-infra-as-a-product-V7.pptx.tmp"])
    assert deck_base.dernier_livrable(d)[1] == 3


def test_base_par_version_et_dossier_vide(tmp_path):
    d = _dossier(tmp_path, [4, 5])
    assert deck_base.resoudre_base("V4", d).endswith(NOM.format(4))
    assert deck_base.resoudre_base(None, d).endswith(NOM.format(5))
    with pytest.raises(ValueError, match="introuvable"):
        deck_base.resoudre_base("V99", d)
    vide = tmp_path / "vide"
    vide.mkdir()
    with pytest.raises(ValueError, match="aucun livrable"):
        deck_base.dernier_livrable(str(vide))


# --- designation par titre exact -------------------------------------------------------------------------------------------------------------------------

def test_titre_exact_refuse_prefixe_absent_et_ambigu(tmp_path):
    prs = Presentation(_mini_deck(tmp_path / "m.pptx"))
    sid = prs.slides[0].slide_id
    assert deck_garde.selecteur_titre(prs, "Prochain pas") == f"id:{sid}"  # exact, pas la slide « Prochain pas : la suite »
    assert deck_garde.selecteur_titre(prs, "  prochain   PAS ") == f"id:{sid}"  # casse et espaces normalises
    assert deck_garde.selecteur_titre(prs, 2) == 2 and deck_garde.selecteur_titre(prs, "id:7") == "id:7"
    with pytest.raises(ValueError, match="0 slide"):
        deck_garde.selecteur_titre(prs, "Prochain")  # un prefixe n'est pas un titre
    with pytest.raises(ValueError, match="0 slide"):
        deck_garde.selecteur_titre(prs, "Inconnue")
    prs.slides[2].shapes.title.text_frame.text = "Prochain pas"
    with pytest.raises(ValueError, match="2 slide"):
        deck_garde.selecteur_titre(prs, "Prochain pas")


def test_slide_par_titre_exact(tmp_path):
    prs = Presentation(_mini_deck(tmp_path / "m.pptx"))
    assert deck_draw.slide_par_titre(prs, "Autre slide") is not None
    with pytest.raises(ValueError):
        deck_draw.slide_par_titre(prs, "Autre")


# --- statut : jamais « valide » ---------------------------------------------------------------------------------------------------------------------------

@pytest.mark.parametrize("pngs,sans", [(["s005.png"], False), ([], False), ([], True)])
def test_statut_ne_dit_jamais_valide(pngs, sans):
    s = deck_rapport.statut(pngs, sans)
    assert "valid" not in s.casefold() and "ok" not in s.casefold().split()
    assert s.startswith("RENDU PRODUIT — À REGARDER") if pngs else s.startswith("NON REGARDÉ")


# --- bout en bout sur un mini-deck, avec refus d'ecraser --------------------------------------------------------------------------------------------------

def _delta(tmp_path, ops_src):
    p = tmp_path / "delta_test.py"
    p.write_text(f"NOTE = 'test'\nOPS = {ops_src}\n", encoding="utf-8")
    return str(p)


def test_bout_en_bout_sans_ecraser(tmp_path):
    base = _mini_deck(tmp_path / "base.pptx")
    delta = _delta(tmp_path, "[{'slide': 'Autre slide', 'op': 'remplacer_texte', 'ancien': 'Autre', 'nouveau': 'Nouvelle'}]")
    sortie = str(tmp_path / "sortie.pptx")
    avant = open(base, "rb").read()
    dry = deck.executer([delta, "--base", base, "--sortie", sortie, "--dry-run"])
    assert not os.path.exists(sortie) and any("SIMULATION" in x for x in dry)
    lignes = deck.executer([delta, "--base", base, "--sortie", sortie, "--sans-rendu"])
    assert open(base, "rb").read() == avant  # la base n'est jamais modifiee
    prs = Presentation(sortie)
    assert len(prs.slides) == 3 and prs.slides[2].shapes.title.text_frame.text == "Nouvelle slide"
    assert lignes[-1].startswith("STATUT : NON REGARDÉ") and "valid" not in lignes[-1].casefold()
    assert not [f for f in os.listdir(tmp_path) if ".pret" in f or f.endswith(".verrou") or f.endswith(".tmp")]  # aucun temporaire
    with pytest.raises(ValueError, match="existe deja"):
        deck.executer([delta, "--base", base, "--sortie", sortie, "--sans-rendu"])  # un livrable existant n'est jamais ecrase
    assert len(Presentation(sortie).slides) == 3


def test_ecriture_refusee_sur_sortie_existante_avant_tout_calcul(tmp_path):
    base = _mini_deck(tmp_path / "base.pptx")
    existant = tmp_path / "livrable.pptx"
    existant.write_bytes(b"livrable existant")
    with pytest.raises(ValueError, match="existe deja"):
        deck.executer([_delta(tmp_path, "[]"), "--base", base, "--sortie", str(existant)])
    assert existant.read_bytes() == b"livrable existant"


def test_titre_ambigu_ou_prefixe_refuse_par_deck(tmp_path):
    base = _mini_deck(tmp_path / "base.pptx")
    delta = _delta(tmp_path, "[{'slide': 'Prochain', 'op': 'remplacer_texte', 'ancien': 'a', 'nouveau': 'b'}]")
    with pytest.raises(ValueError, match="titre exact"):
        deck.executer([delta, "--base", base, "--sortie", str(tmp_path / "s.pptx"), "--dry-run"])


def test_suppression_declaree_apparait_au_rapport(tmp_path):
    base = _mini_deck(tmp_path / "base.pptx")
    delta = _delta(tmp_path, "[{'slide': 'Autre slide', 'op': 'supprimer'}, {'slide': 'Prochain pas', 'op': 'remplacer_texte', 'ancien': 'Prochain pas', 'nouveau': 'X'}]")
    lignes = deck.executer([delta, "--base", base, "--sortie", str(tmp_path / "s.pptx"), "--dry-run"])
    assert any("Slide RETIREE" in x for x in lignes)  # une suppression DECLAREE passe


def test_balise_a_retiree_sans_source_refusee():
    op = {"slide": 1, "op": "remplacer_texte", "ancien": "Prix [À CONFIRMER]", "nouveau": "Prix 12"}
    with pytest.raises(ValueError, match="sans `source`"):
        deck_garde.verifier_balises([op])
    deck_garde.verifier_balises([dict(op, source="devis du 2026-10-01")])  # avec source : admis


def test_rendre_repli_sur_libreoffice_puis_statut_non_regarde(monkeypatch, tmp_path):
    """Si PowerPoint COM ET LibreOffice echouent, `_rendre` renvoie [] (statut « NON REGARDÉ ») au lieu de lever."""
    import sys
    import types

    import deck

    faux = types.ModuleType("rendu_com")
    faux.rendre = lambda *a, **k: (_ for _ in ()).throw(OSError("COM indisponible"))
    monkeypatch.setitem(sys.modules, "rendu_com", faux)
    faux_lo = types.ModuleType("rendu_slide")
    faux_lo.rendre = lambda *a, **k: (_ for _ in ()).throw(OSError("soffice muet"))
    monkeypatch.setitem(sys.modules, "rendu_slide", faux_lo)
    monkeypatch.setattr(deck.time, "sleep", lambda s: None)
    assert deck._rendre(str(tmp_path / "x.pptx"), [1], str(tmp_path)) == []


def test_rendre_prefere_powerpoint_com(monkeypatch, tmp_path):
    import sys
    import types

    import deck

    faux = types.ModuleType("rendu_com")
    faux.rendre = lambda sortie, rangs, dossier: [f"{dossier}/slide_{r:02d}.png" for r in rangs]
    monkeypatch.setitem(sys.modules, "rendu_com", faux)
    assert deck._rendre("x.pptx", [2, 5], "d") == ["d/slide_02.png", "d/slide_05.png"]
