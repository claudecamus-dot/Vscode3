"""Garde du deck « Offre Infra as a Product — V4 » (generate_offre_v4.py).

Cadrage géométrique, squelette en chapitres nommés (intercalaires, bandeau courant,
absence de renvoi par numéro), charte (pas de cyan sur du
texte), absence des mentions retirées, présence des étiquettes obligatoires.
`V4_GEN_DIR` (optionnel) pointe vers une copie mutée du générateur (test de mutation).
"""
import os
import re
import sys

import pytest

CADRAGE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "docs", "cadrage-ppt")
if os.environ.get("V4_GEN_DIR"):
    sys.path.insert(0, os.environ["V4_GEN_DIR"])
sys.path.append(CADRAGE)

import generate_offre_v4 as G  # noqa: E402
import pptx_deck_vscode3 as D  # noqa: E402

EMU = 914400
INTERDITS = ["gaspill", "moins de 30", "15 tickets", "sponsor qualifié", "je ne veux que",
             "8 familles", "huit familles", "mesuré dès le t0"]


@pytest.fixture(scope="module")
def prs():
    return G.construire()


def _titre(slide):
    return slide.shapes.title.text_frame.text.strip()


def _tout_texte(prs):
    return "\n".join(t for _, t in G.textes(prs))


def _runs(prs):
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    yield from p.runs


NOMS_CHAPITRES = ["Ouverture", "Contexte et enjeux", "Nos convictions", "Notre approche",
                  "Démarche", "Preuve", "Dispositif et prix", "Prochain pas", "Annexes"]
ETAPES_BOUCLE = ["Explorer", "Comprendre", "Prioriser", "Expérimenter", "Mesurer", "Ancrer"]
RENVOI = re.compile(r"(?i)\b(?:page|p\.|slide|chapitres?|partie|section|ch\.)\s*(?:n°\s*)?\d"
                    r"|\bvoir\s+(?:la\s+|le\s+)?(?:page|slide|chapitre)")
BANDEAU = re.compile(r"^(?P<nom>.+) · (?P<x>\d+)/(?P<y>\d+)$")


def _n(texte):
    """Espaces insécables (typographie française posée par `rich`) ramenés à l'espace."""
    return texte.replace(" ", " ").replace(" ", " ")


def _est_intercalaire(slide):
    return slide.slide_layout.name.startswith(("50 - Chapitre", "51 - Chapitre"))


def _est_contenu(slide):
    return slide.slide_layout.name.startswith("04 - Titre seul")


def _groupes(prs):
    """[(nom, [slides hors intercalaire])] : Ouverture (avant le 1er intercalaire) puis un
    groupe par intercalaire, nommé d'après le titre de l'intercalaire."""
    groupes = [[G.CHAPITRES[0].nom, []]]
    for slide in prs.slides:
        if _est_intercalaire(slide):
            groupes.append([_titre(slide), []])
        else:
            groupes[-1][1].append(slide)
    return groupes


def _bandeaux(slide):
    return [m for sh in slide.shapes if sh.has_text_frame
            for m in [BANDEAU.match(sh.text_frame.text.strip())] if m]


def test_chapitres_nommes_dans_l_ordre_avec_leur_intercalaire(prs):
    assert [c.nom for c in G.CHAPITRES] == NOMS_CHAPITRES
    titres = [_titre(s) for s in prs.slides]
    assert titres == G.titres_attendus()
    assert len(prs.slides) == len(titres)
    # un intercalaire nommé par chapitre, dans l'ordre, sauf Ouverture (la couverture en tient lieu)
    assert [_titre(s) for s in prs.slides if _est_intercalaire(s)] == NOMS_CHAPITRES[1:]
    # chaque intercalaire précède immédiatement les slides de son chapitre
    for c in G.CHAPITRES:
        if c.intercalaire:
            i = titres.index(c.nom)
            assert titres[i + 1:i + 1 + len(c.pages)] == [p.titre for p in c.pages]


def test_titres_courts_deux_lignes_max():
    assert all(len(p.titre) <= 80 for c in G.CHAPITRES for p in c.pages)


def test_numeros_d_intercalaire_calcules_depuis_la_liste(prs):
    imprimes = []
    for s in prs.slides:
        if _est_intercalaire(s):
            vus = [sh.text_frame.text.strip() for sh in s.shapes if sh.has_text_frame
                   and re.fullmatch(r"\d\d|[A-Z]", sh.text_frame.text.strip())]
            assert len(set(vus)) == 1, vus
            imprimes.append(vus[0])
    assert imprimes == [f"{n:02d}" for n in range(1, len(NOMS_CHAPITRES) - 1)] + ["A"]


def test_synthese_en_une_page_et_annonce_des_convictions(prs):
    par_titre = {_titre(s): s for s in prs.slides}
    synthese = par_titre[G.TITRE_SYNTHESE]
    pilules = [sh.text_frame.text for sh in synthese.shapes if sh.has_text_frame
               and sh.text_frame.text in ("Le constat", "La cible", "Pourquoi maintenant",
                                          "La démarche", "Ce que vous obtenez")]
    assert len(pilules) == 5
    annonce = par_titre[G.TITRE_CONVICTIONS]
    texte = _n("\n".join(sh.text_frame.text for sh in annonce.shapes if sh.has_text_frame))
    assert "LA CIBLE" in texte and "LA DÉMARCHE" in texte
    assert texte.count("dans la démarche : ") == len(G.CONVICTIONS) == 5
    assert all(c["etape"] in ETAPES_BOUCLE for c in G.CONVICTIONS)
    # la substance des cinq convictions tient sur UNE slide, marquée « proposition »
    t = _n(_texte_slide(par_titre[G.TITRE_CONVICTIONS_TABLEAU]))
    for cv in G.CONVICTIONS:
        for morceau in (cv["detail"], cv["principe"], cv["critere"], cv["chantier"]):
            assert morceau in t
    assert "PROPOSITION À VALIDER" in t


def test_equipe_sans_role_invente(prs):
    s = next(s for s in prs.slides if _titre(s) == G.TITRES[14])
    assert _n("\n".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
              ).count("rôle : à fournir") == 4


def test_bandeau_courant_coherent_avec_la_position(prs):
    for nom, slides in _groupes(prs):
        for rang, sl in enumerate(slides, 1):
            trouves = _bandeaux(sl)
            if _est_contenu(sl):
                assert len(trouves) == 1, (nom, _titre(sl))
                m = trouves[0]
                assert m["nom"] == nom.upper(), (nom, m.string)
                assert (int(m["x"]), int(m["y"])) == (rang, len(slides)), (nom, rang, m.string)
            else:
                assert trouves == [], (nom, _titre(sl))


def test_bandeau_9pt_muted_sous_le_titre(prs):
    muted = D.rgb(G.MUTED)
    for sl in prs.slides:
        for sh in sl.shapes:
            if sh.has_text_frame and BANDEAU.match(sh.text_frame.text.strip()):
                run = sh.text_frame.paragraphs[0].runs[0]
                assert run.font.size.pt == 9 and run.font.color.rgb == muted
                assert 0.81 <= sh.top / EMU and (sh.top + sh.height) / EMU <= 1.3


def test_aucun_renvoi_par_numero_de_page_ou_de_chapitre(prs):
    for i, tx in G.textes(prs):
        assert not RENVOI.search(_n(tx)), (i, tx[:60])
    # un chapitre se cite par son nom : le bandeau ne porte pas de numéro de chapitre
    for sl in prs.slides:
        for m in _bandeaux(sl):
            assert not re.search(r"\d", m["nom"]), m.string


def test_geometrie_rien_hors_slide(prs):
    assert D.verifier_geometrie(prs) == []
    w, h = prs.slide_width, prs.slide_height
    for i, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if sh.is_placeholder and sh.shape_type == 14 and not sh.has_text_frame:
                continue
            assert sh.left >= -1 and sh.top >= -1, (i, sh.name)
            assert sh.left + sh.width <= w + 1 and sh.top + sh.height <= h + 1, (i, sh.name)
            if _est_contenu(slide) and sh.has_text_frame and sh.text_frame.text.strip() \
                    and not sh.is_placeholder:
                assert (sh.left + sh.width) / EMU <= G.XR + 0.01, (i, sh.text_frame.text[:30])
                assert (sh.top + sh.height) / EMU <= 5.02, (i, sh.text_frame.text[:30])


def test_pas_de_chevauchement_de_textes(prs):
    for i, slide in enumerate(prs.slides, 1):
        if i == 1:      # couverture : pilules posées sur le cadre du gabarit, rendu relu à l'œil
            continue
        boites = [(sh.left / EMU, sh.top / EMU, (sh.left + sh.width) / EMU,
                   (sh.top + sh.height) / EMU, sh.text_frame.text[:30])
                  for sh in slide.shapes
                  if sh.has_text_frame and sh.text_frame.text.strip()]
        for a in range(len(boites)):
            for b in range(a + 1, len(boites)):
                x0, y0, x1, y1, ta = boites[a]
                u0, v0, u1, v1, tb = boites[b]
                dx, dy = min(x1, u1) - max(x0, u0), min(y1, v1) - max(y0, v0)
                assert not (dx > 0.02 and dy > 0.02), (i, ta, tb)


def test_aucune_mention_retiree(prs):
    texte = _tout_texte(prs).lower()
    for mot in INTERDITS:
        assert mot not in texte, mot


def test_etiquette_scenario_pedagogique(prs):
    scenario = next(s for s in prs.slides if _titre(s) == G.TITRES[10])
    assert G.ETIQUETTE_SCENARIO.lower() in "\n".join(
        sh.text_frame.text.lower() for sh in scenario.shapes if sh.has_text_frame)
    assert "scénario pédagogique" in _tout_texte(prs).lower()


def test_chiffres_de_marche_qualifies(prs):
    texte = _tout_texte(prs)
    assert "prévision Gartner" in texte
    assert "projection Gartner (juin 2025)" in texte
    # R15 : deux chiffres + une ligne de sources ; aucune autre citation non vérifiée
    assert texte.count("Gartner") == 3
    assert _n(G.SOURCES_MARCHE) in _n(texte) and "à confirmer avant diffusion" in texte
    assert "auront une platform team" not in texte      # prévision 2026 : plus au futur


def test_cyan_jamais_sur_du_texte(prs):
    cyan = D.rgb(G.ACCENT)
    for r in _runs(prs):
        if r.text.strip() and r.font.color is not None and r.font.color.type is not None:
            assert r.font.color.rgb != cyan, r.text


def test_taille_minimale_des_textes(prs):
    for i, slide in enumerate(prs.slides, 1):
        if not _est_contenu(slide):
            continue
        for sh in slide.shapes:
            if sh.has_text_frame and not sh.is_placeholder:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.text.strip():
                            assert r.font.size is not None and r.font.size.pt >= 9, (i, r.text)


def _slide(prs, titre_):
    return next(s for s in prs.slides if _titre(s) == titre_)


def _texte_slide(s):
    return _n("\n".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame))


def test_r04_chaque_constat_relie_a_un_levier(prs):
    texte = _texte_slide(_slide(prs, G.TITRES[1]))
    assert texte.count("Levier : ") == 4
    leviers = {nom for nom, _, _ in G.LEVIERS}
    for lien in re.findall(r"Levier : ([^\n]+)", texte):
        assert {p.strip() for p in lien.split(",")} <= leviers, lien


def test_r09_statut_des_agents_du_coach_pas_presente_comme_disponible(prs):
    texte = _texte_slide(_slide(prs, G.TITRES[8]))
    assert "INTENTION, STATUT À VALIDER" in texte
    assert "DISPONIBLES" not in texte


def test_r12_r13_annexe_engagement(prs):
    texte = _texte_slide(_slide(prs, G.TITRES_ANNEXES[2]))
    for attendu in ("PROPOSITION À VALIDER", "Entrer", "Passer au second temps",
                    "Passer la main", "seuils chiffrés à fixer", "Aucun montant"):
        assert attendu in texte, attendu
    assert texte.count("Décide") == 6          # un porteur par décision, aucune sans
    assert "€" not in texte


def test_r16_annexe_agent_ready_option_non_obligatoire(prs):
    texte = _texte_slide(_slide(prs, G.TITRES_ANNEXES[3]))
    for attendu in ("PROPOSITION À VALIDER", "Public visé", "Garde-fous", "jamais un prérequis"):
        assert attendu in texte, attendu


def _pages(nom):
    return [p.titre for ch in G.CHAPITRES if ch.nom == nom for p in ch.pages]


def test_restructuration_a_resserrer_en_place(prs):
    # A : convictions <= 3 slides intercalaire inclus ; contexte fusionné ; coach en tête
    assert 1 + len(_pages("Nos convictions")) <= 3
    assert _pages("Contexte et enjeux")[0] == G.TITRES[1]
    assert len(_pages("Contexte et enjeux")) == 2
    assert _pages("Notre approche")[0] == G.TITRES[6]
    # la gouvernance IA répond à une objection : elle vit dans Dispositif et prix, pas Démarche
    assert G.TITRES[9] in _pages("Dispositif et prix")
    assert G.TITRES[9] not in _pages("Démarche")
    assert not any(c.nom == "Équipe" for c in G.CHAPITRES)


def test_c_trous_boucles_et_visibles(prs):
    risques = _texte_slide(_slide(prs, G.TITRES[2]))
    assert risques.count("« ") >= 5 and "PROPOSITION À VALIDER" in risques
    for mot in ("données", "plateforme", "Run", "IA", "gains"):
        assert mot in risques
    cal = _texte_slide(_slide(prs, G.TITRES[15]))
    assert cal.count("à valider") >= 10
    assert not re.search(r"\d+\s*(?:jours?|semaines?|mois|j/h|k€)", cal.lower())
    cloture = _texte_slide(_slide(prs, G.TITRES[14]))
    assert cloture.count("rôle : à fournir") == 4 and "go / no-go" in cloture


def test_prix_sans_montant_invente(prs):
    texte = " ".join(_tout_texte(prs).replace(" ", " ").replace(" ", " ").split())
    assert texte.count("MONTANT : À COMPLÉTER") == 2
    assert "€" not in texte


def test_prix_forme_jours_x_roles_tout_a_zero(prs):
    texte = _texte_slide(_slide(prs, G.TITRES[16]))
    for attendu in ("TAUX JOUR", "Remise", "Échéancier", "Délai de paiement",
                    "Part à la réussite", "TOUTES VALEURS : 0, À VALIDER"):
        assert attendu in texte, attendu
    # aucune valeur chiffrée autre que 0 (pas de montant, jour, taux ni pourcentage inventé)
    texte = re.sub(r"[A-ZÉ ]+ · \d+/\d+", "", _n(texte))      # bandeau courant
    nombres = re.findall(r"\d+", texte.replace("TOUTES VALEURS : 0", ""))
    assert nombres and set(nombres) == {"0"}, set(nombres)


def test_gouvernance_ia_deux_faces_humaines(prs):
    texte = _texte_slide(_slide(prs, G.TITRES[9]))
    for attendu in ("RACI humain / agent", "Accountability humaine", "Gates humains",
                    "Charte IA", "hallucination et dérive", "propriétaire nommé"):
        assert attendu in texte, attendu


def test_reprises_mecanisme_seulement(prs):
    tab = _texte_slide(_slide(prs, G.TITRE_CONVICTIONS_TABLEAU))
    assert _n(G.REGLE_PASSAGE) in tab and _n(G.CONTROLE_HUMAIN) in tab
    assert "Hors périmètre" in _texte_slide(_slide(prs, G.TITRES[14]))
    cal = _texte_slide(_slide(prs, G.TITRES[15]))
    for attendu in ("COPIL", "Livrable par échéance", "HYPOTHÈSES ET CONDITIONS DE TENUE",
                    "pilotage", "suivi opérationnel", "revue de valeur"):
        assert attendu.lower() in cal.lower(), attendu
    fiche = _texte_slide(_slide(prs, G.TITRES_ANNEXES[4]))
    for attendu in ("Objectifs", "Activités", "Livrables", "Prérequis côté client",
                    "Passage de jalon", "Indicateurs de succès"):
        assert attendu.lower() in fiche.lower(), attendu


def _formes(s):
    return [(sh.auto_shape_type, sh) for sh in s.shapes if sh.shape_type == 1]


def test_v4_forme_avant_apres_relies(prs):
    s = _slide(prs, G.TITRES[3])
    texte = _texte_slide(s)
    assert "AVANT" in texte and "APRÈS" in texte
    for pilier in ("Consommateurs", "Catalogue", "Ownership", "Mesure"):
        assert pilier in texte
    assert "Le Run reste une responsabilité de service" in texte     # contenu V3 conservé


def test_v4_calendrier_gelules_trois_etats_et_instances(prs):
    s = _slide(prs, G.TITRES[15])
    texte = _texte_slide(s)
    assert "plein, contour, clair" in texte.lower()
    from pptx.enum.shapes import MSO_SHAPE
    gelules = [sh for sh in s.shapes if sh.shape_type == 1
               and sh.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE and sh.height / EMU < 0.25]
    etats = set()
    for sh in gelules:
        plein = sh.fill.type == 1 and str(sh.fill.fore_color.rgb) == G.NAVY.lstrip("#")
        etats.add("plein" if plein else ("contour" if sh.line.fill.type == 1 else "clair"))
    assert etats >= {"plein", "contour", "clair"}, etats
    assert texte.count("fréquence : à valider") == 3     # trois instances, même gabarit


def test_v4_fiche_de_phase_navigateur_une_phase_active(prs):
    s = _slide(prs, G.TITRES_ANNEXES[4])
    texte = _texte_slide(s)
    for ph_ in ("Assessment flash", "Restitution", "Pilote Produit + IA", "Ancrage",
                "Extension éventuelle"):
        assert ph_ in texte
    assert "DURÉE : À VALIDER" in texte
    from pptx.enum.shapes import MSO_SHAPE
    onglets = [sh for sh in s.shapes if sh.shape_type == 1
               and sh.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE
               and abs(sh.top / EMU - 1.4) < 0.01]
    assert len(onglets) == 5
    assert sum(1 for sh in onglets if sh.fill.type == 1
               and str(sh.fill.fore_color.rgb) == G.NAVY.lstrip("#")) == 1   # un sur N en accent
