"""Les trois filets portes dans pptx_deck.py le 2026-09-09 (finding robustesse de
l'audit-technique VScode5 : `verifier_geometrie` seule ne voit que les bords des
formes de la slide).

Chaque test est ADVERSARIAL : il commence par prouver que le filet CRIE sur le
defaut qu'il pretend attraper (un texte qui deborde d'un pouce, une forme qui
mord le badge de pagination, un plancher passe sous le badge), puis seulement
ensuite qu'il se tait sur le cas correct. Un test qui n'a jamais vu son filet
rouge ne prouve rien — lecon du hub du 2026-09-09 (« quatre gardes vertes et
contournables le meme matin »).

Les presentations sont construites EN MEMOIRE : aucun fichier ecrit, aucun
rendu, aucun deck touche. Le seul fichier lu est le gabarit versionne
template-octo.pptx, en lecture seule — c'est lui qui porte la zone reelle du
numero de page.
"""
import importlib.util
from pathlib import Path

import pytest
from pptx import Presentation

CADRAGE = Path(__file__).resolve().parents[1] / "docs" / "cadrage-ppt"
TEMPLATE = CADRAGE / "template-octo.pptx"

# Constantes de gabarit du generateur (generate_deck.py) que ces filets
# protegent. Copiees ici a dessein : importer generate_deck ferait charger tout
# le module (5000 lignes, images, reseau) pour trois nombres.
BORD_DROIT = 9.15
CONTENT_BOTTOM = 5.45
ZONE_BADGE = (9.2512, 5.0874, 9.7974, 5.3374)  # mesuree sur template-octo.pptx


@pytest.fixture(scope="module")
def D():
    spec = importlib.util.spec_from_file_location(
        "pptx_deck_sous_test", CADRAGE / "pptx_deck_vscode3.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _prs_vide():
    """Presentation neuve (10 x 7.5 in) avec une slide vierge — aucun gabarit,
    donc aucune zone de numero de page declaree."""
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])  # « Blank »
    return prs


def _prs_gabarit(nom_layout="04 - Titre seul"):
    """Une slide sur le VRAI gabarit OCTO, sans passer par generate_deck."""
    prs = Presentation(str(TEMPLATE))
    for sld in list(prs.slides._sldIdLst):
        prs.part.drop_rel(sld.get(
            "{http://schemas.openxmlformats.org/officeDocument/2006/"
            "relationships}id"))
        prs.slides._sldIdLst.remove(sld)
    layout = next(lay for lay in prs.slide_masters[0].slide_layouts
                  if lay.name == nom_layout)
    prs.slides.add_slide(layout)
    return prs


# --- Filet 1 : le texte deborde de SA PROPRE boite ---------------------------

TEXTE_LONG = ("Une phrase de restitution suffisamment longue pour occuper "
              "plusieurs lignes une fois repliee sur trois pouces de large, "
              "comme les puces de cartes du deck de cadrage IAP.")
TAILLE = 12.0
LARGEUR = 3.0


def _hauteur_estimee(D, texte=TEXTE_LONG, largeur=LARGEUR, taille=TAILLE):
    """Le meme modele de hauteur que le filet, pour dimensionner la boite du
    test a un debordement CHOISI plutot qu'a un debordement espere."""
    lignes = D.estimer_lignes(texte, largeur, taille, cpi_ref=10.7)
    return lignes * (taille * 0.017 + 4 / 72)


def test_debordement_dun_pouce_est_signale(D):
    """La question adversariale n°1 : le filet echoue-t-il VRAIMENT si un texte
    deborde d'un pouce ?"""
    prs = _prs_vide()
    hauteur = _hauteur_estimee(D) - 1.0
    assert hauteur > 0, "le cas de test doit rester une boite de hauteur reelle"
    D.add_text(prs.slides[0], 0.5, 0.5, LARGEUR, hauteur, [(TEXTE_LONG, {"size": TAILLE})])

    problemes = D.verifier_debordements_texte(prs)

    assert len(problemes) == 1, problemes
    assert "slide 1" in problemes[0]
    assert "> boite" in problemes[0]


def test_boite_a_la_bonne_taille_ne_declenche_rien(D):
    """Le pendant : un filet qui crie sur une boite correcte finit debranche."""
    prs = _prs_vide()
    D.add_text(prs.slides[0], 0.5, 0.5, LARGEUR, _hauteur_estimee(D) + 0.5,
               [(TEXTE_LONG, {"size": TAILLE})])

    assert D.verifier_debordements_texte(prs) == []


def test_debordement_juste_sous_la_tolerance_reste_silencieux(D):
    """Ou passe exactement la frontiere : `tolerance_in` (0.15in par defaut) est
    une tolerance ASSUMEE, pas un oubli. Ce test la fixe dans les deux sens."""
    prs = _prs_vide()
    est = _hauteur_estimee(D)
    D.add_text(prs.slides[0], 0.5, 0.5, LARGEUR, est - 0.10,
               [(TEXTE_LONG, {"size": TAILLE})])

    assert D.verifier_debordements_texte(prs) == []
    assert D.verifier_debordements_texte(prs, tolerance_in=0.05) != []


def test_le_compte_dit_combien_de_zones_ont_ete_regardees(D):
    """Un vert qui ne dit pas ce qu'il a couvert laisse croire qu'il couvre
    tout : ici une zone examinee et une zone ecartee (ancrage bas)."""
    from pptx.enum.text import MSO_ANCHOR
    prs = _prs_vide()
    D.add_text(prs.slides[0], 0.5, 0.5, LARGEUR, 2.0, [("court", {})])
    D.add_text(prs.slides[0], 0.5, 3.0, LARGEUR, 2.0, [("court", {})],
               anchor=MSO_ANCHOR.BOTTOM)

    compte = {}
    assert D.verifier_debordements_texte(prs, compte=compte) == []
    assert compte == {"examinees": 1, "ignorees": 1}


# --- Filet 2 : une forme de contenu recouvre le chrome du gabarit ------------

def test_forme_qui_mord_le_badge_de_pagination_est_signalee(D):
    """La question adversariale n°2, cas franc : une forme posee sur le badge de
    pagination du gabarit ne sort PAS de la slide — `verifier_geometrie` la
    laisse passer, ce filet doit la voir."""
    prs = _prs_gabarit()
    # 0.05in (3,6 pt) de recouvrement sur le coin haut-gauche de la zone.
    D.add_rect(prs.slides[0], 8.0, 4.6, 1.30, 0.54, fill="#ffffff")

    assert D.verifier_geometrie(prs) == [], "le cas de test ne doit rien devoir a l'autre filet"
    problemes = D.verifier_chrome_gabarit(prs)
    assert len(problemes) == 1, problemes
    assert "recouvre la zone du numero de page" in problemes[0]


def test_recouvrement_de_1pt_passe_sous_la_tolerance_par_defaut(D):
    """La question adversariale n°2, cas limite : 1 pt = 0.0139in, SOUS la
    tolerance de 0.02in du filet. Repondre « oui, il le voit » serait faux — ce
    test fixe la vraie frontiere : muet a 0.02, rouge des qu'on serre."""
    prs = _prs_gabarit()
    un_point = 1 / 72
    largeur = (ZONE_BADGE[0] + un_point) - 8.0
    D.add_rect(prs.slides[0], 8.0, 4.6, largeur, 0.54, fill="#ffffff")

    assert D.verifier_chrome_gabarit(prs) == []
    assert D.verifier_chrome_gabarit(prs, marge_in=0.0) != []


def test_le_compte_du_filet_chrome_pose_ses_trois_clefs(D):
    """Le docstring promet `examinees`, `ignorees` et `groupes` : un appelant
    qui lit `compte["groupes"]` sur un deck sans groupe doit lire 0, pas lever
    un KeyError (revue bmad-code-review du 2026-09-09)."""
    prs = _prs_gabarit()
    D.add_rect(prs.slides[0], 0.615, 4.9, 1.0, 0.5, fill="#ffffff")

    compte = {}
    D.verifier_chrome_gabarit(prs, compte=compte)

    # 2 formes : le rectangle pose ici + le placeholder de titre que le layout
    # « 04 - Titre seul » recopie sur la slide.
    assert compte == {"examinees": 2, "ignorees": 0, "groupes": 0}


def test_forme_a_gauche_du_badge_ne_declenche_rien(D):
    """Le contenu du deck s'arrete a BORD_DROIT : le cas nominal doit rester
    vert, sinon le filet serait debranche au premier build."""
    prs = _prs_gabarit()
    D.add_rect(prs.slides[0], 0.615, 4.9, BORD_DROIT - 0.615, 0.5, fill="#ffffff")

    assert D.verifier_chrome_gabarit(prs) == []


def test_zone_lue_dans_un_groupe_du_gabarit_et_non_sa_boite_englobante(D):
    """L'adaptation propre a CE gabarit : sur 10 layouts, le champ slidenum est
    porte par un ENFANT d'un groupe qui couvre presque toute la slide. Sans la
    descente, la « zone du numero de page » du layout 63 ferait 9,5 x 5,0in et
    le filet crierait sur toute forme de contenu."""
    prs = _prs_gabarit("63 - Titre, contenu et visuel à droite - cadre blanc")
    zones = D.zones_numero_page(prs.slides[0])

    assert zones, "le gabarit declare bien un numero de page sur ce layout"
    for zone in zones:
        assert zone[2] - zone[0] < 1.0, f"zone trop large, boite de groupe ? {zone}"
        assert zone[3] - zone[1] < 0.5, f"zone trop haute, boite de groupe ? {zone}"
    # et le contenu nominal reste vert sur ce layout aussi
    D.add_rect(prs.slides[0], 0.615, 4.9, BORD_DROIT - 0.615, 0.5, fill="#ffffff")
    assert D.verifier_chrome_gabarit(prs) == []


def test_aucun_des_layouts_du_gabarit_ne_rend_une_boite_de_groupe(D):
    """Le docstring parle de 10 layouts concernes ; un seul teste ne prouve que
    celui-la (revue du 2026-09-09). Les 34 layouts de template-octo.pptx sont
    ici passes en revue : aucune zone ne doit avoir la taille d'un groupe."""
    prs = Presentation(str(TEMPLATE))
    for sld in list(prs.slides._sldIdLst):
        prs.slides._sldIdLst.remove(sld)
    layouts = list(prs.slide_masters[0].slide_layouts)
    assert len(layouts) > 20, "gabarit inattendu : la couverture du test ne veut plus rien dire"

    trop_larges = []
    for layout in layouts:
        for zone in D.zones_numero_page(prs.slides.add_slide(layout)):
            if zone[2] - zone[0] >= 1.0 or zone[3] - zone[1] >= 0.5:
                trop_larges.append((layout.name, zone))

    assert not trop_larges, trop_larges


def test_la_zone_est_lue_sur_le_gabarit_pas_prise_dans_la_constante(D):
    """Le repli ne doit jamais masquer une lecture reelle : sur un gabarit
    quelconque (celui de python-pptx, qui declare son propre numero de page en
    7.17/6.95), la zone rendue est la SIENNE, pas `_ZONE_NUMERO_PAGE_IN`."""
    zones = D.zones_numero_page(_prs_vide().slides[0])

    assert zones and zones != [D._ZONE_NUMERO_PAGE_IN]


def test_repli_si_le_gabarit_ne_declare_aucun_numero_de_page(D, monkeypatch):
    """Un gabarit muet ne doit pas faire disparaitre le filet en silence."""
    monkeypatch.setattr(D, "_zones_numero_page_de", lambda conteneur: [])

    assert D.zones_numero_page(_prs_vide().slides[0]) == [D._ZONE_NUMERO_PAGE_IN]


# --- Filet 3 : le plancher de dessin a decroche du gabarit -------------------

def test_plancher_sous_le_badge_est_signale_si_la_bande_l_atteint(D):
    """La question adversariale n°3 : le filet echoue-t-il si le plancher passe
    sous le badge ? Oui — des lors que la bande dessinee atteint le badge."""
    prs = _prs_gabarit()

    problemes = D.verifier_plancher_de_dessin(prs, CONTENT_BOTTOM,
                                              bord_droit_in=9.30)

    assert len(problemes) == 1, problemes
    assert "plancher de dessin" in problemes[0]
    assert "5.45" in problemes[0]


def test_constantes_reelles_du_generateur_restent_vertes(D):
    """Le cas nominal de CE projet : CONTENT_BOTTOM (5.45) passe bien sous le
    haut du badge (5.09), mais BORD_DROIT (9.15) s'arrete a sa gauche (9.25) —
    aucun recouvrement, donc aucun constat. C'est l'adaptation au canal : sans
    `bord_droit_in`, la comparaison purement verticale de l'homologue VSCode4
    (qui dessine pleine largeur, lui) crierait a chaque build."""
    prs = _prs_gabarit()

    assert D.verifier_plancher_de_dessin(prs, CONTENT_BOTTOM,
                                         bord_droit_in=BORD_DROIT) == []
    assert D.verifier_plancher_de_dessin(prs, CONTENT_BOTTOM) != []


def test_badge_qui_derive_dans_la_bande_est_signale(D, monkeypatch):
    """Le vrai cas rouge de la forme d'appel de CE projet, et le trou que la
    revue du 2026-09-09 a trouve : avec `bord_droit_in=9.15`, le filet est vert
    quel que soit `plancher_in` tant que le badge reste a droite de 9.15 — un
    `== []` ne distingue donc pas « rien ne se recouvre » de « filet inerte ».
    Ce test-ci pose une derive REELLE (badge ramene a gauche de la bande, comme
    un gabarit v7 qui deplacerait son bloc) et exige le rouge."""
    prs = _prs_gabarit()
    monkeypatch.setattr(D, "zones_numero_page",
                        lambda slide, defaut=None: [(9.10, 5.0874, 9.65, 5.3374)])

    problemes = D.verifier_plancher_de_dessin(prs, CONTENT_BOTTOM,
                                              bord_droit_in=BORD_DROIT)

    assert len(problemes) == 1, problemes
    assert "plancher de dessin" in problemes[0]


def test_plancher_au_dessus_du_badge_est_vert_dans_tous_les_cas(D):
    prs = _prs_gabarit()

    assert D.verifier_plancher_de_dessin(prs, 5.00) == []
    assert D.verifier_plancher_de_dessin(prs, 5.00, bord_droit_in=10.0) == []


# --- Le filet central : `verifier_geometrie` --------------------------------
#
# C'est LUI qui decide, via `_controler` puis `build`, si le deck s'ecrit sous
# son nom livrable ou sous `.INVALIDE.pptx`. Il n'avait pourtant aucun test de
# detection POSITIVE : ses deux seules occurrences dans tests/ etaient un
# `== []` en premisse d'un test portant sur un autre filet et une chaine de
# parametrage du cablage — il aurait pu rendre `[]` en dur sans qu'un seul test
# ne rougisse (finding risque_technique de l'audit du 2026-09-13).

def test_forme_hors_cadre_est_signalee(D):
    """La detection qui n'avait jamais ete exigee : une forme posee au-dela du
    bord droit d'une slide de 10in doit rendre exactement un constat."""
    prs = _prs_vide()
    D.add_rect(prs.slides[0], 12.0, 1.0, 3.0, 0.4, fill="#ffffff")

    problemes = D.verifier_geometrie(prs)

    assert len(problemes) == 1, problemes
    assert "hors cadre" in problemes[0]
    assert "slide 1" in problemes[0]


def test_hauteur_negative_est_signalee(D):
    """L'angle mort ferme le 2026-09-13, reproduit tel quel depuis la sonde de
    l'audit : `add_rect(s, 1, 1, 3, -0.4)` fabrique une boite INVERSEE. Elle ne
    depasse aucun bord (son bas `t + h` est au-DESSUS de son haut), donc le
    controle de bords la laissait passer et le deck sortait « CONTROLE: OK »."""
    prs = _prs_vide()
    shp = D.add_rect(prs.slides[0], 1.0, 1.0, 3.0, -0.4, fill="#ffffff")
    assert shp.height < 0, "le cas de test doit vraiment poser une boite inversee"
    l, t, w, h = shp.left, shp.top, shp.width, shp.height
    assert 0 <= l and 0 <= t and (l + w) <= prs.slide_width \
        and (t + h) <= prs.slide_height, \
        "cette boite ne depasse aucun bord : seul un controle de dimension la voit"

    problemes = D.verifier_geometrie(prs)

    assert len(problemes) == 1, problemes
    assert "dimension non positive" in problemes[0]
    assert "h=-0.40" in problemes[0]


def test_hauteur_nulle_est_signalee(D):
    """Le pendant degenere : une soustraction qui tombe pile a zero rend une
    forme invisible au rendu, pas une forme correcte."""
    prs = _prs_vide()
    D.add_rect(prs.slides[0], 1.0, 1.0, 3.0, 0.0, fill="#ffffff")

    problemes = D.verifier_geometrie(prs)

    assert len(problemes) == 1, problemes
    assert "dimension non positive" in problemes[0]


def test_largeur_negative_est_signalee(D):
    """Meme filet, autre axe : la garde ne doit pas ne regarder que la hauteur
    parce que c'est par la que le defaut est arrive."""
    prs = _prs_vide()
    D.add_rect(prs.slides[0], 4.0, 1.0, -2.0, 0.5, fill="#ffffff")

    problemes = D.verifier_geometrie(prs)

    assert len(problemes) == 1, problemes
    assert "dimension non positive" in problemes[0]
    assert "w=-2.00" in problemes[0]


def test_hauteur_calculee_par_soustraction_qui_decroche_est_signalee(D):
    """Le scenario REEL du generateur, pas un cas de laboratoire : une hauteur
    de carte calculee `CONTENT_BOTTOM - top` quand un contenu plus long a
    repousse `top` sous le plancher. Un seul des cinq sites de ce calcul porte
    une garde a l'appel ; le filet, lui, les couvre tous."""
    prs = _prs_vide()
    top = CONTENT_BOTTOM + 0.3          # le contenu a debordé sous le plancher
    card_h = CONTENT_BOTTOM - top       # -> -0.30in
    D.add_rect(prs.slides[0], 0.615, top, 3.0, card_h, fill="#ffffff")

    assert [p for p in D.verifier_geometrie(prs) if "dimension non positive" in p]


def test_formes_correctes_ne_declenchent_rien(D):
    """Le pendant obligatoire : un filet qui crie sur une slide correcte est
    debranche au premier build. Trois formes nominales, zero constat."""
    prs = _prs_vide()
    D.add_rect(prs.slides[0], 0.615, 0.5, 3.0, 1.2, fill="#ffffff")
    D.add_rect(prs.slides[0], 0.615, 2.0, 8.0, 0.02, fill="#ffffff")  # filet fin
    D.add_text(prs.slides[0], 0.615, 3.0, 3.0, 0.6, [("court", {})])

    assert D.verifier_geometrie(prs) == []
