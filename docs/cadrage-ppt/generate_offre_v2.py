"""Deck « Offre Infra as a Product — V2 » : squelette en chapitres NOMMÉS.

Source : plan de refonte et recommandations V2 (docs/imports/recommandations/),
contenu V1 revu (convictions : version retouchée par l'utilisateur).
Réutilise les primitives de `generate_offre.py` (non modifié) et le gabarit OCTO
(`source/offre-iap-source.pptx`) : on garde la couverture ; chaque chapitre a un
intercalaire (gabarit « Chapitre » cloné, photo CC0 en cache) ; toutes les autres
slides sont recréées sur la mise en page « Titre seul ».

UNE SEULE source de vérité : `CHAPITRES` (plus bas). Le sommaire, les intercalaires
(titre et numéro calculés) et le bandeau courant « NOM · x/y » en dérivent. Aucun
numéro de chapitre n'est écrit en dur, aucun renvoi par numéro dans le texte.
Ouverture n'a pas d'intercalaire : la couverture en tient lieu, et un numéro
d'intercalaire décalerait ceux des chapitres numérotés (le gabarit du sommaire
n'imprime que 01 à 07) ; Annexes porte une lettre calculée depuis son nom.

Charte (CLAUDE.md) : navy encre, cyan en aplat seulement, jamais sur du texte ;
« un sur N en accent ». Pas de « gaspillage ». Un chapitre se cite par son nom.
Décisions : la mention « 8 familles » est supprimée (non détaillée en annexe, et
sa source emploie le mot interdit) ; le cas démonstrateur est un scénario
pédagogique sans chiffre ; chiffres de marché = prévision / projection Gartner.

Usage : python generate_offre_v2.py  (depuis docs/cadrage-ppt/)
"""
import copy
import os
import sys
from dataclasses import dataclass
from functools import partial

import pptx_deck_vscode3 as D
from generate_offre import (
    ACCENT,
    MUTED,
    NAVY,
    SOURCE,
    TRACK,
    WHITE,
    X0,
    XR,
    W,
    badge,
    carte,
    chapo,
    chevron,
    cols,
    fleche,
    forme,
    gras,
    ph,
    photo_chapitre,
    pilule,
    retirer,
    rich,
)
from generate_offre import LINE as GRIS
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
# sorties generees : docs/exports/decks/ (jamais a cote des scripts)
SORTIE = os.path.join(os.path.normpath(os.path.join(HERE, "..", "exports", "decks")),
                      "offre-infra-as-a-product-V2.pptx")
TOP = 1.4

TITRES = [  # S01 .. S15, dans l'ordre
    "Faire de l'Infrastructure un levier de valeur, avec le Produit et l'IA",
    "Le guichet Infra atteint ses limites",
    "La pression monte : complexité, fiabilité, coûts et demande IA",
    "Infra as a Product : des capacités pensées pour leurs consommateurs",
    "Six leviers au service de cinq résultats métier",
    "Un point d'entrée adapté à la maturité et à la criticité de chaque entité",
    "Le coach produit augmenté est le moteur de la transformation",
    "Une boucle continue pour décider sur des faits, expérimenter et ancrer",
    "Des agents augmentent le coach ; d'autres améliorent l'Infra du client",
    "L'IA est gouvernée : données, droits, validations et traçabilité",
    "De la douleur au produit : une mission pilote de bout en bout",
    "Nous mesurons la valeur client et l'apport du coaching augmenté",
    "Un dispositif OCTO + client qui construit son propre relais",
    "Deux temps d'engagement : diagnostiquer, puis expérimenter et déployer",
    "Démarrer sur un périmètre utile, prouver, puis étendre",
]
TITRES_ANNEXES = [
    "Dix problèmes, dix apports, dix indicateurs à suivre",
    "Les alternatives sont des compléments que l'offre intègre",
]
TITRE_EQUIPE = "Quatre personnes à contacter pour l'offre"
TITRE_SYNTHESE = "Cinq points pour passer d'un guichet saturé à une Infra créatrice de valeur"
TITRE_SOMMAIRE = "Sommaire"
TITRE_CONVICTIONS = "Cinq convictions : une cible à construire et une démarche pour y conduire"
ETIQUETTE_SCENARIO = "Scénario pédagogique — illustratif, non client"
GARTNER_80 = "prévision Gartner"
GARTNER_40 = "projection Gartner (juin 2025)"
# photos CC0 en cache (_img/_brut_*), une par intercalaire dans l'ordre des chapitres
PHOTOS_INTERCALAIRES = ["mountains_0", "canyon_0", "forest_0", "riverdelta_5", "wheatfield_0",
                        "ocean_0", "sunset_1", "dunes_0"]
# slides « Chapitre » du gabarit source (n° source) clonées pour les intercalaires
MODELES_INTERCALAIRE = (4, 14, 17, 22, 26, 28)
BANDEAU_Y = 0.93      # vide sous le titre (0,81) et au-dessus du chapô (1,3 à 1,4)


# ---------------------------------------------------------------- aides
def titre(s, texte):
    s.shapes.title.text = texte


def t(s, x, y, w, h, texte, size=10, bold=False, color=NAVY, italic=False,
      anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
    return rich(s, x, y, w, h, [(texte, {"size": size, "bold": bold, "color": color,
                                         "italic": italic})], anchor=anchor, align=align)


def msg(s, texte, cle=None, h=0.45, y=TOP):
    chapo(s, y, [gras(texte, cle, size=11)], h=h)


def etiquette(s, x, y, w, texte):
    """Surtitre en capitales, 9 pt (plancher de lisibilité du deck V2)."""
    rich(s, x, y, w, 0.2, [(texte.upper(), {"size": 9, "bold": True, "color": MUTED})])


def cloture(s, y, h, texte, size=11.5):
    """Bandeau navy de clôture ; guillemet blanc (charte : jamais de texte cyan)."""
    forme(s, MSO_SHAPE.ROUNDED_RECTANGLE, X0, y, W, h, NAVY, adj=min(0.5, 0.1 / h))
    rich(s, X0 + 0.15, y, 0.4, h, [("“", {"size": 26, "bold": True, "color": WHITE})],
         anchor=MSO_ANCHOR.MIDDLE)
    rich(s, X0 + 0.65, y, W - 0.85, h, [(texte, {"size": size, "bold": True, "color": WHITE})],
         anchor=MSO_ANCHOR.MIDDLE)


def filet(s, x, y, w):
    forme(s, MSO_SHAPE.RECTANGLE, x, y, w, 0.012, GRIS)


def carte_titre(s, x, y, w, h, titre_, corps, ligne=GRIS, taille=10, hauteur_titre=0.0):
    carte(s, x, y, w, h, line=ligne)
    paras = [[(titre_, {"size": 11, "bold": True, "space_after": 3})]]
    for c in ([corps] if isinstance(corps, str) else corps):
        paras.append([(c, {"size": taille, "space_after": 2})])
    rich(s, x + 0.14, y + 0.1, w - 0.28, h - 0.18 - hauteur_titre, paras)


# ---------------------------------------------------------------- slides
def s01(s):
    """Promesse : couverture."""
    titre(s, TITRES[0])
    for sh in s.shapes:
        if sh.is_placeholder and sh.placeholder_format.idx == 0:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(18)
    x = 1.81
    for i, (nom, acc) in enumerate((("Infra", False), ("Produit", False), ("IA", True))):
        pilule(s, x + i * 1.05, 2.93, 0.95, 0.26, nom, fill=ACCENT if acc else NAVY,
               color=NAVY if acc else WHITE, size=9)
    t(s, x, 3.24, 4.4, 0.3, "orchestrés par des coachs produit augmentés", size=10,
      color=WHITE)


def s02(s):
    titre(s, TITRES[1])
    chapo(s, TOP, [
        gras("Beaucoup d'équipes Infra subissent la demande plus qu'elles ne pilotent la "
             "valeur.", "subissent la demande", size=11),
        [("Comment sortir du guichet sursollicité ? Comment mieux servir les utilisateurs ?",
          {"size": 10, "color": MUTED})]], h=0.5)
    items = [("Des délais qui s'allongent", "L'équipe Infra devient le goulot."),
             ("Un parcours illisible", "Des outils sans lien ; à qui s'adresser ?"),
             ("Des outils contournés", "Une plateforme existe, les équipes passent à côté."),
             ("Des coûts sans valeur visible", "Des coûts cloud qui montent sans lien avec la valeur.")]
    y, h = 2.25, 1.0
    for i, ((x, w), (a, b)) in enumerate(zip(cols(4), items, strict=True)):
        carte(s, x, y, w, h, line=NAVY if i == 0 else GRIS)
        badge(s, x + 0.3, y + 0.28, 0.34, str(i + 1), accent=i == 0, size=10)
        rich(s, x + 0.55, y + 0.08, w - 0.65, 0.45, [(a, {"size": 10.5, "bold": True})],
             anchor=MSO_ANCHOR.MIDDLE)
        rich(s, x + 0.14, y + 0.55, w - 0.28, 0.42, [(b, {"size": 10})])
    etiquette(s, X0, 3.38, W, "Le guichet : un besoin, plusieurs équipes à solliciter (exemple)")
    noms = ["Équipe applicative", "Sécurité", "Réseau", "Plateforme", "Livraison"]
    cw = (W + 0.1 * 4) / 5
    for i, n in enumerate(noms):
        chevron(s, X0 + i * (cw - 0.1), 3.6, cw, 0.36, n, accent=i == 0, size=10)
    etiquette(s, X0, 4.1, W, "Verbatims recueillis — illustratifs, pas une moyenne de marché")
    quotes = ["« Nos développeurs attendent trois semaines pour obtenir leur environnement. »",
              "« Nous avons une plateforme, mais personne ne l'utilise. »",
              "« Pas de vision globale : pas d'outils intégrés et simples. »"]
    for (x, w), q in zip(cols(3), quotes, strict=True):
        carte(s, x, 4.32, w, 0.64, line=GRIS)
        rich(s, x + 0.14, 4.32, w - 0.28, 0.64, [(q, {"size": 10, "italic": True})],
             anchor=MSO_ANCHOR.MIDDLE)


def s03(s):
    titre(s, TITRES[2])
    msg(s, "Sans évolution du fonctionnement Infra, les nouveaux usages IA peuvent amplifier "
        "une demande déjà difficile à absorber.", "peuvent amplifier", h=0.45)
    y, h, wc = 2.0, 1.55, 2.95
    carte(s, X0, y, wc, h, line=NAVY)
    rich(s, X0 + 0.15, y + 0.1, wc - 0.3, h - 0.2, [
        [("PRESSION DE CONSOMMATION", {"size": 9, "bold": True, "color": MUTED,
                                      "space_after": 5})],
        [("Run et legacy", {"size": 10.5, "bold": True, "space_after": 4})],
        [("Interdépendances et fiabilité", {"size": 10.5, "bold": True, "space_after": 4})],
        [("Coûts à relier à la valeur", {"size": 10.5, "bold": True, "space_after": 4})],
        [("Nouveaux usages IA et agents", {"size": 10.5, "bold": True})]])
    xr = XR - wc
    carte(s, xr, y, wc, h, line=NAVY)
    rich(s, xr + 0.15, y + 0.1, wc - 0.3, h - 0.2, [
        [("POTENTIEL D'AUGMENTATION", {"size": 9, "bold": True, "color": MUTED,
                                      "space_after": 5})],
        [("Analyser des signaux plus vite", {"size": 10.5, "bold": True, "space_after": 4})],
        [("Trier et qualifier les demandes", {"size": 10.5, "bold": True, "space_after": 4})],
        [("Documenter capacités et parcours", {"size": 10.5, "bold": True})]])
    xc = X0 + wc + 0.3
    wm = xr - 0.3 - xc
    forme(s, MSO_SHAPE.ROUNDED_RECTANGLE, xc, y + 0.2, wm, h - 0.4, NAVY, adj=0.1)
    rich(s, xc + 0.1, y + 0.2, wm - 0.2, h - 0.4, [
        [("Infra pour l'IA", {"size": 11.5, "bold": True, "color": WHITE, "space_after": 6})],
        [("et", {"size": 10, "color": WHITE, "space_after": 6})],
        [("IA pour l'Infra", {"size": 11.5, "bold": True, "color": WHITE})]],
        anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    fleche(s, X0 + wc + 0.05, y + h / 2 - 0.13, 0.2, 0.26)
    forme(s, MSO_SHAPE.CHEVRON, xr - 0.25, y + h / 2 - 0.13, 0.2, 0.26, NAVY, adj=0.5).rotation = 180
    stats = [("80 %", "des grandes organisations auront une platform team en 2026 — "
              + GARTNER_80 + ".", "Une platform team ne garantit pas l'adoption : contournements, "
              "délais et coûts restent à observer."),
             ("> 40 %", "des projets agentic seraient abandonnés d'ici 2027 — "
              + GARTNER_40 + ".", "D'où un cadre gouverné avant d'industrialiser.")]
    for (x, w), (g, a, b) in zip(cols(2), stats, strict=True):
        carte(s, x, 3.75, w, 1.2, line=GRIS)
        t(s, x + 0.15, 3.75, 1.0, 1.2, g, size=20, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        rich(s, x + 1.2, 3.8, w - 1.3, 1.1, [
            [(a, {"size": 10, "bold": True, "space_after": 3})], [(b, {"size": 10})]],
            anchor=MSO_ANCHOR.MIDDLE)


def s04(s):
    titre(s, TITRES[3])
    msg(s, "On ne transforme pas tout en produit : on pilote comme des produits les "
        "capacités qui le justifient.", "les capacités qui le justifient", h=0.45)
    piliers = [("Consommateurs", "Partir des équipes qui consomment l'Infra et de leurs "
                "douleurs mesurées."),
               ("Catalogue", "Des capacités lisibles, en self-service quand c'est pertinent, "
                "avec un niveau de service explicite."),
               ("Ownership", "Un ownership clair par produit ou capacité ; la forme du rôle "
                "dépend de l'organisation."),
               ("Mesure", "Les décisions combinent expérience, usage, fiabilité, risques, "
                "coût et valeur.")]
    for i, ((x, w), (a, b)) in enumerate(zip(cols(4), piliers, strict=True)):
        carte_titre(s, x, 2.0, w, 1.2, a, b, ligne=NAVY if i == 0 else GRIS)
    etiquette(s, X0, 3.58, W, "Exemple : environnement applicatif prêt à l'emploi")
    etapes = [("Consommateur", "Équipe de développement"),
              ("Capacité", "Environnement prêt à l'emploi"),
              ("Résultats observés", "Délai, autonomie, conformité")]
    for i, ((x, w), (a, b)) in enumerate(zip(cols(3, gap=0.0), etapes, strict=True)):
        chevron(s, x, 3.82, w + (0.12 if i < 2 else 0), 0.42, a, accent=i == 1, size=10.5)
        t(s, x + 0.2, 4.3, w - 0.4, 0.3, b, size=10, align=PP_ALIGN.CENTER)
    t(s, X0, 4.7, W, 0.25, "Le Run reste une responsabilité de service : tout ne devient "
      "pas produit.", size=10, italic=True, color=MUTED)


LEVIERS = [("Expérience", "« trop compliquée »", (0, 1)),
           ("Delivery", "« nous sommes trop lents »", (0, 2, 4)),
           ("Organisation", "« tout remonte à l'équipe centrale »", (0, 4)),
           ("Valeur", "« on dépense sans savoir pourquoi »", (0, 3)),
           ("Fiabilité", "« nos systèmes sont fragiles »", (1, 2)),
           ("Stratégie", "« un levier stratégique »", (3, 4))]
RESULTATS = ["Efficacité opérationnelle", "Expérience utilisateur", "Robustesse",
             "Maîtrise des coûts", "Anticipation de l'IA"]


def s05(s):
    titre(s, TITRES[4])
    msg(s, "Nous activons les bons leviers selon le diagnostic, pas un framework identique "
        "pour tous. Liens principaux, à confirmer.", "les bons leviers", h=0.45)
    wl = 2.9
    wc = (W - wl) / 5
    for j, r in enumerate(RESULTATS):
        pilule(s, X0 + wl + j * wc + 0.04, 2.0, wc - 0.08, 0.44, r, fill=NAVY, size=9.5)
    h = 0.4
    for i, (nom, constat, liens) in enumerate(LEVIERS):
        y = 2.55 + i * h
        filet(s, X0, y, W)
        rich(s, X0, y, wl, h, [[(f"{i + 1}  {nom}", {"size": 10.5, "bold": True}),
                                (f"  {constat}", {"size": 10, "color": MUTED})]],
             anchor=MSO_ANCHOR.MIDDLE)
        for j in liens:
            cx = X0 + wl + j * wc + wc / 2
            forme(s, MSO_SHAPE.OVAL, cx - 0.1, y + h / 2 - 0.1, 0.2, 0.2,
                  ACCENT if (i, j) == (0, 0) else NAVY, NAVY if (i, j) == (0, 0) else None,
                  lw=1.25)
    filet(s, X0, 2.55 + 6 * h, W)


def s06(s):
    titre(s, TITRES[5])
    msg(s, "Même objectif de valeur, mais un parcours d'accompagnement différent selon le "
        "point de départ.", "parcours d'accompagnement différent", h=0.45)
    etiquette(s, X0, 1.98, 2.5, "4 critères du diagnostic")
    crit = ["Maturité Produit", "Capacité IA et données", "Criticité opérationnelle",
            "Complexité de l'écosystème"]
    for i, c in enumerate(crit):
        y = 2.22 + i * 0.57
        carte(s, X0, y, 2.5, 0.49, line=NAVY if i == 0 else GRIS)
        badge(s, X0 + 0.28, y + 0.245, 0.3, str(i + 1), accent=i == 0, size=9)
        rich(s, X0 + 0.55, y, 1.9, 0.49, [(c, {"size": 10, "bold": True})],
             anchor=MSO_ANCHOR.MIDDLE)
    modes = [("Reveal & Simplify", "Run très réactif, demandes en guichet",
              "Repérer les demandes évitables, standardiser les parcours."),
             ("Service to Product", "Infra structurée en services et SLA",
              "Clarifier consommateur, promesse et ownership."),
             ("Product & AI Acceleration", "Plateforme ou produit déjà formalisé",
              "Renforcer discovery et mesure ; tester l'IA ciblée."),
             ("Agent-ready Infrastructure", "Infra très automatisée, exposée par API",
              "Gouverner des capacités consommables par des agents.")]
    xg = X0 + 2.5 + 0.3
    wg = (XR - xg - 0.14) / 2
    etiquette(s, xg, 1.98, 4, "4 modes d'entrée (noms de travail)")
    for k, (nom, sit, trav) in enumerate(modes):
        x = xg + (k % 2) * (wg + 0.14)
        y = 2.22 + (k // 2) * 1.2
        carte(s, x, y, wg, 1.08, line=NAVY if k == 0 else GRIS)
        rich(s, x + 0.14, y + 0.08, wg - 0.28, 0.95, [
            [(nom, {"size": 11, "bold": True, "space_after": 2})],
            [(sit, {"size": 10, "color": MUTED, "space_after": 2})],
            [(trav, {"size": 10})]])
    cloture(s, 4.57, 0.41, "Une segmentation de l'accompagnement, pas un classement de "
            "maturité : plusieurs situations coexistent, sans scoring calibré.", size=10)


def s07(s):
    titre(s, TITRES[6])
    msg(s, "Le coach exploite des signaux et orchestre des agents ; les décisions "
        "d'organisation, de priorité et de production restent humaines.",
        "restent humaines", h=0.5)
    wl = 1.5
    wc = (W - wl - 0.2) / 3
    moments = ["Avant une séance", "Pendant", "Entre deux séances"]
    for j, m in enumerate(moments):
        chevron(s, X0 + wl + 0.1 + j * (wc + 0.0), 2.0, wc + 0.1, 0.4, m, accent=j == 1,
                size=10)
    lignes = [("COACH", NAVY, WHITE,
               ["Choisit questions, sources et hypothèses ; vérifie le contexte.",
                "Conduit les entretiens, facilite les arbitrages, teste les hypothèses.",
                "Interprète les résultats, lève les obstacles, développe l'autonomie."]),
              ("AGENTS DU COACH", ACCENT, NAVY,
               ["Catégorisent les signaux des sources autorisées ; préparent la trame.",
                "Restituent des faits sourcés ; ne tranchent pas.",
                "Capitalisent les décisions validées et consolident les indicateurs."]),
              ("CLIENT", TRACK, NAVY,
               ["Autorise les sources et précise ses besoins.",
                "Partage la réalité opérationnelle et décide.",
                "Expérimente, mesure, ajuste ses pratiques."])]
    for i, (lab, fill, coul, cells) in enumerate(lignes):
        y = 2.5 + i * 0.67
        pilule(s, X0, y + 0.17, wl, 0.28, lab, fill=fill, color=coul, size=9)
        for j, c in enumerate(cells):
            x = X0 + wl + 0.1 + j * wc
            carte(s, x + 0.05, y, wc - 0.1, 0.61, line=GRIS)
            rich(s, x + 0.15, y, wc - 0.3, 0.61, [(c, {"size": 10})], anchor=MSO_ANCHOR.MIDDLE)
    cloture(s, 4.55, 0.43, "Une transformation Produit + IA contextualisée, opérée par des "
            "coachs augmentés et mesurée dans la durée.", size=10)


def s08(s):
    titre(s, TITRES[7])
    import math
    cx, cy, rx, ry = 2.85, 2.72, 1.6, 0.95
    forme(s, MSO_SHAPE.OVAL, cx - rx, cy - ry, 2 * rx, 2 * ry, None, GRIS, lw=3)
    etapes = ["Explorer", "Comprendre", "Prioriser", "Expérimenter", "Mesurer", "Ancrer"]
    for k, nom in enumerate(etapes):
        a = math.radians(-90 + 60 * k)
        px, py = cx + rx * math.cos(a), cy + ry * math.sin(a)
        forme(s, MSO_SHAPE.ROUNDED_RECTANGLE, px - 0.62, py - 0.2, 1.24, 0.4,
              ACCENT if k == 0 else NAVY, adj=0.5)
        rich(s, px - 0.58, py - 0.2, 1.16, 0.4, [(f"{k + 1}  {nom}", {
            "size": 10, "bold": True, "color": NAVY if k == 0 else WHITE})],
            anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    rich(s, cx - 0.7, cy - 0.3, 1.4, 0.6, [
        [("puis nouveau", {"size": 10, "color": MUTED})],
        [("cycle", {"size": 10, "color": MUTED})]], anchor=MSO_ANCHOR.MIDDLE,
        align=PP_ALIGN.CENTER)
    xr = 5.15
    wr = XR - xr
    carte(s, xr, 1.5, wr, 1.45, line=NAVY)
    for k, (a, wp, xp) in enumerate((("Assainissement", 1.3, 0.15), ("Évolution technique", 1.5, 1.5),
                                     ("Formation", 0.8, 3.1))):
        pilule(s, xr + xp, 2.45, wp, 0.3, a, fill=ACCENT if k == 0 else NAVY,
               color=NAVY if k == 0 else WHITE, size=9)
    rich(s, xr + 0.15, 1.58, wr - 0.3, 0.8, [
        [("Actions possibles dans la boucle", {"size": 11, "bold": True, "space_after": 4})],
        [("Des leviers choisis par le diagnostic, pas un préalable :", {"size": 10})]])
    carte(s, xr, 3.05, wr, 0.95, line=GRIS)
    rich(s, xr + 0.15, 3.1, wr - 0.3, 0.85, [
        [("« Repérer, prioriser, tenir »", {"size": 10.5, "bold": True, "space_after": 2})],
        [("La logique simple d'un pilote, à l'intérieur de la boucle : pas une deuxième "
          "méthode.", {"size": 10})]], anchor=MSO_ANCHOR.MIDDLE)
    fils = [("Personnes", "engager, expérimenter, relayer"),
            ("Technologies et Run", "cartographier, standardiser, décommissionner si pertinent"),
            ("IA", "assister, recommander, déléguer sous contrôle")]
    etiquette(s, X0, 4.17, W, "Trois fils transverses")
    for i, ((x, w), (a, b)) in enumerate(zip(cols(3), fils, strict=True)):
        carte(s, x, 4.4, w, 0.58, line=NAVY if i == 0 else GRIS)
        rich(s, x + 0.12, 4.4, w - 0.24, 0.58, [[(a, {"size": 10, "bold": True})],
             [(b, {"size": 10})]], anchor=MSO_ANCHOR.MIDDLE)


def s09(s):
    titre(s, TITRES[8])
    msg(s, "Les agents augmentent d'abord notre accompagnement ; certains cas d'usage "
        "peuvent ensuite augmenter les activités du client.", "augmentent d'abord", h=0.45)
    wc = 3.85
    xr = XR - wc
    y, h = 2.15, 1.85
    carte(s, X0, y, wc, h, line=NAVY, lw=1.5)
    carte(s, xr, y, wc, h, line=GRIS)
    pilule(s, X0 + 0.15, y - 0.13, 2.7, 0.26, "AGENTS DU COACH — DISPONIBLES", size=9)
    pilule(s, xr + 0.15, y - 0.13, 3.0, 0.26, "AGENTS DU CLIENT — EXEMPLES", fill=TRACK,
           color=NAVY, size=9)
    agents = [("Infra Signal Miner", "repère irritants et opportunités dans les signaux"),
              ("Product Discovery Agent", "prépare la compréhension des consommateurs"),
              ("Value Intelligence Agent", "prépare arbitrages de valeur et mesure"),
              ("Coaching Companion", "prépare, documente et capitalise le travail du coach")]
    rich(s, X0 + 0.15, y + 0.2, wc - 0.3, h - 0.25,
         [[(a, {"size": 10.5, "bold": True}), (" — " + b, {"size": 10})]
          for a, b in agents])
    clients = [("Triage de tickets", "classe et route des demandes bornées"),
               ("Veille FinOps", "propose une liste à valider, ne décommissionne jamais seul"),
               ("Assistant documentaire", "répond avec la source citée")]
    rich(s, xr + 0.15, y + 0.2, wc - 0.3, h - 0.25,
         [[(a, {"size": 10.5, "bold": True}), (" — " + b, {"size": 10})]
          for a, b in clients] + [[("Chantier distinct, propre à chaque client.",
                                    {"size": 10, "italic": True, "color": MUTED})]])
    forme(s, MSO_SHAPE.RECTANGLE, (X0 + wc + xr) / 2 - 0.015, y - 0.1, 0.03, h + 0.1, NAVY)
    for x, txt_ in ((X0, "Le coach humain valide chaque sortie avant usage."),
                    (xr, "Les équipes opérationnelles restent responsables de leurs agents.")):
        carte(s, x, 4.45, wc, 0.5, line=GRIS)
        rich(s, x + 0.15, 4.45, wc - 0.3, 0.5, [(txt_, {"size": 10, "bold": True})],
             anchor=MSO_ANCHOR.MIDDLE)
    t(s, X0, 4.12, W, 0.2, "Une frontière d'accès aux données sépare les deux univers.",
      size=10, italic=True, color=MUTED, align=PP_ALIGN.CENTER)


def s10(s):
    titre(s, TITRES[9])
    msg(s, "Le dispositif agentique respecte la criticité et les politiques du client.",
        "criticité et les politiques", h=0.4)
    gates = [("Données", ["Sources autorisées", "Classification", "Minimisation"]),
             ("Autorisations", ["Lecture seule par défaut", "Droits bornés par agent"]),
             ("Qualité", ["Évaluation avant usage", "Règles et limites testées avant toute "
                          "délégation"]),
             ("Supervision", ["Validation humaine des décisions", "Traçabilité",
                              "Réversibilité et arrêt"])]
    gap = 0.3
    gw = (W - 3 * gap) / 4
    for i, (nom, pts) in enumerate(gates):
        x = X0 + i * (gw + gap)
        carte(s, x, 2.15, gw, 1.45, line=NAVY if i == 0 else GRIS)
        badge(s, x + 0.28, 2.15, 0.4, str(i + 1), accent=i == 0, size=11)
        rich(s, x + 0.14, 2.42, gw - 0.28, 1.15, [[(nom, {"size": 11, "bold": True,
             "space_after": 4})]] + [[("– " + p, {"size": 10, "space_after": 2})] for p in pts])
        if i < 3:
            fleche(s, x + gw + 0.05, 2.75, 0.2, 0.26)
    etiquette(s, X0, 3.74, W, "Quatre niveaux d'autonomie de l'IA")
    niveaux = ["Assistance", "Recommandation", "Exécution bornée", "Action Infra validée"]
    cw = (W + 0.1 * 3) / 4
    for i, n in enumerate(niveaux):
        chevron(s, X0 + i * (cw - 0.1), 3.95, cw, 0.42, n, accent=i == 3, size=10)
    cloture(s, 4.5, 0.46, "L'IA assiste et recommande ; des tâches bornées peuvent être "
            "déléguées sous contrôle ; les arbitrages structurants restent humains.", size=10)


def s11(s):
    titre(s, TITRES[10])
    pilule(s, X0, TOP, 4.3, 0.28, ETIQUETTE_SCENARIO.upper(), fill=ACCENT, color=NAVY, size=9)
    t(s, X0 + 4.45, TOP, W - 4.45, 0.4, "Connectivité applicative sécurisée — T0 et T1 à "
      "mesurer, aucune valeur réelle", size=10, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
    etapes = ["Problème et T0", "Signaux et interviews", "Hypothèse produit",
              "Expérimentation", "Mesure au T1", "Décision"]
    cw = (W + 0.1 * 5) / 6
    for i, e in enumerate(etapes):
        chevron(s, X0 + i * (cw - 0.1), 1.9, cw, 0.5, e, accent=i == 0, size=9.5)
    coach = ["Qualifie le périmètre, fixe la référence.",
             "Conduit les interviews.",
             "Facilite le cadrage du produit.",
             "Accompagne l'équipe client.",
             "Interprète les résultats.",
             "Éclaire l'arbitrage du client."]
    agents = ["Repèrent les motifs récurrents.",
              "Synthétisent avec les sources.",
              "Préparent des alternatives.",
              "Aucune action en production.",
              "Consolident les faits.",
              "Capitalisent la décision validée."]
    for (lab, fill, coul, cells, y) in (
            ("CE QUE FAIT LE COACH", NAVY, WHITE, coach, 2.5),
            ("CE QUE PRÉPARENT LES AGENTS", TRACK, NAVY, agents, 3.6)):
        pilule(s, X0, y, 2.6, 0.26, lab, fill=fill, color=coul, size=9)
        for i, c in enumerate(cells):
            x = X0 + i * (cw - 0.1) + 0.04
            carte(s, x, y + 0.34, cw - 0.2, 0.62, line=GRIS)
            rich(s, x + 0.1, y + 0.34, cw - 0.4, 0.62, [(c, {"size": 10})],
                 anchor=MSO_ANCHOR.MIDDLE)
    t(s, X0, 4.7, W, 0.26, "Le client décide : généraliser, adapter ou arrêter.", size=10.5,
      bold=True, align=PP_ALIGN.CENTER)


def s12(s):
    titre(s, TITRES[11])
    msg(s, "Référence documentée au T0, résultats et écarts évalués au T1 ; la contribution "
        "de l'IA est vérifiée, jamais présumée.", "vérifiée, jamais présumée", h=0.45)
    xa, xb, xc = X0, X0 + 2.05, X0 + 5.95
    wa, wb, wc = 1.95, 3.8, XR - xc
    pilule(s, xa, 1.95, 1.9, 0.24, "NIVEAU 1 · CLIENT", size=9)
    t(s, xb + 0.0, 1.95, wb, 0.24, "Indicateurs possibles (T0 → T1)", size=9, bold=True,
      color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
    t(s, xc, 1.95, wc, 0.24, "Garde-fou", size=9, bold=True, color=MUTED,
      anchor=MSO_ANCHOR.MIDDLE)
    lignes1 = [
        ("Efficacité opérationnelle", "Délai de mise à disposition, demandes complètes",
         "Ne pas compter un simple déplacement de charge"),
        ("Expérience utilisateur", "Effort perçu, succès du self-service",
         "Échantillon représentatif, clients captifs"),
        ("Robustesse et risque", "SLO, incidents liés aux changements",
         "Ne jamais sacrifier la fiabilité à la vitesse"),
        ("Coûts et valeur", "Coût unitaire par capacité utile, prévisibilité",
         "Distinguer économies, coûts évités, temps libéré"),
        ("Anticipation de l'IA", "Délai d'accès à une capacité IA gouvernée",
         "Pas de métrique IA hors périmètre concerné")]
    lignes2 = [("Efficacité du coaching augmenté", "Temps de préparation, preuves sourcées, "
                "délai signal → arbitrage, autonomie transférée",
                "Vérifier la contribution de l'IA")]

    def rangee(y, h, a, b, c):
        filet(s, X0, y, W)
        t(s, xa, y, wa, h, a, size=10, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        t(s, xb, y, wb - 0.1, h, b, size=10, anchor=MSO_ANCHOR.MIDDLE)
        t(s, xc, y, wc, h, c, size=10, anchor=MSO_ANCHOR.MIDDLE)
    y = 2.25
    for a, b, c in lignes1:
        rangee(y, 0.38, a, b, c)
        y += 0.38
    pilule(s, xa, y + 0.07, 2.8, 0.24, "NIVEAU 2 · COACHING AUGMENTÉ", fill=ACCENT, color=NAVY,
           size=9)
    y += 0.32
    for a, b, c in lignes2:
        rangee(y, 0.48, a, b, c)
        y += 0.48
    filet(s, X0, y, W)


def s13(s):
    titre(s, TITRES[12])
    etapes = ["Sources autorisées", "Agents du coach", "Coach augmenté + experts",
              "Décisions du client", "Mesure et transfert"]
    cw = (W + 0.1 * 4) / 5
    for i, e in enumerate(etapes):
        chevron(s, X0 + i * (cw - 0.1), TOP, cw, 0.5, e, accent=i == 2, size=10)
    t(s, X0, 2.0, W, 0.22, "Aucun raccourci « agent → action en production » dans la mission "
      "de conseil.", size=10, italic=True, color=MUTED)
    y, h = 2.4, 2.35
    wc = (W - 0.4) / 3
    xs = [X0 + i * (wc + 0.2) for i in range(3)]
    carte(s, xs[0], y, wc, h, line=NAVY)
    pilule(s, xs[0] + 0.15, y - 0.13, 1.1, 0.26, "CÔTÉ OCTO", size=9)
    roles = [("CP", "Coach produit / plateforme augmenté", "porte la démarche et les rituels"),
             ("EP", "Expert plateforme et Infra", "diagnostic, catalogue, automatisation"),
             ("IA", "Expert IA, selon le besoin", "cadre et outille les usages")]
    for i, (ini, a, b) in enumerate(roles):
        yy = y + 0.3 + i * 0.7
        badge(s, xs[0] + 0.4, yy + 0.3, 0.46, ini, accent=i == 0, size=9)
        rich(s, xs[0] + 0.8, yy, wc - 0.9, 0.6, [[(a, {"size": 10, "bold": True})],
             [(b, {"size": 10})]], anchor=MSO_ANCHOR.MIDDLE)
    carte(s, xs[1], y, wc, h, line=GRIS)
    pilule(s, xs[1] + 0.15, y - 0.13, 1.3, 0.26, "CÔTÉ CLIENT", fill=TRACK, color=NAVY, size=9)
    rich(s, xs[1] + 0.15, y + 0.25, wc - 0.3, h - 0.3, [
        [("Un sponsor", {"size": 10, "bold": True}),
         (" : DSI ou direction Infra.", {"size": 10, "space_after": 6})],
        [("Des propriétaires", {"size": 10, "bold": True}),
         (" de produits et de capacités.", {"size": 10, "space_after": 6})],
        [("Des experts", {"size": 10, "bold": True}),
         (" Infra, applicatifs, sécurité.", {"size": 10, "space_after": 6})],
        [("Décisions, rituels, accès : ", {"size": 10, "bold": True}),
         ("répartition arbitrée ensemble.", {"size": 10})]])
    carte(s, xs[2], y, wc, h, line=NAVY)
    pilule(s, xs[2] + 0.15, y - 0.13, 1.9, 0.26, "CRITÈRES D'AUTONOMIE", fill=ACCENT,
           color=NAVY, size=9)
    rich(s, xs[2] + 0.15, y + 0.25, wc - 0.3, h - 0.3, [
        [("1  Le client lit et tient ses indicateurs.", {"size": 10, "space_after": 6})],
        [("2  Il anime lui-même discovery et arbitrages.", {"size": 10, "space_after": 6})],
        [("3  Des relais internes formés portent les rituels.", {"size": 10,
                                                                 "space_after": 8})],
        [("Effort attendu côté client et dimensionnement : à compléter selon le périmètre.",
          {"size": 10, "italic": True, "color": MUTED})]])


def s14(s):
    titre(s, TITRES[13])
    y1, h1 = 1.5, 1.0
    carte(s, X0, y1, W, h1, line=NAVY, lw=1.5)
    pilule(s, X0 + 0.15, y1 - 0.13, 0.9, 0.26, "TEMPS 1", size=9)
    pilule(s, XR - 2.1, y1 - 0.13, 1.95, 0.26, "MONTANT : À COMPLÉTER", fill=TRACK,
           color=NAVY, size=9)
    rich(s, X0 + 0.2, y1 + 0.2, W - 0.4, h1 - 0.25, [
        [("Assessment flash — forfait d'entrée", {"size": 12, "bold": True, "space_after": 4})],
        [("Diagnostic et cartographie des opportunités, baseline documentée, coût du statu quo "
          "estimé quand il est calculable, choix d'un pilote. Aucun ROI promis.",
          {"size": 10})]])
    pilule(s, X0 + 2.0, 2.62, W - 4.0, 0.26, "JALON : RESTITUTION ET ARBITRAGE DU SPONSOR",
           fill=ACCENT, color=NAVY, size=9)
    y2, h2 = 3.1, 1.4
    carte(s, X0, y2, W, h2, line=GRIS)
    pilule(s, X0 + 0.15, y2 - 0.13, 0.9, 0.26, "TEMPS 2", size=9)
    pilule(s, XR - 2.1, y2 - 0.13, 1.95, 0.26, "MONTANT : À COMPLÉTER", fill=TRACK,
           color=NAVY, size=9)
    rich(s, X0 + 0.2, y2 + 0.15, W - 0.4, 0.75, [
        [("Trajectoire modulaire", {"size": 12, "bold": True, "space_after": 3})],
        [("Dimensionnée à l'issue de l'Assessment flash : on n'enchaîne que ce que le pilote "
          "justifie. L'assainissement est déclenché par le diagnostic.", {"size": 10})]])
    opts = ["Pilote Produit + IA", "Ancrage", "Extension éventuelle"]
    cw = (W - 0.4 + 0.1 * 2) / 3
    for i, o in enumerate(opts):
        chevron(s, X0 + 0.2 + i * (cw - 0.1), y2 + 0.92, cw, 0.36, o, accent=i == 0, size=10)
    t(s, X0, 4.65, W, 0.3, "Financement : modalités budgétaires convenues avec le client "
      "selon l'objet du travail.", size=10, italic=True, color=MUTED)


def s15(s):
    titre(s, TITRES[14])
    msg(s, "Choisir un périmètre prioritaire, mesurer une situation initiale et lancer une "
        "première boucle de transformation.", "première boucle", h=0.45)
    steps = [("1  Choisir", "Un sponsor et des objectifs clairs pour un périmètre utile."),
             ("2  Diagnostiquer", "Accès aux personnes et aux données autorisées ; référence "
              "mesurée au T0."),
             ("3  Tester", "Une capacité dédiée aux décisions et à l'expérimentation.")]
    for i, ((x, w), (a, b)) in enumerate(zip(cols(3, gap=0.0), steps, strict=True)):
        chevron(s, x, 2.05, w + (0.12 if i < 2 else 0), 0.6, a, accent=i == 0, size=12)
        carte(s, x + 0.08, 2.8, w - 0.16, 1.15, line=NAVY if i == 0 else GRIS)
        rich(s, x + 0.22, 2.8, w - 0.44, 1.15, [(b, {"size": 10.5})], anchor=MSO_ANCHOR.MIDDLE)
    cloture(s, 4.2, 0.7, "Première décision commerciale : la restitution du diagnostic, "
            "avec des critères de go / no-go partagés.", size=11)


def a_problemes(s):
    titre(s, TITRES_ANNEXES[0])
    lignes = [
        ("L'infra est perçue comme un centre de coûts", "Clarifier la proposition de valeur",
         "Valeur délivrée par capacité"),
        ("Les utilisateurs ne comprennent pas l'offre", "Formaliser un catalogue de capacités",
         "Demandes couvertes par le catalogue"),
        ("Les demandes sont trop spécifiques", "Identifier les besoins récurrents",
         "Part de demandes standardisées"),
        ("Les outils sont peu adoptés", "Travailler l'UX et l'adoption", "Taux d'adoption"),
        ("Les équipes applicatives contournent l'infra", "Rendre l'offre simple et désirable",
         "Usages hors plateforme"),
        ("Les priorités sont politiques", "Prioriser par valeur, usage, risque, adoption",
         "Backlog priorisé sur critères partagés"),
        ("Les capacités livrées ne sont pas utilisées", "Mesurer l'usage réel",
         "Usage par capacité livrée"),
        ("L'automatisation manque de financement", "Construire un business case produit",
         "Capacité Run libérée, puis réallouée"),
        ("L'obsolescence est subie", "Piloter le cycle de vie des capacités",
         "Décommissionnements tenus"),
        ("La plateforme est trop technique", "Définir des parcours compréhensibles",
         "Délai de mise à disposition")]
    xs = [X0, X0 + 3.15, X0 + 6.2]
    ws = [3.05, 2.95, XR - xs[2]]
    for x, w, h_ in zip(xs, ws, ("Le problème", "L'apport de l'approche produit",
                                 "L'indicateur à suivre"), strict=True):
        t(s, x, 1.3, w, 0.22, h_.upper(), size=9, bold=True, color=MUTED)
    for i, ligne in enumerate(lignes):
        y = 1.55 + i * 0.34
        filet(s, X0, y, W)
        for k, (x, w, c) in enumerate(zip(xs, ws, ligne, strict=True)):
            t(s, x, y, w - 0.08, 0.34, c, size=10, bold=k == 0, anchor=MSO_ANCHOR.MIDDLE)
    filet(s, X0, 1.55 + 10 * 0.34, W)


def a_alternatives(s):
    titre(s, TITRES_ANNEXES[1])
    lignes = [
        ("Ne rien faire", "Aucun investissement immédiat.",
         "Rendre le coût du statu quo explicite et fondé sur des données, sans le supposer "
         "systématiquement croissant."),
        ("FinOps", "Un complément efficace pour piloter dépenses et valeur économique.",
         "Ajoute parcours, ownership et arbitrages produit lorsqu'ils sont nécessaires."),
        ("Platform engineering", "Un ensemble pertinent de pratiques et de plateformes.",
         "Ne couvre pas mécaniquement toutes les entités Infra ni l'accompagnement du "
         "changement : l'offre y ajoute coaching et adaptation au contexte."),
        ("AIOps et agents opérationnels", "Un levier d'amélioration des opérations.",
         "Se distingue par le diagnostic produit, le coaching, l'organisation et la mesure.")]
    xs = [X0, X0 + 2.2, X0 + 5.0]
    ws = [2.1, 2.7, XR - xs[2]]
    for x, w, h_ in zip(xs, ws, ("Alternative", "Ce qu'elle apporte",
                                 "Ce que l'offre ajoute"), strict=True):
        t(s, x, 1.3, w, 0.22, h_.upper(), size=9, bold=True, color=MUTED)
    for i, ligne in enumerate(lignes):
        y = 1.6 + i * 0.7
        filet(s, X0, y, W)
        for k, (x, w, c) in enumerate(zip(xs, ws, ligne, strict=True)):
            t(s, x, y, w - 0.1, 0.7, c, size=10.5 if k == 0 else 10, bold=k == 0,
              anchor=MSO_ANCHOR.MIDDLE)
    filet(s, X0, 1.6 + 4 * 0.7, W)
    cloture(s, 4.6, 0.4, "Des disciplines complémentaires, intégrables dans une transformation "
            "orientée valeur.", size=10)


# ---------------------------------------------------------------- slides neuves
def synthese(s):
    """Synthèse en une page : cinq blocs courts (constat, cible, pourquoi, démarche, résultat)."""
    titre(s, TITRE_SYNTHESE)
    rows = [
        ("Le constat", "Le guichet Infra est saturé : délais qui s'allongent, parcours "
         "illisible, outils contournés, coûts sans valeur visible.", "guichet Infra est saturé"),
        ("La cible", "Infra as a Product : des consommateurs identifiés, un catalogue de "
         "capacités, un propriétaire et une valeur mesurée.", "Infra as a Product"),
        ("Pourquoi maintenant", "La pression monte et l'IA peut amplifier une demande déjà "
         "difficile à absorber : mieux vaut un cadre gouverné d'abord.", "cadre gouverné"),
        ("La démarche", "Une boucle continue menée par des coachs produit augmentés : un "
         "Assessment flash, puis expérimenter et déployer.", "coachs produit augmentés"),
        ("Ce que vous obtenez", "Des décisions fondées sur des faits, une valeur mesurée et "
         "des équipes qui portent le modèle sans nous.", "valeur mesurée"),
    ]
    h, gap, y0 = 0.6, 0.1, 1.42
    for i, (lab, corps, cle) in enumerate(rows):
        y = y0 + i * (h + gap)
        carte(s, X0 + 1.0, y, W - 1.0, h)
        pilule(s, X0, y + 0.12, 2.0, h - 0.24, lab, fill=ACCENT if i == 0 else NAVY,
               color=NAVY if i == 0 else WHITE, size=10)
        rich(s, X0 + 2.2, y, W - 2.35, h, [gras(corps, cle, size=10.5)],
             anchor=MSO_ANCHOR.MIDDLE)


# Convictions : texte du V1 revu (arbitrage de l'utilisateur). `principe`, `critere` et
# `chantier` sont une PROPOSITION de notre part (affichée comme telle), pas un acquis.
# `etape` : étape de la boucle de la slide « Une boucle continue... » (S08), citée par nom.
CONVICTIONS = [
    {"groupe": "cible", "court": "Partir des utilisateurs et de leurs douleurs",
     "titre": "Partir des utilisateurs et de leurs douleurs, pas du catalogue d'outils",
     "detail": "Une plateforme n'est adoptée que si elle résout une douleur mesurée chez "
               "ceux qui la consomment.",
     "etape": "Explorer",
     "principe": "Les consommateurs sont identifiés ; leurs douleurs sont mesurées avant "
                 "tout choix de catalogue.",
     "critere": "Les principales douleurs de chaque consommateur sont documentées, avec "
                "une mesure.",
     "chantier": "Interviews des consommateurs et situation de référence mesurée."},
    {"groupe": "cible", "court": "Un propriétaire, une roadmap, une valeur mesurée",
     "titre": "Un propriétaire, une roadmap, une valeur mesurée pour chaque capacité",
     "detail": "Sans porteur produit, l'infra retourne au guichet.",
     "etape": "Ancrer",
     "principe": "Chaque capacité clé a un porteur produit, une trajectoire et une valeur "
                 "suivie.",
     "critere": "Les capacités clés ont un propriétaire nommé, une roadmap et un "
                "indicateur de valeur.",
     "chantier": "Désigner les propriétaires, bâtir la roadmap, définir l'indicateur de "
                 "valeur."},
    {"groupe": "demarche", "court": "Décider sur des mesures et des critères partagés",
     "titre": "Décider sur des mesures et des critères partagés",
     "detail": "Choisir les indicateurs à mesurer ; nommer, scorer, prioriser (valeur, "
               "coût, complexité) ; des décisions ritualisées, reliées au budget.",
     "etape": "Prioriser",
     "principe": "Les priorités se décident sur des faits mesurés, pas sur le rapport de "
                 "force.",
     "critere": "Une grille de priorisation partagée existe et alimente les arbitrages "
                "budgétaires.",
     "chantier": "Choisir les indicateurs, scorer et prioriser, ritualiser la décision."},
    {"groupe": "demarche", "court": "S'améliorer en continu",
     "titre": "S'améliorer en continu, au rythme des saisons",
     "detail": "Mesure systématique et transparente des progrès, à tous les niveaux. Des "
               "équipes redevables de leurs améliorations ; succès et erreurs partagés.",
     "etape": "Mesurer",
     "principe": "Le progrès se mesure de façon régulière et transparente, à chaque niveau.",
     "critere": "Des indicateurs de progrès sont suivis à rythme régulier et partagés, "
                "succès comme erreurs.",
     "chantier": "Installer la mesure et un rituel de revue qui rend les équipes "
                 "redevables."},
    {"groupe": "demarche", "court": "L'IA gagne du temps là où il est perdu",
     "titre": "L'IA fait gagner du temps là où il est perdu, pas là où l'on décide",
     "detail": "Elle accélère le diagnostic, le tri des demandes et la restitution : le "
               "temps du Run est rendu aux équipes, mesuré puis vérifié, pas promis.",
     "etape": "Comprendre",
     "principe": "L'IA assiste le diagnostic et le triage ; les décisions restent "
                 "humaines. Jamais un prérequis.",
     "critere": "Les gisements de temps perdu sont repérés et les données utilisables "
                "sont autorisées.",
     "chantier": "Un cas d'usage borné, après le contrôle de confidentialité, au gain "
                 "mesuré puis vérifié."},
]
GROUPES = [("cible", "LA CIBLE · CE QUE LE CLIENT CONSTRUIT"),
           ("demarche", "LA DÉMARCHE · COMMENT NOUS Y CONDUISONS")]


def annonce_convictions(s):
    """Annonce : les convictions regroupées en cible et démarche, avec leur étape de boucle."""
    titre(s, TITRE_CONVICTIONS)
    msg(s, "Chaque conviction se retrouve dans une étape de la boucle : c'est là qu'elle "
        "devient un chantier.", "une étape de la boucle", h=0.4)
    y0, wl = 2.15, 3.95
    xr = X0 + wl + 0.25
    boites = {"cible": (X0, wl), "demarche": (xr, XR - xr)}
    n = 0
    for gid, label in GROUPES:
        x, w = boites[gid]
        mine = [c for c in CONVICTIONS if c["groupe"] == gid]
        carte(s, x, y0, w, 0.2 + 0.78 * len(mine), line=NAVY if gid == "cible" else GRIS)
        pilule(s, x + 0.15, y0 - 0.13, 3.2 if gid == "cible" else 3.55, 0.26, label,
               fill=NAVY if gid == "cible" else TRACK,
               color=WHITE if gid == "cible" else NAVY, size=9)
        for k, c in enumerate(mine):
            n += 1
            y = y0 + 0.12 + k * 0.78
            badge(s, x + 0.3, y + 0.39, 0.34, str(n), accent=n == 1, size=10)
            rich(s, x + 0.6, y, w - 0.75, 0.78, [
                [(c["court"], {"size": 10.5, "bold": True, "space_after": 2})],
                [("dans la démarche : ", {"size": 10, "color": MUTED}),
                 (c["etape"], {"size": 10, "bold": True})]], anchor=MSO_ANCHOR.MIDDLE)
    yp = y0 + 0.2 + 0.78 * 2 + 0.14
    carte(s, X0, yp, wl, 4.69 - yp, line=GRIS)
    pilule(s, X0 + 0.15, yp - 0.13, 1.1, 0.26, "PROPOSITION", fill=ACCENT, color=NAVY, size=9)
    rich(s, X0 + 0.15, yp + 0.15, wl - 0.3, 4.69 - yp - 0.2,
         [("Ce regroupement et le lien avec la boucle sont notre proposition, à valider "
           "avec vous.", {"size": 10})], anchor=MSO_ANCHOR.MIDDLE)


def focus(s, conv):
    """Focus d'une conviction : conviction, principe, critère au diagnostic, chantier."""
    titre(s, conv["titre"])
    y, h = 1.65, 2.0
    colonnes = cols(4, gap=0.3)
    cartes = [("LA CONVICTION", conv["detail"], False),
              ("LE PRINCIPE QU'ELLE FONDE", conv["principe"], True),
              ("LE CRITÈRE OBSERVABLE AU DIAGNOSTIC", conv["critere"], True),
              ("LE CHANTIER SI LE CRITÈRE MANQUE", conv["chantier"], True)]
    for i, ((x, w), (lab, corps, hypothese)) in enumerate(zip(colonnes, cartes, strict=True)):
        carte(s, x, y, w, h, line=NAVY if i == 0 else GRIS)
        if hypothese:
            pilule(s, x + 0.15, y - 0.13, 1.1, 0.26, "PROPOSITION", fill=TRACK, color=NAVY,
                   size=9)
        rich(s, x + 0.14, y + 0.18, w - 0.28, 0.42,
             [(lab, {"size": 9, "bold": True, "color": MUTED})])
        rich(s, x + 0.14, y + 0.62, w - 0.28, h - 0.7, [(corps, {"size": 10})])
        if i < 3:
            fleche(s, x + w + 0.05, y + h / 2 - 0.13)
    (xa, wa), (xb, wb) = cols(2)
    yb, hb = 3.98, 0.85
    carte(s, xa, yb, wa, hb, line=GRIS)
    rich(s, xa + 0.15, yb, wa - 0.3, hb, [
        [("DANS LA DÉMARCHE", {"size": 9, "bold": True, "color": MUTED, "space_after": 3})],
        [("Étape « " + conv["etape"] + " » de la boucle continue",
          {"size": 11, "bold": True})]], anchor=MSO_ANCHOR.MIDDLE)
    carte(s, xb, yb, wb, hb, line=GRIS)
    rich(s, xb + 0.15, yb, wb - 0.3, hb, [
        [("Principe, critère et chantier sont notre hypothèse de travail, à valider avec "
          "le client avant toute décision.", {"size": 10})]], anchor=MSO_ANCHOR.MIDDLE)


def equipe(s):
    titre(s, TITRE_EQUIPE)
    gens = [("Gaëlle MILOCHEVITCH", "MILE"), ("Claude CAMUS", "CACL"),
            ("Lan LEVY", "LLE"), ("Cédric CHEVALÉRIAS", "CECH")]
    for i, ((x, w), (nom, tri)) in enumerate(zip(cols(4), gens, strict=True)):
        carte(s, x, 1.9, w, 2.1, line=NAVY if i == 0 else GRIS)
        badge(s, x + w / 2, 2.55, 0.9, tri[:2], accent=i == 0, size=16)
        rich(s, x + 0.1, 3.15, w - 0.2, 0.8, [
            [(nom, {"size": 10.5, "bold": True})],
            [(f"({tri})", {"size": 10, "color": MUTED, "space_after": 3})],
            [("rôle : à fournir", {"size": 10, "italic": True, "color": MUTED})]],
            align=PP_ALIGN.CENTER)


# ---------------------------------------------------------------- chapitres
@dataclass(frozen=True)
class Page:
    titre: str
    fabrique: object = None
    genre: str = "contenu"          # couverture | sommaire | contenu


@dataclass(frozen=True)
class Chapitre:
    nom: str
    pages: tuple
    intercalaire: bool = True
    lettre: bool = False            # annexes : lettre (initiale du nom) au lieu d'un numéro


# SOURCE DE VÉRITÉ du squelette : sommaire, intercalaires et bandeaux en dérivent.
CHAPITRES = [
    Chapitre("Ouverture", (Page(TITRES[0], s01, "couverture"), Page(TITRE_SYNTHESE, synthese),
                           Page(TITRE_SOMMAIRE, None, "sommaire")), intercalaire=False),
    Chapitre("Contexte et enjeux", (Page(TITRES[1], s02), Page(TITRES[2], s03),
                                    Page(TITRES[3], s04))),
    Chapitre("Nos convictions", (Page(TITRE_CONVICTIONS, annonce_convictions),
                                 *(Page(c["titre"], partial(focus, conv=c))
                                   for c in CONVICTIONS))),
    Chapitre("Notre approche", (Page(TITRES[4], s05), Page(TITRES[5], s06),
                                Page(TITRES[6], s07))),
    Chapitre("Démarche", (Page(TITRES[7], s08), Page(TITRES[8], s09), Page(TITRES[9], s10))),
    Chapitre("Preuve", (Page(TITRES[10], s11), Page(TITRES[11], s12))),
    Chapitre("Dispositif et prix", (Page(TITRES[12], s13), Page(TITRES[13], s14),
                                    Page(TITRES[14], s15))),
    Chapitre("Équipe", (Page(TITRE_EQUIPE, equipe),)),
    Chapitre("Annexes", (Page(TITRES_ANNEXES[0], a_problemes),
                         Page(TITRES_ANNEXES[1], a_alternatives)), lettre=True),
]


def numeros_intercalaires():
    """Numéro (ou lettre) imprimé sur chaque intercalaire : CALCULÉ depuis CHAPITRES."""
    out, n = {}, 0
    for c in CHAPITRES:
        if not c.intercalaire:
            continue
        if c.lettre:
            out[c.nom] = c.nom[0].upper()
        else:
            n += 1
            out[c.nom] = f"{n:02d}"
    return out


def titres_attendus():
    """Titres de toutes les slides dans l'ordre (intercalaire = nom du chapitre)."""
    out = []
    for c in CHAPITRES:
        if c.intercalaire:
            out.append(c.nom)
        out += [p.titre for p in c.pages]
    return out


def texte_bandeau(nom, rang, taille):
    return f"{nom.upper()} · {rang}/{taille}"


def bandeau(s, nom, rang, taille):
    """Bandeau courant : NOM DU CHAPITRE · x/y (9 pt muted), sous le titre."""
    rich(s, X0, BANDEAU_Y, 6.0, 0.2, [(texte_bandeau(nom, rang, taille),
                                       {"size": 9, "color": MUTED})])


# ---------------------------------------------------------------- assemblage
def _layout(prs, debut):
    return next(lay for lay in prs.slide_layouts if lay.name.startswith(debut))


def _remplacer_texte(sh, texte):
    """Remplace le texte en gardant la mise en forme du premier run."""
    tf = sh.text_frame
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    runs = tf.paragraphs[0].runs
    if runs:
        runs[0].text = texte
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        tf.text = texte


def intercalaire(prs, modele, nom, numero, scene):
    """Clone une slide « Chapitre » du gabarit : titre = nom, numéro calculé, photo CC0."""
    sl = prs.slides.add_slide(modele.slide_layout)
    for sh in list(sl.shapes):
        retirer(sh)
    arbre = sl.shapes._spTree
    for el in list(modele.shapes._spTree)[2:]:       # sauf nvGrpSpPr / grpSpPr
        arbre.append(copy.deepcopy(el))
    for sh in list(sl.shapes):
        if not sh.has_text_frame:
            continue
        txt = sh.text_frame.text.strip()
        if txt in ("01", "02"):
            _remplacer_texte(sh, numero)
        elif txt == "Chapitre.":
            retirer(sh)                              # reliquat du gabarit
        elif sh.is_placeholder:
            _remplacer_texte(sh, nom)
    photo_chapitre(sl, scene)
    return sl


def sommaire(prs):
    """Sommaire sur le gabarit natif « Table des matières [7] », alimenté par CHAPITRES.

    Le gabarit a 7 emplacements numérotés 01 à 07 ; un 8e (Annexes, lettre) est ajouté
    dans la 4e rangée de la colonne de droite, vide dans le gabarit."""
    layout = _layout(prs, "95 - ")
    s = prs.slides.add_slide(layout)
    s.shapes.title.text = TITRE_SOMMAIRE
    numeros = numeros_intercalaires()
    listes = [c for c in CHAPITRES if c.intercalaire]
    assert len(listes) <= 8, "sommaire : 8 entrées au plus (7 du gabarit + 1 ajoutée)"
    for k, c in enumerate(listes, 1):
        if k <= 7:
            assert numeros[c.nom] == f"{k:02d}", "le gabarit imprime 01-07 : ordre incohérent"
            ph(s, k).text_frame.text = c.nom
        else:
            _entree_supplementaire(s, layout, c.nom, numeros[c.nom])
    return s


def _entree_supplementaire(s, layout, nom, repere):
    """8e entrée du sommaire : pastille + repère du gabarit (rangée 4) décalés à droite."""
    emu = 914400
    cercle = next(sh for sh in layout.shapes if sh.shape_type == 1
                  and abs(sh.top / emu - 3.91) < 0.02 and sh.left / emu < 2)
    num = next(sh for sh in layout.shapes if sh.has_text_frame
               and sh.text_frame.text.strip() == "04")
    for modele, x in ((cercle, 5.24), (num, 5.26)):
        el = copy.deepcopy(modele._element)
        el.spPr.find(qn("a:xfrm")).find(qn("a:off")).set("x", str(int(x * emu)))
        if modele is num:
            next(el.iter(qn("a:t"))).text = repere
        s.shapes._spTree.append(el)
    rich(s, 6.4, 3.91, 2.7, 0.65, [(nom, {"size": 12})], anchor=MSO_ANCHOR.MIDDLE)


def construire():
    prs = Presentation(SOURCE)
    sld = prs.slides._sldIdLst
    ids = list(sld)
    src = list(prs.slides)
    couv = src[0]
    modeles = [src[n - 1] for n in MODELES_INTERCALAIRE]
    layout = _layout(prs, "04 - Titre seul")
    numeros = numeros_intercalaires()
    ordre, k = [], 0
    for chap in CHAPITRES:
        if chap.intercalaire:
            intercalaire(prs, modeles[k % len(modeles)], chap.nom, numeros[chap.nom],
                         PHOTOS_INTERCALAIRES[k])
            ordre.append(sld[-1])
            k += 1
        for rang, page in enumerate(chap.pages, 1):
            if page.genre == "couverture":
                page.fabrique(couv)
                ordre.append(ids[0])
            elif page.genre == "sommaire":
                sommaire(prs)
                ordre.append(sld[-1])
            else:
                sl = prs.slides.add_slide(layout)
                page.fabrique(sl)
                bandeau(sl, chap.nom, rang, len(chap.pages))
                ordre.append(sld[-1])
    assert k <= len(PHOTOS_INTERCALAIRES)
    for el in list(sld):
        if not any(el is o for o in ordre):
            prs.part.drop_rel(el.rId)
        sld.remove(el)
    for el in ordre:
        sld.append(el)
    return prs


def textes(prs):
    out = []
    for i, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if sh.has_text_frame:
                out.append((i, sh.text_frame.text))
    return out


def main():
    prs = construire()
    defauts = [f"slide {i}: « gaspillage » subsiste" for i, tx in textes(prs)
               if "gaspill" in tx.lower()]
    defauts += D.verifier_geometrie(prs)
    debord = D.verifier_debordements_texte(prs, tolerance_in=0.4)
    os.makedirs(os.path.dirname(SORTIE), exist_ok=True)
    prs.save(SORTIE)
    print(f"{len(prs.slides)} slides, {len(CHAPITRES)} chapitres -> {SORTIE}")
    for d in debord:
        print("  debordement estime :", d)
    if defauts:
        print("DEFAUTS :\n  " + "\n  ".join(defauts))
        sys.exit(1)
    print("CONTROLE: OK")


if __name__ == "__main__":
    main()
