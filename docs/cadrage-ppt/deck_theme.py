"""Thème du deck (extrait mécaniquement de generate_deck.py) : gabarit,
constantes de mise en page, couleurs lues dans le thème du template,
rayon de coin branché sur D.add_rect, new_prs.
"""
import os
import sys

import pptx_deck_vscode3 as D
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(__file__)
TEMPLATE = os.path.join(HERE, "template-octo.pptx")

LAYOUT_COUVERTURE = 8   # "40 - Couverture [1]" — idx0 titre, idx1 sous-titre, idx2/idx3 crédit+date
LAYOUT_TITRE_SEUL = 5   # "04 - Titre seul" — idx0 titre, garde logo/pied de page/n° de slide
LAYOUT_VIDE = 0         # "06 - Slide vide" — pas de placeholder, juste logo + badge de pagination
LAYOUT_CHAPITRE = 2     # "50 - Chapitre [1]" — idx0 titre (grand), idx1 numéro ; cadre photo teardrop
LAYOUT_VISUEL_DROITE = 15  # "63 - Titre, contenu et visuel à droite - cadre blanc"

# --- Géométrie du template OCTO réel (10 x 5.625 in, 16:9) — cf.
# docs/vscode1-export/template-octo.md §4-5, vérifiée localement contre
# template-octo.pptx (mêmes dims/layouts/thème). Contenu dessiné dans la
# zone de contenu du layout « Titre seul » (sous le titre, au-dessus du
# pied de page), marge gauche alignée sur le placeholder titre (0.615 in),
# marge droite plafonnée avant le badge de pagination bas-droit.
SLIDE_W, SLIDE_H = 10.0, 5.625
MARGIN = 0.615
BORD_DROIT = 9.15
CONTENT_TOP = 1.15
CONTENT_BOTTOM = 5.45
CONTENT_W = BORD_DROIT - MARGIN
CONTENT_H = CONTENT_BOTTOM - CONTENT_TOP
GAP = 0.2

def _exiger_template():
    """Garde à l'import (finding robustesse, audit 2026-07-23) : sans elle, un
    template absent remontait en FileNotFoundError brute depuis python-pptx. Le
    générateur est lancé à la main — l'échec doit nommer le fichier attendu et
    son emplacement, sans exiger de lire la stack."""
    if not os.path.isfile(TEMPLATE):
        raise SystemExit(
            "generate_deck : template introuvable — placer template-octo.pptx "
            f"à côté du générateur (attendu : {TEMPLATE})"
        )


_exiger_template()
TH = D.theme_colors(Presentation(TEMPLATE))
NAVY = TH.get("dk1", D.INK)          # #0E2356 — texte principal, titres
DK2 = TH.get("dk2", NAVY)            # #3E4F78 — navy secondaire (palier de gradient sans PALETTE)
WHITE = TH.get("lt1", "#FFFFFF")
ACCENT = TH.get("accent3", NAVY)    # #00D2DD — cyan OCTO, identité du deck
MUTED = TH.get("lt2", D.MUTED)       # #586586 — slate 600, texte secondaire
ACCENT1 = TH.get("accent1", MUTED)   # #6E7B9A — bleu-gris clair (palier de gradient sans PALETTE)
ACCENT2 = TH.get("accent2", ACCENT1)  # #9FA7BB — bleu-gris très clair, le plus clair du thème
LINE = TH.get("accent5", D.LINE)     # #CFD3DD — slate 200, bordures de cards
TRACK = TH.get("accent6", D.TRACK)   # #E7E9EE — slate 100, fonds d'encarts

# D0..D4 : rampe MONOCHROME de la famille navy, du clair au foncé. Une
# échelle ordonnée se rend par une rampe, pas par des teintes étrangères —
# le vert->rouge d'avant lisait comme un feu tricolore sur un thème qui n'a
# ni vert ni rouge. Les 5 tons portent tous du texte blanc (chip() écrit en
# blanc par défaut) : le plus clair, #586586, tient 5,80:1 — un chiffre estimé
# à 5,0 de tête ici le 2026-09-10, puis mesuré. Le rejouer plutôt que le citer :
#   ratio = (L1+0,05)/(L2+0,05), L = luminance relative WCAG 2.x.
# ATTENTION, mesuré aussi : la rampe ne discrimine plus ses paliers adjacents
# (1,18 / 1,18 / 1,40 / 1,33). D0 et D2 — « CONFIRMÉ » et « DÉDUIT » — ne se
# distinguent QUE par leur texte. C'est assumé tant qu'un libellé les
# accompagne ; une échelle lue à l'œil seul demanderait plus d'amplitude.
# --- v2.45 (demande utilisateur : « des formes plus jolies avec plus
# d'arrondi ») : UN rayon de coin absolu pour tout le deck, au lieu de ~12
# ajustements relatifs épars (0.06 à 0.35 × le petit côté — un encart de 0,5in
# à 0.12 n'avait que 0,06in de coin, une carte de 2in à 0.08 en avait 0,16).
# Toute forme arrondie non pilule reçoit AU MOINS ce rayon, exprimé en pouces ;
# les pilules (0.5) restent des pilules. Branché sur D.add_rect, donc aussi sur
# D.add_card et tous les helpers du générateur qui dessinent par lui.
RAYON_COIN_IN = 0.20
_add_rect_brut = D.add_rect


def _add_rect_arrondi(slide, l, t, w, h, fill=None, line=None, line_w=1.0, rounded=False,
                      radius=0.12):
    if rounded and radius < 0.5 and min(w, h) > 0:
        radius = max(radius, min(0.5, RAYON_COIN_IN / min(w, h)))
    return _add_rect_brut(slide, l, t, w, h, fill=fill, line=line, line_w=line_w,
                          rounded=rounded, radius=radius)


D.add_rect = _add_rect_arrondi


SEVERITE = ["#586586", "#4A5A80", "#3E4F78", "#26386A", "#0E2356"]


def _rgb(hexcolor):
    return RGBColor.from_string(hexcolor.lstrip("#").upper())


def new_prs():
    _exiger_template()
    prs = Presentation(TEMPLATE)
    # Retire les 9 slides d'exemple du template — masters/layouts/thème conservés.
    # Il faut aussi supprimer la relation (drop_rel), sinon les parties
    # ppt/slides/slideN.xml orphelines entrent en collision de nom avec les
    # nouvelles slides ajoutées ensuite (même numérotation réutilisée).
    xml_slides = prs.slides._sldIdLst
    for sld in list(xml_slides):
        rId = sld.get(D.qn("r:id"))
        prs.part.drop_rel(rId)
        xml_slides.remove(sld)
    return prs


# Il n'y a PLUS de couleur par chapitre. Un mecanisme `PALETTE_CHAPITRES` /
# `couleur_chapitre()` a existe quelques heures le 2026-09-10, entre la mesure
# de l'ecart a la charte et l'arbitrage qui a suivi : il cyclait 4 tons du
# theme sur les chapitres. L'arbitrage — la couleur ne porte pas le sens — l'a
# rendu caduc le jour meme. Retire plutot que laisse en place : ses 10 appels
# reels passaient deja `ENCRE` en dur, et ses 2 seuls appelants residuels
# donnaient un kicker GRIS a deux slides dont l'intercalaire est navy.
#
# --- Vocabulaire de différenciation SANS code couleur (arbitrage du 2026-09-10)
#
# La couleur ne porte plus le sens. Les 160 sites qui appelaient `D.PALETTE[n]`
# — un bleu pour l'infra, un teal pour l'utilisateur, un or pour le management,
# un violet pour le sponsor — pointent tous sur `ENCRE`. Ce qui différencie
# désormais deux éléments de même niveau, dans l'ordre où le catalogue des decks
# OCTO réels les emploie (`deck-design-library`) :
#
#   1. « UN SUR N EN ACCENT » — dans une série d'éléments égaux, un SEUL reçoit
#      un aplat plein (cyan ou navy), les autres restent blancs à contour. Le
#      catalogue le donne comme le mécanisme de hiérarchie le plus systématique
#      du deck, avant même la taille de police.
#   2. La NUMÉROTATION (badge « goutte » + connecteur) et la POSITION (quinconce
#      plutôt qu'alignement en tableau).
#   3. La FORME-SIGNATURE : coins arrondis + un coin coupé pour le contenu
#      riche, pilule pour les chips et étiquettes.
#   4. La TYPOGRAPHIE : accroche grasse + complément régulier, sous-en-têtes en
#      majuscules 8-9pt comme rupture sans bordure.
#
# RÈGLE DURE, mesurée au rendu du 2026-09-10 : le cyan ne porte JAMAIS de texte
# sur blanc — il plafonne à ~1,9:1 et le libellé se délave (constaté sur
# « Infra & RUN » de la slide personas). Cyan = aplat, badge, chip, connecteur.
ENCRE = NAVY                 # tout ce qui porte du sens : texte, contours, filets
ACCENT_PLEIN = ACCENT        # cyan — l'élément mis en avant d'une série, EN APLAT
SUPPORT = TRACK              # fond neutre d'encart, jamais porteur de sens
SUPPORT_LIGNE = LINE         # bordures discrètes


def encre_de(color):
    """Couleur de TEXTE sûre pour un élément dont l'accent est `color`.

    Beaucoup de renderers font piloter la bordure, la barre d'accent ET le
    libellé par une seule variable `color`. C'est commode tant que la couleur
    est sombre — et faux dès qu'elle vaut le cyan : le texte se délave à
    ~1,9:1 sur blanc (constaté au rendu du 2026-09-10 sur l'exec summary et sur
    « Infra & RUN »). Cette garde laisse passer les tons sombres et rabat le
    seul cyan sur l'encre, pour que l'accent reste VISIBLE (bordure, barre,
    aplat) sans que le libellé devienne illisible.
    """
    return ENCRE if str(color).lower() == str(ACCENT).lower() else color


# v2.40 : le kicker d'une slide de contenu EST le nom du chapitre qui la porte.
# Il était écrit en dur à chaque appel (« IA », « PROPOSITION », « BESOINS &
# DOULEURS »...) et survivait aux fusions de chapitres — cinq kickers citaient
# des chapitres disparus au rendu v2.39. build() pose ici le chapitre courant
# (intercalaire, ou « Executive summary » / « Annexe » sans intercalaire) ;
# les appelants passent `kicker=None`. Contrôlé par test_generate_deck.py.
