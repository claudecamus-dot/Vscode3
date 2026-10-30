"""Delta V28 (base V27) : retours utilisateur du 2026-10-30 — « chapitre Prochain pas : on supprime » et « les formes sont toujours similaires et ne correspondent pas aux formes que
j'ai demande d'investiguer dans le PDF » (OCTO_GROUND).

- chapitre « Prochain pas » supprime (intercalaire + slide), pastille de navigation et ligne du sommaire retirees, pastille « Annexes » remontee ;
- formes reprises du vocabulaire d'OCTO_GROUND (forme_ground.py) : fond blanc et trait navy fin, goutte a un coin pointu, colonnes a onglet reliees par un point cyan, barres-pilules,
  anneaux enchaines, frise a cartouches, pastilles a coche, grands pictogrammes au trait ; appliquees a l'introduction des convictions, aux douze convictions, aux deux temps, au planning,
  aux livrables et au chiffrage ; les autres cartes grises passent en cartes blanches a trait navy.
Usage (depuis docs/cadrage-ppt/) : py deck_delta/deck.py delta_V28 --base V27
"""
import re

import deck_base
import deck_draw
from delta_V20 import bandeau
from delta_V21 import sections
from forme_ground import *  # noqa: F401,F403
from forme_v26 import minicarte
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE

BASE = "V27"
NOTE = "V28 : chapitre Prochain pas supprime ; formes OCTO_GROUND (goutte, onglets, anneaux, pilules, cartouches, coches, pictos)"

T_INTRO = "Douze convictions en quatre thèmes, déclinées une à une ensuite"
T_CULTURE = "Conviction 1 : le mode produit est d'abord une culture, pas un outil"
T_TROIS = "Trois convictions, chacune vérifiée par un critère et un chantier"
T_DEUX = "Deux convictions de plus, vérifiées de la même façon"
T_ORG = "Conviction 7 : financer des produits et des équipes stables"
T_EQUIPE = "Conviction 8 : une petite équipe produit qui construit et exploite"
T_PILOT = "Conviction 9 : trois instances, chacune avec sa décision"
T_DEMANDE = "Conviction 10 : un flux d'entrée unique qui qualifie avant de servir"
T_SERVICE = "Conviction 11 : piloter le service vécu, pas seulement le SLA"
T_RUN = "Conviction 12 : aucun produit livré sans RUN prêt"
T_DEUX_TEMPS = "Deux temps d'engagement : diagnostiquer, puis expérimenter et déployer"
T_PLANNING = "Un planning hypothétique en trois phases, avant tout chiffrage"
T_LIV1 = "Exemples de livrables (1/2) : comprendre et décider"
T_LIV2 = "Exemples de livrables (2/2) : piloter et ancrer"
T_CHIFFRAGE = "Chiffrage : la forme du tableau jours × rôles, toutes valeurs à valider"
T_PROCHAIN = "Prochain pas : choisir, diagnostiquer, tester, puis nous contacter"

_prs0 = Presentation(deck_base.resoudre_base(BASE))


def _est_inter(s):
    return s.slide_layout.name.startswith(("50 -", "51 -"))


_SUPPR = [s.slide_id for s in _prs0.slides if s.shapes.title is not None and (
    s.shapes.title.text_frame.text == T_PROCHAIN or (_est_inter(s) and s.shapes.title.text_frame.text == "Prochain pas"))]
OPS = [{"slide": f"id:{i}", "op": "supprimer"} for i in _SUPPR]
FX, FW = 190, 469


def T(sh):
    return sh.text_frame.text


def textes(s):
    return [T(sh) for sh in s.shapes if not garder(sh) and sh.has_text_frame and sh.text_frame.text.strip()]


def nettoyer(s):
    for sh in list(s.shapes):
        if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and sh.width == 8 * 12700 and sh.height == 8 * 12700:
            sh._element.getparent().remove(sh._element)
    vider(s)


# ---------------------------------------------------------------- suppression du chapitre Prochain pas
def retirer_prochain_pas(prs):
    for s in prs.slides:
        for sh in list(s.shapes):
            if sh.name == "nav-Prochain pas":
                sh._element.getparent().remove(sh._element)
            elif sh.name == "nav-Annexes":
                sh.left = 562 * 12700
    s = deck_draw.slide_par_titre(prs, "Sommaire")
    for sh in list(s.shapes):
        if sh.has_text_frame and T(sh) == "Prochain pas":
            sh._element.getparent().remove(sh._element)
        elif sh.has_text_frame and T(sh) == "A" and sh.top > 280 * 12700:
            sh.top = 230 * 12700
        elif sh.has_text_frame and T(sh) == "Annexes":
            sh.top = 225 * 12700
        elif sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and abs(sh.left - 377 * 12700) < 12700 and abs(sh.top - 282 * 12700) < 12700:
            sh.top = 224 * 12700
    cache = forme(s, MSO_SHAPE.RECTANGLE, 373, 218, 57, 57, BLANC, None)  # masque le « 07 » du gabarit, derriere le disque « A »
    s.shapes._spTree.remove(cache._element)
    s.shapes._spTree.insert(2, cache._element)


# ---------------------------------------------------------------- introduction des convictions : quatre colonnes a onglet
THEMES = [
    ("Culture et organisation", [(1, "La culture produit avant les outils", None), (7, "Financer des produits et des équipes stables", None),
                                 (8, "Une petite équipe qui construit et exploite", None)]),
    ("Le produit et l'usage", [(2, "Partir des douleurs des utilisateurs", "Boucle : Explorer"), (3, "Un porteur produit pour chaque capacité", "Boucle : Ancrer")]),
    ("Pilotage et demande", [(4, "Décider sur des faits mesurés", "Boucle : Prioriser"), (9, "Trois instances, trois décisions", None),
                            (10, "Un flux d'entrée qui qualifie avant de servir", None)]),
    ("RUN, progrès et IA", [(5, "Un progrès mesuré, régulier, transparent", "Boucle : Mesurer"), (6, "L'IA gagne du temps là où il est perdu", "Boucle : Comprendre"),
                           (11, "Piloter le service vécu, pas le seul SLA", None), (12, "Aucun produit livré sans RUN prêt", None)]),
]


def introduction(prs):
    s = deck_draw.slide_par_titre(prs, T_INTRO)
    nettoyer(s)
    signature_cyan(s)
    W, G = 143, 14
    for k, (theme, lignes) in enumerate(THEMES):
        x = X0 + k * (W + G)
        onglet(s, x, 98, W, 264, theme, 11.5)
        if k < 3:
            point_cyan(s, x + W + G / 2, 98 + 24)
        for r, (n, these, boucle) in enumerate(lignes):
            y = 156 + r * 51
            forme(s, MSO_SHAPE.OVAL, x + 10, y + 12, 22, 22, NAVY, None)
            tx(s, x + 10, y + 12, 22, 22, [[(str(n), 9.5, True, BLANC)]], A.MIDDLE, P.CENTER)
            paras = [[(these, 10, True, NAVY)]] + ([[(boucle, 9.5, False, MUTED)]] if boucle else [])
            tx(s, x + 38, y, W - 46, 46, paras, A.MIDDLE)


# ---------------------------------------------------------------- pied commun (critere / chantier) et cadre des convictions a disque
def bas(s, critere, chantier):
    for k, (lab, texte_, style) in enumerate((("LE CRITÈRE OBSERVABLE", critere, "navy"), ("LE CHANTIER", chantier, "cyan"))):
        x = X0 + k * 314
        carte(s, x, 284, 300, 76, 22)
        pilule(s, x + 16, 273, 170 if k == 0 else 112, 22, lab, style, 10)
        coche(s, x + 24, 328, 18)
        tx(s, x + 42, 300, 246, 56, [[(texte_, 10.5, False, NAVY)]], A.MIDDLE)


def cadre(s, n, genre, form, critere, chantier):
    nettoyer(s)
    signature_cyan(s)
    disque_picto(s, 110, 158, genre, 112)
    pilule(s, 45, 226, 130, 22, f"CONVICTION {n}", "navy", 10)
    form(s)
    bas(s, critere, chantier)
    minicarte(s, {n})


def trois_colonnes(s, cols, h=164):
    W, G = 143, 20
    for k, (titre, lignes) in enumerate(cols):
        x = FX + k * (W + G)
        onglet(s, x, 98, W, h, titre, 11.5)
        if k < 2:
            point_cyan(s, x + W + G / 2, 98 + 24)
        tx(s, x + 12, 152, W - 24, h - 62, [[(t, 10.5, False, NAVY)] for t in lignes], espace=5)


def f_culture(s):
    trois_colonnes(s, [("Une culture, pas un outil", ["Le mode produit change les décisions, les rôles et les habitudes.", "Les outils viennent ensuite."]),
                       ("Une règle posée d'en haut", ["Chaque investissement est discuté en termes de valeur pour l'usager.", "Le sponsor porte ce message."]),
                       ("Des managers formés", ["Les managers apprennent à piloter par la valeur.", "La valeur est mesurée au T0, puis au T1."])])


def f_org(s):
    trois_colonnes(s, [("Un budget par projet", ["L'équipe est dissoute à la fin du projet.", "Le RUN est transféré à d'autres équipes."]),
                       ("Une capacité par produit", ["Une équipe stable, qui construit et exploite.", "Un budget annuel rattaché au produit."]),
                       ("Produit ou composant ?", ["Produits : Kubernetes, CI/CD, IAM.", "Composants : VM, serveur, SAN.", "Puis-je décrire l'offre, la valeur, les engagements ?"])])


def f_equipe(s):
    for k, (genre, titre, texte_) in enumerate((("cycle", "Construire et exploiter", "Une même équipe conçoit, livre et assure le service. Du temps de construction est protégé."),
                                                 ("personnes", "Un binôme porte le produit", "Un responsable produit et un référent technique. Pas un PO posé dans l'Infra."),
                                                 ("fleche", "Moins de tickets", "Les demandes récurrentes sont automatisées et rendues en libre-service."))):
        x = FX + k * 163
        goutte_contour(s, x + 27, 98, 88)
        picto(s, genre, x + 71, 142, 50)
        tx(s, x, 192, 143, 16, [[(titre, 11.5, True, NAVY)]], A.TOP, P.CENTER)
        tx(s, x + 4, 210, 135, 50, [[(texte_, 9.5, False, NAVY)]], A.TOP, P.CENTER)


def f_instances(s):
    lignes = [("Comité d'investissement", "Décide : arrêter, étendre ou pivoter une capacité", "Avec : sponsor, DSI, responsables produit · cadence trimestrielle, à valider"),
              ("Revue produit", "Décide : backlog, roadmap et budget du produit", "Avec : responsable produit, métiers, équipe · cadence à fixer avec vous"),
              ("Suivi d'exécution", "Décide : ajustements, risques, escalades", "Avec : équipe, coach, relais client · cadence à fixer avec vous")]
    for r, (nom, decide, avec) in enumerate(lignes):
        y = 102 + r * 54
        pilule(s, FX, y, 150, 24, nom, "navy", 10.5)
        pilule(s, FX + 158, y, 311, 24, decide, "contour", 10)
        tx(s, FX + 160, y + 28, 309, 14, [[(avec, 9.5, False, MUTED)]])


def f_demande(s):
    etapes = [("Recevoir", "Un seul point d'entrée ; une fiche : problème, valeur, indicateur."), ("Qualifier", "Valeur, faisabilité et charge estimées avant d'engager l'équipe."),
              ("Prioriser", "Une grille explicite ; au-delà d'un seuil de charge, un ROI est demandé."), ("Servir", "Statut visible ; les demandes récurrentes passent en self-service.")]
    for k, (nom, texte_) in enumerate(etapes):
        x = FX + k * 119
        anneau(s, x + 16, 100, 80, cyan=(k % 2 == 0))
        tx(s, x + 16, 100, 80, 80, [[(str(k + 1), 20, True, NAVY)]], A.MIDDLE, P.CENTER)
        if k < 3:
            forme(s, MSO_SHAPE.RIGHT_ARROW, x + 98, 134, 18, 12, CYAN, NAVY, lw=0.75)
        pilule(s, x + 2, 190, 108, 20, nom, "contour", 10)
        tx(s, x + 2, 214, 112, 50, [[(texte_, 9.5, False, NAVY)]], A.TOP, P.CENTER)


def f_service(s):
    cw = [150, 160, 159]
    cx = [FX, FX + 150, FX + 310]
    forme(s, MSO_SHAPE.ROUNDED_RECTANGLE, FX, 100, FW, 26, CYAN, None, adj=0.5)
    for j, h_ in enumerate(("Ce que mesure le SLA", "Ce que vit le consommateur", "Ce que nous suivons, T0 → T1")):
        tx(s, cx[j] + 10, 100, cw[j] - 14, 26, [[(h_, 10, True, NAVY)]], A.MIDDLE)
    lignes = [("Disponibilité du service", "Une demande qui attend, même quand le service est « vert »", "Objectif de service publié par capacité"),
              ("Incidents clos dans le délai", "Le même incident qui revient", "Taux d'incidents récurrents"),
              ("Tickets traités", "Du temps passé à subir plutôt qu'à améliorer", "Part de RUN subi et maîtrisé")]
    for r, ligne in enumerate(lignes):
        y = 130 + r * 45
        forme(s, MSO_SHAPE.RECTANGLE, FX, y, FW, 42, CLAIR if r % 2 == 0 else BLANC, None)
        for j, c in enumerate(ligne):
            tx(s, cx[j] + 10, y, cw[j] - 14, 42, [[(c, 10, j == 2, NAVY)]], A.MIDDLE)


def f_run(s):
    for k, nom in enumerate(("Construire", "Porte", "Exploiter", "Apprendre")):
        cartouche(s, FX + k * 118, 100, f"0{k + 1}", nom, w2=72, h=24)
    panneau(s, FX, 138, 229, 124, 34)
    tx(s, FX + 14, 146, 200, 16, [[("LA PORTE « RUN PRÊT »", 10, True, NAVY)]])
    puces_coche(s, FX + 14, 168, 205, ["Deux personnes capables par module", "Supervision et astreinte définies", "Budget RUN dans le business case", "Procédure d'incident documentée"], sz=9.5, pas=22)
    panneau(s, FX + 240, 138, 229, 124, 34)
    tx(s, FX + 254, 146, 200, 16, [[("LA BOUCLE D'APPRENTISSAGE", 10, True, NAVY)]])
    tx(s, FX + 254, 168, 205, 90, [[("Chaque incident majeur : un post-mortem, puis une action au backlog.", 10, False, NAVY)],
                                  [("L'IA aide à diagnostiquer ; la décision et la mise en œuvre restent humaines et tracées.", 10, False, NAVY)]], espace=6)


CONV = {  # titre -> (numero, picto, forme, critere, chantier)
    T_CULTURE: (1, "ampoule", f_culture, "Chaque équipe formule la valeur rendue à ses consommateurs et la revoit à chaque cycle ; les décisions citent un indicateur.",
                "Atelier de cadrage de la valeur par capacité, message du sponsor, coaching des managers et des équipes sur le terrain."),
    T_ORG: (7, "classeurs", f_org, "Part du budget infra rattachée à des équipes ou produits nommés, sur plusieurs trimestres.",
            "Définir le produit infra, nommer son responsable, passer d'un budget de projet à une capacité annuelle."),
    T_EQUIPE: (8, "personnes", f_equipe, "Temps de construction protégé ; demandes récurrentes passées en self-service.",
               "Cadrer la composition de l'équipe, le binôme produit et la répartition entre construire et exploiter."),
    T_PILOT: (9, "calendrier", f_instances, "Chaque instance a une carte d'identité écrite ; chaque revue trimestrielle produit une décision tracée.",
              "Atelier des cartes d'identité, seuils de délégation, calendrier annuel, premier cycle sur 2 ou 3 capacités."),
    T_DEMANDE: (10, "fleche", f_demande, "Toute demande a une fiche et un statut visibles ; le délai de qualification est mesuré.",
                "Cartographier les flux entrants, fiche de demande standard, grille de priorisation, capacité répartie par décision."),
    T_SERVICE: (11, "horloge", f_service, "Écart entre SLA et expérience mesuré au T0 puis au T1 ; part de RUN subi suivie.",
                "Objectifs de service par capacité, revue trimestrielle des écarts, charge de RUN rendue visible."),
    T_RUN: (12, "coche", f_run, "Contrôle « RUN prêt » validé avant chaque mise en service ; deux personnes capables par module.",
            "Contrôle de préparation au RUN sur un pilote, post-mortems reliés au backlog, IA d'aide au diagnostic."),
}


def convictions_a_disque(prs):
    for titre, (n, genre, form, critere, chantier) in CONV.items():
        cadre(deck_draw.slide_par_titre(prs, titre), n, genre, form, critere, chantier)


# ---------------------------------------------------------------- convictions 2 a 6 : cartes a pictogramme en tete
def convictions_multiples(prs):
    pictos = {T_TROIS: ["personnes", "classeurs", "cible"], T_DEUX: ["jauge", "cycle"]}
    nums = {T_TROIS: [2, 3, 4], T_DEUX: [5, 6]}
    for titre in (T_TROIS, T_DEUX):
        s = deck_draw.slide_par_titre(prs, titre)
        ts = textes(s)
        n_cartes = len(pictos[titre])
        data = [ts[6 * k:6 * k + 6] for k in range(n_cartes)]
        regle = ts[6 * n_cartes:6 * n_cartes + 2]
        nettoyer(s)
        signature_cyan(s)
        for k, (head, pill, b0, b1, b2, b3) in enumerate(data):
            x = X0 + k * 212
            carte(s, x, 118, 190, 244, 30)
            disque_picto(s, x + 95, 112, pictos[titre][k], 56, goutte=False)
            tx(s, x + 14, 146, 162, 62, [[(head, 11, True, NAVY)]], A.MIDDLE, P.CENTER)
            pilule(s, x + 40, 212, 110, 18, pill, "navy", 9.5)
            tx(s, x + 14, 236, 162, 12, [[(b0, 9.5, True, MUTED)]])
            tx(s, x + 14, 249, 162, 42, [[(b1, 10, False, NAVY)]])
            tx(s, x + 14, 296, 162, 12, [[(b2, 9.5, True, MUTED)]])
            tx(s, x + 14, 309, 162, 48, [[(b3, 10, False, NAVY)]])
        if regle:
            x = X0 + 424
            panneau(s, x, 118, 190, 244, 44)
            pilule(s, x + 20, 132, 130, 22, regle[0], "cyan", 10)
            tx(s, x + 18, 170, 156, 180, [[(regle[1], 11, False, NAVY)]])
        minicarte(s, set(nums[titre]))


# ---------------------------------------------------------------- deux temps : anneaux enchaines
def deux_temps(prs):
    s = deck_draw.slide_par_titre(prs, T_DEUX_TEMPS)
    ts = textes(s)
    bandes = ts[0:2]
    cols = [ts[2 + 6 * k:8 + 6 * k] for k in range(3)]
    note = ts[20]
    nettoyer(s)
    signature_cyan(s)
    pilule(s, X0, 96, 190, 22, bandes[0], "navy", 10)
    pilule(s, X0 + 212, 96, 402, 22, bandes[1], "cyan", 10)
    for k, (nu, ti, co, cd, du, mo) in enumerate(cols):
        x = X0 + k * 212
        anneau(s, x + 60, 124, 70, cyan=(k != 1))
        picto(s, ("cible", "cycle", "personnes")[k], x + 95, 159, 34)
        if k < 2:
            forme(s, MSO_SHAPE.RIGHT_ARROW, x + 160, 152, 24, 14, CYAN, NAVY, lw=0.75)
        tx(s, x, 200, 190, 18, [[(ti, 13, True, NAVY)]], A.TOP, P.CENTER)
        tx(s, x, 222, 190, 66, [[(co, 10.5, False, NAVY)]], A.TOP, P.CENTER)
        tx(s, x, 290, 190, 28, [[(cd, 10.5, True, NAVY)]], A.TOP, P.CENTER)
        pilule(s, x + 10, 318, 170, 18, du, "contour", 9.5)
        pilule(s, x + 10, 339, 170, 22, mo, "contour", 9)
    tx(s, X0, 366, 614, 14, [[(note, 10.5, False, MUTED)]])


# ---------------------------------------------------------------- planning : frise a cartouches
def planning(prs):
    s = deck_draw.slide_par_titre(prs, T_PLANNING)
    nettoyer(s)
    signature_cyan(s)
    pilule(s, 489, 94, 170, 20, "HYPOTHÈSE · SANS CHIFFRAGE", "contour", 10)
    TX0, TW = 165, 494
    for k in range(4):
        cartouche(s, TX0 + k * 124, 120, f"0{k + 1}", f"Trimestre {k + 1}", w2=78, h=24)
    trait(s, TX0, 158, TX0 + TW, 158, 1)
    for x, lab in ((TX0 + 100, "choix du pilote"), (TX0 + 270, "comité de décision"), (TX0 + TW - 8, "passage de main")):
        noeud(s, x, 158, 16)
        dernier = x >= TX0 + TW - 20
        tx(s, (X1 - 120) if dernier else x - 60, 168, 120, 14, [[(lab, 10, False, MUTED)]], A.TOP, P.RIGHT if dernier else P.CENTER)
    for r, (nom, x, w) in enumerate((("Assessment flash", TX0, 100), ("Pilote Produit + IA", TX0 + 108, 160), ("Ancrage", TX0 + 276, TW - 276))):
        y = 192 + r * 27
        tx(s, X0, y, 116, 22, [[(nom, 11, True, NAVY)]], A.MIDDLE)
        pilule(s, x, y, w, 22, "durée à valider", "navy", 10)
    for k, (lab, lignes, style) in enumerate((("LIVRABLES CLÉS", ("Assessment flash : baseline, pilote choisi", "Pilote : produit Infra livré, indicateurs T0 → T1", "Ancrage : relais formés, passage de main"), "navy"),
                                                ("HYPOTHÈSES DE CALENDRIER", ("Démarrage après le choix du périmètre", "Durées à caler avec vous", "Le chiffrage est traité à part"), "cyan"))):
        x = X0 + k * 314
        carte(s, x, 284, 300, 76, 22)
        pilule(s, x + 16, 273, 190 if k else 130, 22, lab, style, 10)
        puces_coche(s, x + 16, 298, 276, lignes, sz=10, pas=18)


# ---------------------------------------------------------------- livrables : colonnes a onglet
LIVRABLES = {
    T_LIV1: [("Cartographie des capacités", "grille", "Voir les capacités, leurs consommateurs et leurs propriétaires.", "tableau", "ensemble", "début"),
             ("Baseline des indicateurs", "jauge", "Fixer l'état de départ (T0) pour mesurer ensuite l'apport.", "tableau", "OCTO", "début"),
             ("Fiche produit Infra", "document", "Décrire l'offre, la valeur, les engagements et le propriétaire.", "fiche", "ensemble", "pilote"),
             ("Grille de priorisation", "cible", "Arbitrer les demandes sur la valeur, la faisabilité et la charge.", "grille", "ensemble", "pilote")],
    T_LIV2: [("Roadmap produit", "calendrier", "Planifier les évolutions et l'arrêt des capacités.", "frise", "client", "pilote"),
             ("Tableau de bord", "classeurs", "Suivre l'usage, la qualité et le coût, du T0 au T1.", "écran", "client", "continu"),
             ("Charte IA et RACI", "personnes", "Cadrer qui prépare, qui valide et qui décide quand l'IA aide.", "charte", "ensemble", "pilote"),
             ("Plan de passage de main", "fleche", "Transférer l'animation aux relais formés du client.", "plan", "OCTO", "fin")],
}


def livrables(prs):
    for titre, liste in LIVRABLES.items():
        s = deck_draw.slide_par_titre(prs, titre)
        nettoyer(s)
        signature_cyan(s)
        for k, (nom, genre, sert, f, qui, quand) in enumerate(liste):
            x = X0 + k * 157
            onglet(s, x, 130, 143, 224, nom, 11)
            if k < 3:
                point_cyan(s, x + 143 + 7, 130 + 24)
            picto(s, genre, x + 71, 112, 40)
            tx(s, x + 12, 184, 119, 60, [[(sert, 10, False, NAVY)]], A.TOP, P.CENTER)
            for j, (lab, val) in enumerate((("FORMAT", f), ("QUI", qui), ("QUAND", quand))):
                tx(s, x + 12, 246 + j * 36, 119, 11, [[(lab, 9, True, MUTED)]], A.TOP, P.CENTER)
                pilule(s, x + 22, 258 + j * 36, 99, 20, val, "contour", 9.5)
        tx(s, X0, 364, 614, 14, [[("Exemples illustratifs : formats et contenus à adapter en atelier avec vous.", 10, False, MUTED)]])


# ---------------------------------------------------------------- chiffrage : tableau a bandeaux
def chiffrage(prs):
    s = deck_draw.slide_par_titre(prs, T_CHIFFRAGE)
    ts = textes(s)
    hdr, rows, total, note, cond_t = ts[0:4], [ts[4 + 4 * r:8 + 4 * r] for r in range(3)], ts[16:20], ts[20], ts[21]
    cond = [ts[22 + 3 * k:25 + 3 * k] for k in range(4)]
    nettoyer(s)
    signature_cyan(s)
    CX, CW = [62, 296, 396, 512], [220, 90, 100, 140]
    forme(s, MSO_SHAPE.ROUNDED_RECTANGLE, X0, 98, 614, 30, CYAN, None, adj=0.5)
    for j in range(4):
        tx(s, CX[j], 98, CW[j], 30, [[(hdr[j], 10.5, True, NAVY)]], A.MIDDLE)
    for r in range(3):
        y = 132 + r * 34
        forme(s, MSO_SHAPE.RECTANGLE, X0, y, 614, 32, CLAIR if r % 2 == 0 else BLANC, None)
        for j in range(4):
            tx(s, CX[j], y, CW[j], 32, [[(rows[r][j], 12 if j == 0 else 14, True, NAVY)]], A.MIDDLE)
    forme(s, MSO_SHAPE.ROUNDED_RECTANGLE, X0, 238, 614, 32, NAVY, None, adj=0.5)
    for j in range(4):
        tx(s, CX[j], 238, CW[j], 32, [[(total[j], 12, True, BLANC)]], A.MIDDLE)
    tx(s, X0, 276, 614, 14, [[(note, 10.5, False, MUTED)]])
    carte(s, X0, 306, 614, 54, 22)
    pilule(s, X0 + 16, 295, 196, 22, cond_t, "navy", 10)
    for k, (lab, val, st) in enumerate(cond):
        x = X0 + 18 + k * 148
        tx(s, x, 322, 138, 12, [[(lab, 10, True, MUTED)]])
        tx(s, x, 334, 138, 20, [[(val, 14, True, NAVY)]])
        tx(s, x + 74, 336, 62, 14, [[(st, 9.5, False, MUTED)]], A.TOP, P.RIGHT)


# ---------------------------------------------------------------- autres slides : cartes grises -> cartes blanches a trait navy
A_RESTYLER = (MSO_SHAPE.ROUNDED_RECTANGLE, MSO_SHAPE.ROUND_2_DIAG_RECTANGLE, MSO_SHAPE.ROUND_1_RECTANGLE, MSO_SHAPE.ROUND_2_SAME_RECTANGLE)


def restyler(prs, deja):
    for s in prs.slides:
        if _est_inter(s) or s.slide_id in deja:
            continue
        for sh in s.shapes:
            if sh.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE or sh.width < 60 * 12700 or sh.height < 30 * 12700:
                continue
            try:
                if sh.auto_shape_type not in A_RESTYLER or sh.fill.type != 1 or sh.fill.fore_color.rgb != RGBColor(0xE7, 0xE9, 0xEE):
                    continue
            except (AttributeError, ValueError):
                continue
            sh.fill.solid()
            sh.fill.fore_color.rgb = BLANC
            sh.line.color.rgb = NAVY
            sh.line.width = Pt(0.75)


def retouche(prs):
    retirer_prochain_pas(prs)
    ids = []
    for t in (T_INTRO, T_TROIS, T_DEUX, T_DEUX_TEMPS, T_PLANNING, T_LIV1, T_LIV2, T_CHIFFRAGE, *CONV):
        ids.append(deck_draw.slide_par_titre(prs, t).slide_id)
    introduction(prs)
    convictions_a_disque(prs)
    convictions_multiples(prs)
    deux_temps(prs)
    planning(prs)
    livrables(prs)
    chiffrage(prs)
    restyler(prs, set(ids))
    bandeau(prs)
    sections(prs)


def _ids():
    return [f"id:{s.slide_id}" for s in _prs0.slides if s.slide_id not in _SUPPR]


ETAPES = [{"declarer": _ids(), "f": retouche, "nouvelles_max": 0, "nom": "chapitre Prochain pas supprime, formes OCTO_GROUND"}]
