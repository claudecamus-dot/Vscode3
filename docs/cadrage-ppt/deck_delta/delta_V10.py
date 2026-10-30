"""Delta V10 (base V9) : retours utilisateur du 2026-10-29 apres la V9.

1. police : les placeholders (titres, sommaire...) heritent de leur gabarit (Outfit SemiBold/Medium/Light, comme le deck Naval) au lieu d'un « Outfit » force par la V9 ;
   slide 1 intacte ; les formes hors placeholder gardent Outfit explicite (le theme dit Arial) ; navigation a 7 pt (Naval : 6-7 pt).
2. formes : en-tete flush avec son corps de carte (meme largeur, meme rayon, coins bas seuls arrondis au corps), liseres d'accent rentres dans la carte
   (modele : cartes a en-tete et pastille d'AGIMA 2026 et du deck Naval Group).
3. « Dispositif et prix » ne garde que le dispositif de mission et le chiffrage : « Deux temps d'engagement » et « Risques et objections » passent en « Demarche detaillee » ;
   compteurs « n/N » et pastille de navigation recalcules.
4. sections PowerPoint : une par chapitre (ouverture, 7 chapitres + demarche generale/detaillee, annexes).
Usage (depuis docs/cadrage-ppt/) : py deck_delta/deck.py delta_V10 --base V9
"""
import re
import uuid

import deck_draw
from delta_V9 import BLANC, CLAIR, EMU, MUTED, NAVY, POLICE, _chapitres, _est_intercalaire, _ids_base, _nom_chapitre  # noqa: F401
from lxml import etree
from pptx.util import Pt

BASE = "V9"
NOTE = "V10 : polices heritees du gabarit, en-tetes de cartes alignes, Dispositif et prix recentre, sections par chapitre"
OPS = []

T_DEUX_TEMPS = "Deux temps d'engagement : diagnostiquer, puis expérimenter et déployer"
T_RISQUES = "Risques et objections : ce que nous répondons, ce que nous mesurons"
RAYON = 12 * EMU


def police_heritee(prs):
    """Placeholders des slides 2..N : retire le nom de police force (le gabarit porte Outfit) ; nav a 7 pt."""
    for s in list(prs.slides)[1:]:
        for sh in s.shapes:
            if sh.is_placeholder and sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    for r in para.runs:
                        if r.font.name == POLICE:
                            r.font.name = None
            elif sh.name.startswith("nav-"):
                for r in sh.text_frame.paragraphs[0].runs:
                    r.font.size = Pt(7)


def _rond(sh):
    return sh.shape_type == 1 and sh.name.startswith("Rounded Rectangle")


def en_tetes_alignes(prs):
    """En-tete (Round Same Side...) pose sur le haut d'une carte : meme largeur, meme rayon ; le corps n'arrondit plus que ses coins bas."""
    for s in prs.slides:
        formes = list(s.shapes)
        for h in formes:
            if not (h.shape_type == 1 and h.name.startswith("Round Same Side") and h.width > 60 * EMU and h.height < 40 * EMU):
                continue
            corps = [b for b in formes if _rond(b) and b.height > 2 * h.height and abs(b.top - h.top - h.height) <= 2 * EMU
                     and b.left <= h.left + 2 * EMU and h.left + h.width <= b.left + b.width + 2 * EMU]
            if not corps:
                continue
            b = corps[0]
            h.left, h.width = b.left, b.width
            h.adjustments[0] = min(0.5, RAYON / h.height)
            geom = b._element.spPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom")
            geom.set("prst", "round2SameRect")
            for gd in list(geom.find("{http://schemas.openxmlformats.org/drawingml/2006/main}avLst")):
                geom.find("{http://schemas.openxmlformats.org/drawingml/2006/main}avLst").remove(gd)
            b.adjustments[0] = 0
            b.adjustments[1] = min(0.5, RAYON / min(b.width, b.height))


def liseres_rentres(prs):
    """Lisere d'accent colle au bord gauche d'une carte arrondie : rentre de 8 pt dans la carte (il ne mord plus sur le rayon)."""
    for s in prs.slides:
        formes = list(s.shapes)
        for ba in formes:
            if not (_rond(ba) and ba.width <= 7 * EMU and ba.height >= 20 * EMU):
                continue
            carte = next((c for c in formes if _rond(c) and c is not ba and c.width > 40 * EMU and c.left == ba.left and c.top == ba.top
                          and c.height == ba.height), None)
            if carte is None:
                continue
            ba.left = carte.left + 8 * EMU
            ba.top = carte.top + 8 * EMU
            ba.height = carte.height - 16 * EMU
            ba.width = 4 * EMU


def _kicker(s):
    return next((sh for sh in s.shapes if sh.has_text_frame and not sh.name.startswith("nav-")
                 and re.fullmatch(r"[^\d]+ · \d+/\d+", sh.text_frame.text)), None)


def _navigation(s, chapitre):
    for sh in s.shapes:
        if sh.name.startswith("nav-"):
            actif = sh.name == "nav-" + chapitre
            sh.fill.solid()
            sh.fill.fore_color.rgb = NAVY if actif else CLAIR
            r = sh.text_frame.paragraphs[0].runs[0]
            r.font.bold = actif
            r.font.color.rgb = BLANC if actif else MUTED


def deplacer_et_compter(prs):
    """Deplace les 2 slides de demarche, puis recompte les « n/N » et la pastille de navigation de chaque slide de chapitre."""
    sl = list(prs.slides)
    deux_temps = deck_draw.slide_par_titre(prs, T_DEUX_TEMPS)
    risques = deck_draw.slide_par_titre(prs, T_RISQUES)
    intercalaire_detaillee, derniere_detaillee = sl[20], sl[24]
    deck_draw.deplacer_apres(prs, deux_temps, intercalaire_detaillee)
    deck_draw.deplacer_apres(prs, risques, derniere_detaillee)
    par_chapitre = {}
    for s, ch in _chapitres(prs):
        if ch:
            par_chapitre.setdefault(ch, []).append(s)
    for ch, slides in par_chapitre.items():
        for i, s in enumerate(slides, 1):
            k = _kicker(s)
            nouveau = f"{ch.upper()} · {i}/{len(slides)}"
            if k is not None and k.text_frame.text != nouveau:
                deck_draw.set_run(k, nouveau)
            _navigation(s, ch)


def sections(prs):
    """Une section PowerPoint par chapitre : ouverture (jusqu'au sommaire), puis une section a chaque intercalaire."""
    groupes, courant = [], ("Ouverture", [])
    for s in prs.slides:
        if _est_intercalaire(s):
            groupes.append(courant)
            courant = (_nom_chapitre(s) or "Chapitre", [])
        courant[1].append(deck_draw.id_de(prs, s))
    groupes.append(courant)
    pres = prs.part._element
    ext_lst = pres.find("{http://schemas.openxmlformats.org/presentationml/2006/main}extLst")
    if ext_lst is None:
        ext_lst = etree.SubElement(pres, "{http://schemas.openxmlformats.org/presentationml/2006/main}extLst")
    for ext in ext_lst.findall("{http://schemas.openxmlformats.org/presentationml/2006/main}ext"):
        if ext.get("uri") == "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}":
            ext_lst.remove(ext)
    ext = etree.SubElement(ext_lst, "{http://schemas.openxmlformats.org/presentationml/2006/main}ext", uri="{521415D9-36F7-43E2-AB2F-B90AF26B5E84}")
    liste = etree.SubElement(ext, f"{{{deck_draw.P14}}}sectionLst", nsmap={"p14": deck_draw.P14})
    for nom, ids in groupes:
        sec = etree.SubElement(liste, f"{{{deck_draw.P14}}}section", name=nom, id="{" + str(uuid.uuid4()).upper() + "}")
        sl = etree.SubElement(sec, f"{{{deck_draw.P14}}}sldIdLst")
        for i in ids:
            etree.SubElement(sl, f"{{{deck_draw.P14}}}sldId", id=i)


def retouche(prs):
    """Une seule etape (la simulation se paie par etape) : ordre = deplacements/compteurs, formes, polices, sections."""
    deplacer_et_compter(prs)
    en_tetes_alignes(prs)
    liseres_rentres(prs)
    police_heritee(prs)
    sections(prs)


def _ids():
    from pptx import Presentation

    import deck_base
    return [f"id:{s.slide_id}" for s in Presentation(deck_base.resoudre_base(BASE)).slides]


ETAPES = [{"declarer": _ids(), "f": retouche, "nouvelles_max": 0, "nom": "polices, formes, Dispositif et prix, sections"}]
