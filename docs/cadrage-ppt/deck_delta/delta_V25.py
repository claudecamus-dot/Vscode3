"""Delta V25 (base V24) : applique les recommandations du comite a 10 bras (comite10/B1..B10.md, 2026-10-30) qui ne demandent aucune donnee d'OCTO.

Demande utilisateur : « applique les points, sauf : garder les slides de chapitre mais les rendre plus simples (retirer le gros numero de chapitre) ».
- intercalaires : le gros numero fantome et les huit reperes sont retires ; titre aligne sur le trait et la question ;
- ordre : « Six leviers... cinq resultats » rejoint Contexte et enjeux (apres la definition) ; « Deux temps d'engagement » ouvre Demarche generale ; le scenario etape par etape,
  les agents du coach et les trois principes passent en Annexes (contenu intact) ; pastilles de navigation, surtitres numerotes et sections recalcules ;
- texte (sans information nouvelle) : T0 et T1 definis a leur premiere occurrence (slide d'ouverture) ; « coach augmente » -> « coach produit augmente » ; « infra » -> « Infra » ;
  quatre titres-etiquettes (« : 1 a 3 ») deviennent des phrases ; le titre de « Prochain pas » dit « diagnostiquer » comme ses cartes ; les « 0 » du tableau de chiffrage deviennent « — » ;
- forme : pastilles de rubrique a 9 pt ; rayon d'angle des cartes ramene a 10 pt ; cartes des convictions etirees jusqu'a y = 340 ; cartes « dix problemes » reajustees ; corps des
  trois modes centre dans la carte ; bandes zebrees du tableau d'indicateurs alignees sur son en-tete ; cyan -> navy sauf liserés fins et intercalaires ; texte sous 9 pt remonte a 9 pt.
Non applique (decision utilisateur ou donnee manquante) : fusion des convictions 3 et 4, chiffres Gartner (a re-sourcer), suppression des marqueurs « a valider », prix / equipe / references / durees.
Usage (depuis docs/cadrage-ppt/) : py deck_delta/deck.py delta_V25 --base V24
"""
import deck_base
import deck_draw
from delta_V20 import bandeau, chapitre_de  # noqa: F401
from delta_V21 import sections
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt

BASE = "V24"
NOTE = "V25 : comite a 10 bras (ordre, texte, forme), intercalaires simplifies"
OPS = []

EMU = 12700
N = RGBColor(0x0E, 0x23, 0x56)
L = RGBColor(0xE7, 0xE9, 0xEE)
M = RGBColor(0x58, 0x65, 0x86)
W = RGBColor(0xFF, 0xFF, 0xFF)
CYAN = "00D2DD"

T_LEVIERS = "Six leviers au service de cinq résultats métier"
T_PIVOT = "Deux temps d'engagement : diagnostiquer, puis expérimenter et déployer"
T_SCENARIO = "Le scénario déroule six étapes, du problème au choix du client"
T_AGENTS = "Des agents augmentent le coach ; d'autres améliorent l'Infra du client"
T_PRINCIPES = "Trois principes de conduite : le service, l'adoption, le binôme"
T_DEFINITION = "Infra as a Product : des capacités pensées pour leurs consommateurs"
T_BOUCLE = "Une boucle continue pour décider sur des faits, expérimenter et ancrer"
T_PROCHAIN = "Prochain pas : choisir un périmètre, mesurer, tester, puis nous contacter"
T_PROCHAIN_NEUF = "Prochain pas : choisir, diagnostiquer, tester, puis nous contacter"
T_MODES = "Du pilote à l'autonomie : trois modes d'interaction entre OCTO et le client"
T_INDIC = "Nous mesurons la valeur client et l'apport du coaching augmenté"
T_OUVERTURE = "En trois temps et trois engagements, votre Infra passe du guichet au produit"
T_PB1 = "Dix problèmes, dix apports, dix indicateurs : 1 à 5"
T_PB2 = "Dix problèmes, dix apports, dix indicateurs : 6 à 10"
T_CHIFFRAGE = "Chiffrage : la forme du tableau jours × rôles, toutes valeurs à valider"
T_CONV13 = "Une conviction, un critère, un chantier : 1 à 3"
T_CONV45 = "Une conviction, un critère, un chantier : 4 et 5"
TITRES = {
    T_CONV13: "Trois convictions, chacune vérifiée par un critère et un chantier",
    T_CONV45: "Deux convictions de plus, vérifiées de la même façon",
    T_PB1: "Problèmes 1 à 5 : chacun trouve son apport et son indicateur",
    T_PB2: "Problèmes 6 à 10 : chacun trouve son apport et son indicateur",
    T_PROCHAIN: T_PROCHAIN_NEUF,
}


def _ids():
    return [f"id:{s.slide_id}" for s in Presentation(deck_base.resoudre_base(BASE)).slides]


def _est_inter(s):
    return s.slide_layout.name.startswith(("50 -", "51 -"))


def _rattacher(s, nom):
    """Pastille de navigation : le chapitre `nom` en aplat navy, les autres en clair."""
    for sh in s.shapes:
        if sh.name.startswith("nav-"):
            actif = sh.name == "nav-" + nom
            sh.fill.solid()
            sh.fill.fore_color.rgb = N if actif else L
            for r in sh.text_frame.paragraphs[0].runs:
                r.font.bold = actif
                r.font.color.rgb = W if actif else M


def intercalaires(prs):
    for s in prs.slides:
        if not _est_inter(s):
            continue
        for sh in list(s.shapes):
            if sh.name == "num-fantome" or sh.name.startswith("repere-"):
                sh._element.getparent().remove(sh._element)
        titre = next((sh for sh in s.shapes if sh.is_placeholder and sh.has_text_frame and sh.top > 100 * EMU and sh.text_frame.text), None)
        if titre is not None:
            titre.text_frame.margin_left = 0


def ordre(prs):
    leviers = deck_draw.slide_par_titre(prs, T_LEVIERS)
    deck_draw.deplacer_apres(prs, leviers, deck_draw.slide_par_titre(prs, T_DEFINITION))
    _rattacher(leviers, "Contexte et enjeux")
    pivot = deck_draw.slide_par_titre(prs, T_PIVOT)
    boucle = deck_draw.slide_par_titre(prs, T_BOUCLE)
    rang = list(prs.slides).index(boucle)
    avant = list(prs.slides)[rang - 1]  # l'intercalaire « Demarche generale »
    deck_draw.deplacer_apres(prs, pivot, avant)
    _rattacher(pivot, "Démarche générale")
    for titre in (T_SCENARIO, T_AGENTS, T_PRINCIPES):
        s = deck_draw.slide_par_titre(prs, titre)
        deck_draw.deplacer_apres(prs, s, list(prs.slides)[-1])
        _rattacher(s, "Annexes")


def _runs(s):
    for sh in s.shapes:
        if sh.has_text_frame and not sh.name.startswith("nav-"):
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    yield sh, r


def textes(prs):
    import re
    ouv = deck_draw.slide_par_titre(prs, T_OUVERTURE)
    carte = next(sh for sh in ouv.shapes if sh.has_text_frame and sh.text_frame.text.startswith("Aucun ROI promis"))
    deck_draw.set_run(carte, "Aucun ROI promis : référence au T0 (état initial mesuré), résultats au T1 (mesure après le pilote).")
    colonne = next(sh for sh in ouv.shapes if sh.has_text_frame and sh.text_frame.text.startswith("Sortir d'un guichet saturé"))
    for r in (r for p in colonne.text_frame.paragraphs for r in p.runs):
        if r.font.size is not None:
            r.font.size = Pt(round(r.font.size.pt * 0.8 * 2) / 2)  # retour a la taille d'avant la V21 : le test de charte estime le texte plus haut que sa boite
    from pptx.enum.text import MSO_ANCHOR
    colonne.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for titre, neuf in TITRES.items():
        s = deck_draw.slide_par_titre(prs, titre)
        deck_draw.set_run(s.shapes.title, neuf)
    for s in prs.slides:
        if _est_inter(s):
            continue
        for _, r in _runs(s):
            t = r.text
            n = re.sub(r"\bcoach augmenté\b", "coach produit augmenté", t)
            n = re.sub(r"\binfra\b", "Infra", n)
            if n != t:
                r.text = n
    ch = deck_draw.slide_par_titre(prs, T_CHIFFRAGE)
    for _, r in _runs(ch):
        t = r.text.strip()
        if t in ("0", "—"):
            r.text = "—"
        elif t == "0 %":
            r.text = "— %"
        elif t == "0 jour":
            r.text = "— jour"
        elif t.startswith("MONTANT : 0"):
            r.text = r.text.replace("MONTANT : 0", "MONTANT : —", 1)


def _round_rect(sh):
    g = sh._element.spPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom")
    return g is not None and g.get("prst") == "roundRect"


def rayons(prs):
    for s in list(prs.slides)[1:]:
        if _est_inter(s):
            continue
        for sh in s.shapes:
            if sh.shape_type == 1 and _round_rect(sh) and min(sh.width, sh.height) >= 36 * EMU:
                sh.adjustments[0] = min(0.5, 10 * EMU / min(sh.width, sh.height))


def pastilles(prs):
    for s in prs.slides:
        for sh in s.shapes:
            if sh.name == "pastille-rubrique" and sh.has_text_frame:
                droite = sh.left + sh.width
                for r in (r for p in sh.text_frame.paragraphs for r in p.runs):
                    r.font.size = Pt(9)
                largeur = int((14 + len(sh.text_frame.text) * 5.9) * EMU)
                sh.width = largeur
                sh.left = droite - largeur
                sh.height = int(15 * EMU)
                sh.top = sh.top - int(0.5 * EMU)


def convictions(prs):
    for titre in (TITRES[T_CONV13], TITRES[T_CONV45]):
        s = deck_draw.slide_par_titre(prs, titre)
        for sh in s.shapes:
            if sh.name in ("corps-carte", "texte-carte", "panneau-regle"):
                sh.height = int(sh.height + 36 * EMU)
        for sh in s.shapes:
            if sh.name == "texte-carte":
                for r in (r for p in sh.text_frame.paragraphs for r in p.runs):
                    if r.font.size is not None and r.font.size.pt >= 10:
                        r.font.size = Pt(12.5)
                    elif r.font.size is not None and r.font.size.pt >= 8:
                        r.font.size = Pt(9.5)


def problemes(prs):
    for titre in (TITRES[T_PB1], TITRES[T_PB2]):
        s = deck_draw.slide_par_titre(prs, titre)
        for sh in s.shapes:
            if sh.name in ("corps-carte", "texte-carte"):
                sh.height = int(sh.height + 12 * EMU)
            elif sh.name == "bande-lecture":
                sh.top = int(330 * EMU)
                sh.height = int(30 * EMU)


def modes(prs):
    s = deck_draw.slide_par_titre(prs, T_MODES)
    for sh in s.shapes:
        if sh.has_text_frame and abs(sh.top / EMU - 171) < 2 and sh.height > 80 * EMU:
            from pptx.enum.text import MSO_ANCHOR
            sh.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE


def indicateurs(prs):
    s = deck_draw.slide_par_titre(prs, T_INDIC)
    for sh in s.shapes:
        if sh.name == "zebre":
            sh.left, sh.width = int(45 * EMU), int(614 * EMU)
        elif sh.has_text_frame and sh.text_frame.text.startswith("NIVEAU 2"):
            sh.top = sh.top + int(3 * EMU)
        elif sh.name.startswith("Rounded Rectangle") and abs(sh.top / EMU - 290) < 2 and sh.width > 150 * EMU:
            sh.top = sh.top + int(3 * EMU)


def cyan(prs):
    """Cyan -> navy hors liserés fins, intercalaires et couverture ; le texte posé dessus (navy) passe en blanc."""
    for s in list(prs.slides)[1:]:
        if _est_inter(s):
            continue
        formes = list(s.shapes)
        cibles = []
        for sh in formes:
            if sh.shape_type != 1 or sh.name.startswith("nav-"):
                continue
            try:
                ok = sh.fill.type == 1 and str(sh.fill.fore_color.rgb) == CYAN
            except Exception:
                ok = False
            if ok and sh.width >= 8 * EMU and sh.height >= 8 * EMU:
                cibles.append(sh)
        for sh in cibles:
            sh.fill.fore_color.rgb = N
            zone = (sh.left, sh.top, sh.left + sh.width, sh.top + sh.height)
            for t in formes:
                if t is sh or not t.has_text_frame or not t.text_frame.text.strip():
                    continue
                cx, cy = t.left + t.width // 2, t.top + t.height // 2
                if zone[0] <= cx <= zone[2] and zone[1] <= cy <= zone[3] or t is sh:
                    for r in (r for p in t.text_frame.paragraphs for r in p.runs):
                        if r.font.color is not None and r.font.color.type is not None and str(r.font.color.rgb) == str(N):
                            r.font.color.rgb = W
            if sh.has_text_frame and sh.text_frame.text.strip():
                for r in (r for p in sh.text_frame.paragraphs for r in p.runs):
                    if r.font.color is not None and r.font.color.type is not None and str(r.font.color.rgb) == str(N):
                        r.font.color.rgb = W


def plancher(prs):
    """Texte entre 6 et 9 pt -> 9 pt (navigation, pastilles de rubrique et intercalaires exclus)."""
    for s in list(prs.slides)[1:]:
        if _est_inter(s):
            continue
        for sh, r in _runs(s):
            if r.font.size is not None and 6 <= r.font.size.pt < 9:
                r.font.size = Pt(9)


def retouche(prs):
    intercalaires(prs)
    ordre(prs)
    textes(prs)
    rayons(prs)
    pastilles(prs)
    convictions(prs)
    problemes(prs)
    modes(prs)
    indicateurs(prs)
    cyan(prs)
    plancher(prs)
    bandeau(prs)
    sections(prs)


ETAPES = [{"declarer": _ids(), "f": retouche, "nouvelles_max": 0, "nom": "comite 10 bras : ordre, texte, forme, intercalaires"}]
