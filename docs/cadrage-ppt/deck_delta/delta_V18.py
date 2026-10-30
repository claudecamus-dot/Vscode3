"""Delta V18 (base V17) : « Decommissionnements » se coupait en milieu de mot dans les cartes etroites des problemes : corps a 9,5 pt et marges de 6 pt sur ces deux slides."""
import deck_draw
from delta_V17 import TITRES
from pptx.util import Pt

BASE = "V17"
NOTE = "V18 : plus de mot coupe dans les cartes de problemes"
OPS = []
PROBLEMES = TITRES[2:]


def corps(prs):
    for titre in PROBLEMES:
        for sh in deck_draw.slide_par_titre(prs, titre).shapes:
            if sh.name == "texte-carte":
                sh.text_frame.margin_left = sh.text_frame.margin_right = Pt(6)
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.size == Pt(10.5):
                            r.font.size = Pt(9.5)


ETAPES = [{"declarer": PROBLEMES, "f": corps, "nouvelles_max": 0, "nom": "corps des cartes de problemes"}]
