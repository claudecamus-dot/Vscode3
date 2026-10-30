"""Delta V11 (base V10) : le chapitre « Preuve » est ventile dans « Demarche detaillee » (demande utilisateur du 2026-10-29).

- les 3 slides de Preuve passent dans « Demarche detaillee », juste apres le calendrier du pilote, dans leur ordre ; l'intercalaire « Preuve » disparait ;
- numerotation : les intercalaires gardent leurs numeros 01..07 ; « Demarche detaillee » prend le 05 (celui de Preuve) et le sommaire liste « Demarche generale » (04) puis
  « Demarche detaillee » (05) : les pastilles 01..07 du gabarit de sommaire restent toutes occupees ;
- navigation en haut reconstruite sur 8 chapitres, compteurs « n/N » et sections PowerPoint recalcules ; le renvoi « Preuve » de la slide « Chaque enjeu cite... » pointe vers « Demarche detaillee ».
Usage (depuis docs/cadrage-ppt/) : py deck_delta/deck.py delta_V11 --base V10
"""
import re

import deck_base
import deck_draw
import delta_V9
from delta_V9 import EMU, _chapitres, _est_intercalaire, _nom_chapitre
from delta_V10 import _kicker, _navigation, sections
from pptx import Presentation
from pptx.util import Pt

BASE = "V10"
NOTE = "V11 : chapitre Preuve ventile dans Demarche detaillee, 8 chapitres, sommaire et sections recalcules"
CHAPITRES = ["Contexte et enjeux", "Nos convictions", "Notre approche", "Démarche générale", "Démarche détaillée", "Dispositif et prix", "Prochain pas", "Annexes"]
T_CALENDRIER = "Calendrier du pilote : étapes et jalons, durées à valider"


def _intercalaire(prs, nom):
    return next(s for s in prs.slides if _est_intercalaire(s) and _nom_chapitre(s) == nom)


def ventiler(prs):
    """Les slides de Preuve (ids releves sur la base) suivent le calendrier du pilote ; l'intercalaire part par l'op `supprimer`."""
    par_id = {s.slide_id: s for s in prs.slides}
    ref = deck_draw.slide_par_titre(prs, T_CALENDRIER)
    for sid in PREUVE_IDS:
        deck_draw.deplacer_apres(prs, par_id[sid], ref)
        ref = par_id[sid]


def numeros_et_sommaire(prs):
    """« Demarche detaillee » reprend le numero 05 ; sommaire : 04 Demarche generale, 05 Demarche detaillee ; renvoi par nom de la slide « Chaque enjeu... »."""
    num = next(sh for sh in _intercalaire(prs, "Démarche détaillée").shapes if sh.has_text_frame and re.fullmatch(r"\d{2}", sh.text_frame.text))
    if num.text_frame.text != "05":
        deck_draw.set_run(num, "05")
    sommaire = next(s for s in prs.slides if s.slide_layout.name.startswith("95 -"))
    for sh in sommaire.shapes:
        if not sh.has_text_frame:
            continue
        if sh.text_frame.text == "Démarche générale et détaillée":
            deck_draw.set_run(sh, "Démarche générale")
        elif sh.text_frame.text == "Preuve":
            deck_draw.set_run(sh, "Démarche détaillée")
    for s in prs.slides:
        if s.slide_layout.name.startswith("95 -") or _est_intercalaire(s):
            continue
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text == "Preuve" and not sh.name.startswith("nav-"):
                deck_draw.set_run(sh, "Démarche détaillée")


def navigation_8(prs):
    """Retire la navigation de chaque slide, la redessine sur 8 chapitres (7 pt), puis recompte les kickers."""
    for s in prs.slides:
        for sh in [x for x in s.shapes if x.name.startswith("nav-")]:
            sh._element.getparent().remove(sh._element)
    delta_V9.CHAPITRES_NAV[:] = CHAPITRES
    delta_V9.navigation(prs)
    par_chapitre = {}
    for s, ch in _chapitres(prs):
        if ch:
            par_chapitre.setdefault(ch, []).append(s)
    for ch, slides in par_chapitre.items():
        for i, s in enumerate(slides, 1):
            for sh in s.shapes:
                if sh.name.startswith("nav-"):
                    for r in sh.text_frame.paragraphs[0].runs:
                        r.font.size = Pt(7)
            k = _kicker(s)
            nouveau = f"{ch.upper()} · {i}/{len(slides)}"
            if k is not None and k.text_frame.text != nouveau:
                deck_draw.set_run(k, nouveau)
            _navigation(s, ch)


def retouche(prs):
    ventiler(prs)
    numeros_et_sommaire(prs)
    navigation_8(prs)
    sections(prs)


def _base():
    prs = Presentation(deck_base.resoudre_base(BASE))
    preuve = [s.slide_id for s, ch in _chapitres(prs) if ch == "Preuve"]
    interc = _intercalaire(prs, "Preuve").slide_id
    return [f"id:{s.slide_id}" for s in prs.slides if s.slide_id != interc], preuve, interc  # l'op `supprimer` passe avant l'etape : ne pas declarer la slide supprimee


_TOUTES, PREUVE_IDS, _ID_INTERCALAIRE = _base()
OPS = [{"slide": f"id:{_ID_INTERCALAIRE}", "op": "supprimer"}]


ETAPES = [{"declarer": _TOUTES, "f": retouche, "nouvelles_max": 0, "nom": "Preuve ventile dans Demarche detaillee"}]
