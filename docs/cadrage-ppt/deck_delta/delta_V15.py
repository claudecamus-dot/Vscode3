"""Delta V15 (base V14) : finitions des cartes a en-tete (revue du rendu V14).

- le corps de carte commence sous l'en-tete (son trait navy ne depasse plus des coins arrondis de l'en-tete) ;
- les pictos du deck (poses a cheval sur le coin haut droit en V9) se retrouvent sous la pastille de rubrique : ils passent dans le corps de la carte, a droite.
"""
import deck_draw
from delta_V9 import EMU
from delta_V14 import HAUT_ENTETE, T_CONV, T_IA, T_REUSSIR
from pptx.util import Pt

BASE = "V14"
NOTE = "V15 : corps de carte sous l'en-tete, pictos dans le corps des cartes"
OPS = []


def finitions(prs):
    for titre in (T_CONV, T_REUSSIR, T_IA):
        s = deck_draw.slide_par_titre(prs, titre)
        formes = list(s.shapes)
        entetes = [h for h in formes if h.name == "entete-carte"]
        if len(entetes) != 5:
            raise ValueError(f"{titre!r} : {len(entetes)} en-tete(s), 5 exiges")
        for h in entetes:
            corps = next(b for b in formes if b.name.startswith("Rounded Rectangle") and b.left == h.left and b.width == h.width and b.top == h.top)
            corps.top, corps.height = h.top + HAUT_ENTETE, corps.height - HAUT_ENTETE
            cible_x, cible_y = h.left + h.width - 18 * EMU, h.top  # coin haut droit de la carte : ou V9 a pose le picto
            disque = next((d for d in formes if d.name.startswith("picto-") and not d.name.endswith("-glyphe")
                           and abs(d.left + d.width // 2 - cible_x) < 14 * EMU and abs(d.top + d.height // 2 - cible_y) < 14 * EMU), None)
            if disque is None:
                continue
            genre = disque.name
            dx = (h.left + h.width - 22 * EMU) - (disque.left + disque.width // 2)
            dy = (h.top + HAUT_ENTETE + 26 * EMU) - (disque.top + disque.height // 2)
            for f in formes:
                if f.name in (genre, genre + "-glyphe"):
                    f.left, f.top = f.left + dx, f.top + dy
            for t in formes:  # le texte du corps ne passe pas sous le picto
                if t.has_text_frame and t.text_frame.text.startswith("dans la démarche") and abs(t.left - (h.left + 50 * EMU)) < 3 * EMU and t.top < h.top + 60 * EMU:
                    t.width = h.width - 58 * EMU - 36 * EMU


ETAPES = [{"declarer": [T_CONV, T_REUSSIR, T_IA], "f": finitions, "nouvelles_max": 0, "nom": "finitions des cartes a en-tete"}]
