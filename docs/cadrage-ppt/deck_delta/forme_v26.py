"""Formes de la refonte V26 : A (cartes goutte), D (processus en chevrons), E (tableaux en bandes). Dimensions en points (slide 720x405, zone utile x 45..659).
Un seul cyan par slide (aplat decoratif, jamais sur du texte). Police Outfit. Repris des prototypes valides (comite forme, styles A, D, E)."""
from pptx.dml.color import RGBColor as C
from pptx.enum.shapes import MSO_SHAPE as S
from pptx.enum.text import MSO_ANCHOR as A
from pptx.enum.text import PP_ALIGN as P
from pptx.util import Emu, Pt

NAVY, CYAN, MUTED, CLAIR, BLANC = C(0x0E, 0x23, 0x56), C(0x00, 0xD2, 0xDD), C(0x58, 0x65, 0x86), C(0xE7, 0xE9, 0xEE), C(255, 255, 255)
FONT = "Outfit"
X0, X1 = 45, 659


def pt(v):
    return Emu(int(round(v * 12700)))


def garder(sh):
    """Formes conservees quand on refait le corps d'une slide : titre/placeholders, navigation, surtitre numerote, etiquette d'etat (haut droite)."""
    if sh.is_placeholder or sh.name.startswith("nav-"):
        return True
    if sh.top < 85 * 12700 and sh.top > 60 * 12700 and (sh.left > 500 * 12700 or (sh.has_text_frame and " · " in sh.text_frame.text and sh.text_frame.text[:2].isdigit())):
        return True
    return sh.has_text_frame and sh.text_frame.text[:2].isdigit() and " · " in sh.text_frame.text and sh.top < 80 * 12700


def vider(slide):
    for sh in list(slide.shapes):
        if not garder(sh):
            sh._element.getparent().remove(sh._element)


def forme(s, kind, x, y, w, h, fill, ligne=None, adj=None, lw=1.25):
    r = s.shapes.add_shape(kind, pt(x), pt(y), pt(w), pt(h))
    if fill is None:
        r.fill.background()
    else:
        r.fill.solid()
        r.fill.fore_color.rgb = fill
    if ligne is None:
        r.line.fill.background()
    else:
        r.line.color.rgb = ligne
        r.line.width = Pt(lw)
    r.shadow.inherit = False
    if adj is not None:
        r.adjustments[0] = adj
    return r


def goutte(s, x, y, w, h, fill, rayon=28, ligne=None):
    return forme(s, S.ROUND_2_DIAG_RECTANGLE, x, y, w, h, fill, ligne, adj=min(0.5, rayon / min(w, h)))


def capsule(s, x, y, w, h, fill, ligne=None):
    return forme(s, S.ROUNDED_RECTANGLE, x, y, w, h, fill, ligne, adj=0.5)


def rond(s, cx, cy, d, fill, ligne=None, lw=2.25):
    return forme(s, S.OVAL, cx - d / 2, cy - d / 2, d, d, fill, ligne, lw=lw)


def tx(s, x, y, w, h, paras, ancre=A.TOP, align=P.LEFT, interligne=None, espace=0):
    """paras = [[(texte, taille, gras, couleur), ...], ...] ; aucune marge."""
    tb = s.shapes.add_textbox(pt(x), pt(y), pt(w), pt(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = ancre
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if interligne:
            p.line_spacing = interligne
        if i and espace:
            p.space_before = Pt(espace)
        for t, sz, b, col in runs:
            r = p.add_run()
            r.text = t
            r.font.size, r.font.bold, r.font.name = Pt(sz), b, FONT
            r.font.color.rgb = col
    return tb


def signature_cyan(s, x=X0, y=86):
    return capsule(s, x, y, 44, 5, CYAN)


def pastille(s, cx, cy, d, num, sombre=True):
    rond(s, cx, cy, d, NAVY if sombre else BLANC, BLANC if sombre else NAVY)
    return tx(s, cx - d / 2, cy - d / 2, d, d, [[(str(num), round(d * 0.42, 1), True, BLANC if sombre else NAVY)]], A.MIDDLE, P.CENTER)


def etiquette(s, x, y, w, texte_, h=22, sz=10.5):
    capsule(s, x, y, w, h, BLANC, NAVY)
    return tx(s, x, y, w, h, [[(texte_, sz, True, NAVY)]], A.MIDDLE, P.CENTER)


def chevron(s, x, y, w, h, fill, premier=False):
    return forme(s, S.PENTAGON if premier else S.CHEVRON, x, y, w, h, fill, None, adj=0.35)


def bas_de_conviction(s, critere, chantier):
    """Pied commun des slides de conviction : le critere observable (carte claire) et le chantier (carte navy), a y 284."""
    goutte(s, X0, 284, 300, 76, CLAIR, 20)
    etiquette(s, X0 + 14, 273, 168, "LE CRITÈRE OBSERVABLE")
    tx(s, X0 + 16, 302, 270, 52, [[(critere, 11, False, NAVY)]], A.MIDDLE)
    goutte(s, X0 + 314, 284, 300, 76, NAVY, 20)
    etiquette(s, X0 + 328, 273, 110, "LE CHANTIER")
    tx(s, X0 + 330, 302, 270, 52, [[(chantier, 11, False, BLANC)]], A.MIDDLE)


def trois_cartes(s, cartes, y=106, h=156):
    """Trois cartes goutte (la deuxieme en accent navy). cartes = [(etiquette, titre, [lignes])]."""
    W, G = 190, 22
    for k, (lab, titre, lignes) in enumerate(cartes):
        x = X0 + k * (W + G)
        sombre = k == 1
        col = BLANC if sombre else NAVY
        goutte(s, x, y, W, h, NAVY if sombre else CLAIR, 28)
        etiquette(s, x + 14, y - 11, 118, lab)
        tx(s, x + 16, y + 20, W - 32, 34, [[(titre, 13, True, col)]])
        tx(s, x + 16, y + 58, W - 32, h - 64, [[(t, 11, False, col)] for t in lignes], espace=4)


# ---------------------------------------------------------------- pictogrammes natifs (blanc sur disque navy) et mini-carte des convictions
def disque_picto(s, cx, cy, glyphe, d=34):
    """Disque navy d = 34 portant un pictogramme dessine en formes natives blanches (aucune image)."""
    rond(s, cx, cy, d, NAVY, None)
    b = BLANC
    if glyphe == "grille":
        for i in range(2):
            for j in range(2):
                forme(s, S.ROUNDED_RECTANGLE, cx - 8 + i * 9, cy - 8 + j * 9, 7, 7, b, adj=0.2)
    elif glyphe == "barres":
        for i, h in enumerate((6, 10, 14)):
            forme(s, S.RECTANGLE, cx - 8 + i * 6, cy + 7 - h, 4, h, b)
    elif glyphe == "document":
        forme(s, S.FLOWCHART_DOCUMENT, cx - 6, cy - 8, 12, 16, b)
    elif glyphe == "cible":
        rond(s, cx, cy, 18, b, None)
        rond(s, cx, cy, 11, NAVY, None)
        rond(s, cx, cy, 4, b, None)
    elif glyphe == "calendrier":
        forme(s, S.ROUNDED_RECTANGLE, cx - 8, cy - 7, 16, 15, b, adj=0.15)
        forme(s, S.RECTANGLE, cx - 8, cy - 7, 16, 4, NAVY)
        for i in range(3):
            forme(s, S.RECTANGLE, cx - 5 + i * 4, cy + 1, 2, 2, NAVY)
    elif glyphe == "jauge":
        forme(s, S.BLOCK_ARC, cx - 9, cy - 8, 18, 18, b)
        forme(s, S.RECTANGLE, cx - 1, cy - 1, 2, 8, b)
    elif glyphe == "personnes":
        rond(s, cx, cy - 4, 8, b, None)
        forme(s, S.ROUNDED_RECTANGLE, cx - 8, cy + 2, 16, 8, b, adj=0.5)
    elif glyphe == "fleche":
        forme(s, S.RIGHT_ARROW, cx - 9, cy - 6, 18, 12, b)


def minicarte(s, actifs, total=12, x=388, y=70, pas=11, d=8):
    """Repere de lecture des slides de conviction : une pastille par conviction, navy pour celle(s) de la slide, claire pour les autres."""
    for k in range(1, total + 1):
        forme(s, S.OVAL, x + (k - 1) * pas, y, d, d, NAVY if k in actifs else CLAIR, None)
