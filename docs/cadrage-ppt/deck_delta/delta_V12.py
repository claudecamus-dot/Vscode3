"""Delta V12 (base V11) : reprises du comite sur les PPT de Telechargements (accord utilisateur « ok continue », 2026-10-29).

- nouvelle slide de « Demarche detaillee » : « Du pilote a l'autonomie : trois modes d'interaction entre OCTO et le client » (collaboration, facilitation, X-as-a-Service :
  les trois modes d'interaction entre equipes de Team Topologies ; idee relevee dans un support BPCE IT, aucun contenu client repris), sur le gabarit de « Deux temps d'engagement ».
- slide « Fiche reference » : champs precises d'apres les fiches OCTO lues (bandeau d'identite, mission en 3 etapes, resultats avec valeurs estimees et horizon, dispositif) ;
  aucune reference n'est inscrite : elles restent « a fournir par OCTO » (confidentialite a valider).
Usage (depuis docs/cadrage-ppt/) : py deck_delta/deck.py delta_V12 --base V11
"""
import deck_base
import deck_draw
from delta_V10 import sections
from delta_V11 import navigation_8
from pptx import Presentation

BASE = "V11"
NOTE = "V12 : slide des trois modes d'interaction (passage de main), fiche reference precisee"
OPS = []

T_MODELE = "Deux temps d'engagement : diagnostiquer, puis expérimenter et déployer"
T_APRES = "Ce que nous attendons de vous, et comment nous passons la main"
T_REFERENCE = "Fiche référence : un gabarit prêt à recevoir les références d'OCTO"
T_MODES = "Du pilote à l'autonomie : trois modes d'interaction entre OCTO et le client"
NB = "\xa0"


def _textes(formes, remplacements):
    """Remplace par INDICE de forme (l'ordre est celui de la slide modele ; liste relevee AVANT toute suppression)."""
    for j, texte in remplacements.items():
        deck_draw.set_run(formes[j], texte)


def slide_modes(prs):
    modele = deck_draw.slide_par_titre(prs, T_MODELE)
    rang = list(prs.slides).index(modele)
    neuve = deck_base.core().copier_slide(prs, modele, index=rang + 1)
    formes = list(neuve.shapes)
    _textes(formes, {
        0: T_MODES,
        2: "OCTO PORTE LE DÉBUT",
        4: "LE CLIENT PREND LA MAIN, PAS À PAS",
        7: "Collaboration", 17: "Facilitation", 27: "X-as-a-Service",
        10: "Le coach OCTO et les équipes du client construisent le pilote ensemble, le temps d'apprendre à décider sur des faits.",
        20: "Les relais formés du client portent l'accompagnement ; OCTO outille, conseille et reste en appui.",
        30: "L'équipe plateforme du client fournit la capacité comme un service ; les consommateurs l'utilisent sans passer par OCTO.",
        11: f"Conduit par{NB}: OCTO et le client ensemble.",
        21: f"Conduit par{NB}: le client, OCTO en appui.",
        31: f"Conduit par{NB}: le client, avec ses relais.",
        14: "pilote", 24: "extension éventuelle", 33: "autonomie",
        40: "Modes de Team Topologies ; répartition à valider avec le client.",
    })
    for sh in formes[34:40]:  # rangee des montants : sans objet ici
        sh._element.getparent().remove(sh._element)
    deck_draw.deplacer_apres(prs, neuve, deck_draw.slide_par_titre(prs, T_APRES))


def fiche_reference(prs):
    s = deck_draw.slide_par_titre(prs, T_REFERENCE)
    _textes(list(s.shapes), {
        4: f"Une référence par fiche, client anonymisé par secteur tant qu'OCTO ne l'a pas validé{NB}; aucun cas n'est présenté tant qu'OCTO ne l'a pas fourni.",
        9: "Contexte, enjeux et demande (4 lignes au plus)",
        16: "Mission OCTO en trois étapes",
        23: "Résultats et bénéfices, valeurs estimées avec horizon",
        30: "Dispositif : équipe, durée et périmètre",
    })


def retouche(prs):
    slide_modes(prs)
    fiche_reference(prs)
    navigation_8(prs)
    sections(prs)


def _ids():
    return [f"id:{s.slide_id}" for s in Presentation(deck_base.resoudre_base(BASE)).slides]


ETAPES = [{"declarer": _ids(), "f": retouche, "nouvelles_max": 1, "nom": "trois modes d'interaction, fiche reference"}]
