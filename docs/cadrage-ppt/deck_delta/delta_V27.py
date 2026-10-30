"""Delta V27 (base V26) : demande utilisateur du 2026-10-30 — formes d'OCTO_GROUND (comite G1), slide d'introduction des convictions, conviction « culture produit »,
planning hypothetique sans cout, exemples de livrables en encarts (modeles releves dans le deck de Vscode8, comite G2).

- chapitre « Nos convictions » : la slide d'ouverture devient l'introduction des DOUZE convictions regroupees en quatre themes ; une conviction de culture produit est ajoutee en tete
  (les autres sont renumerotees 2 a 12, titres et pastilles) ; une mini-carte de douze pastilles rappelle, sur chaque slide de conviction, laquelle est declinee ;
- Demarche generale : planning hypothetique en trois phases sur quatre trimestres, SANS cout (tampon « hypothese », durees a valider, encarts livrables et hypotheses) ;
- Demarche detaillee : deux slides d'exemples de livrables (huit encarts : finalite, format, qui, quand), pictogrammes dessines en formes natives.
Usage (depuis docs/cadrage-ppt/) : py deck_delta/deck.py delta_V27 --base V26
"""
import re

import deck_base
import deck_draw
from delta_V20 import bandeau
from delta_V21 import sections
from forme_v26 import *  # noqa: F401,F403
from pptx import Presentation

BASE = "V26"
NOTE = "V27 : introduction des convictions, conviction culture produit, planning hypothetique, exemples de livrables"
OPS = []

T_INTRO_ANCIEN = "Cinq convictions : une cible à construire et une démarche pour y conduire"
T_INTRO = "Douze convictions en quatre thèmes, déclinées une à une ensuite"
T_MODELE = "Deux convictions de plus, vérifiées de la même façon"
T_TROIS = "Trois convictions, chacune vérifiée par un critère et un chantier"
T_DEUX_TEMPS = "Deux temps d'engagement : diagnostiquer, puis expérimenter et déployer"
T_PILOTE = "De la douleur au produit : une mission pilote de bout en bout"
T_CULTURE = "Conviction 1 : le mode produit est d'abord une culture, pas un outil"
T_PLANNING = "Un planning hypothétique en trois phases, avant tout chiffrage"
T_LIV1 = "Exemples de livrables (1/2) : comprendre et décider"
T_LIV2 = "Exemples de livrables (2/2) : piloter et ancrer"
ANCIENS = ["Conviction 6 : financer des produits et des équipes stables", "Conviction 7 : une petite équipe produit qui construit et exploite",
           "Conviction 8 : trois instances, chacune avec sa décision", "Conviction 9 : un flux d'entrée unique qui qualifie avant de servir",
           "Conviction 10 : piloter le service vécu, pas seulement le SLA", "Conviction 11 : aucun produit livré sans RUN prêt"]


def _ids():
    return [f"id:{s.slide_id}" for s in Presentation(deck_base.resoudre_base(BASE)).slides]


# ---------------------------------------------------------------- renumerotation et mini-carte
def renumeroter(prs):
    for ancien in ANCIENS:
        s = deck_draw.slide_par_titre(prs, ancien)
        n = int(re.search(r"Conviction (\d+)", ancien).group(1)) + 1
        deck_draw.set_run(s.shapes.title, re.sub(r"Conviction \d+", f"Conviction {n}", ancien))
    for titre in (T_TROIS, T_MODELE):
        s = deck_draw.slide_par_titre(prs, titre)
        for sh in s.shapes:
            if sh.has_text_frame and re.fullmatch(r"CONVICTION \d", sh.text_frame.text):
                deck_draw.set_run(sh, f"CONVICTION {int(sh.text_frame.text[-1]) + 1}")


def mini_cartes(prs):
    cibles = {T_CULTURE: {1}, T_TROIS: {2, 3, 4}, T_MODELE: {5, 6}}
    for k, ancien in enumerate(ANCIENS):
        cibles[re.sub(r"Conviction \d+", f"Conviction {k + 7}", ancien)] = {k + 7}
    for titre, actifs in cibles.items():
        minicarte(deck_draw.slide_par_titre(prs, titre), actifs)


# ---------------------------------------------------------------- introduction des douze convictions
THEMES = [
    ("CULTURE ET ORGANISATION", [(1, "La culture produit avant les outils", None), (7, "Financer des produits et des équipes stables", None),
                                 (8, "Une petite équipe qui construit et exploite", None)]),
    ("LE PRODUIT ET L'USAGE", [(2, "Partir des douleurs des utilisateurs", "Boucle : Explorer"), (3, "Un porteur produit pour chaque capacité", "Boucle : Ancrer")]),
    ("PILOTAGE ET DEMANDE", [(4, "Décider sur des faits mesurés", "Boucle : Prioriser"), (9, "Trois instances, trois décisions", None),
                            (10, "Un flux d'entrée qui qualifie avant de servir", None)]),
    ("RUN, PROGRÈS ET IA", [(5, "Un progrès mesuré, régulier, transparent", "Boucle : Mesurer"), (6, "L'IA gagne du temps là où il est perdu", "Boucle : Comprendre"),
                           (11, "Piloter le service vécu, pas le seul SLA", None), (12, "Aucun produit livré sans RUN prêt", None)]),
]


def introduction(prs):
    s = deck_draw.slide_par_titre(prs, T_INTRO_ANCIEN)
    deck_draw.set_run(s.shapes.title, T_INTRO)
    vider(s)
    signature_cyan(s)
    capsule(s, 536, 67, 122, 17, CLAIR)
    tx(s, 536, 67, 122, 17, [[("PROPOSITION À VALIDER", 9, True, NAVY)]], A.MIDDLE, P.CENTER)
    W, G = 143, 14
    for k, (theme, lignes) in enumerate(THEMES):
        x = X0 + k * (W + G)
        capsule(s, x, 98, W, 24, NAVY)
        tx(s, x, 98, W, 24, [[(theme, 9.5, True, BLANC)]], A.MIDDLE, P.CENTER)
        for r, (n, these, boucle) in enumerate(lignes):
            y = 132 + r * 57
            goutte(s, x, y, W, 54, CLAIR, 14)
            pastille(s, x + 17, y + 27, 22, n, sombre=True)
            paras = [[(these, 10, True, NAVY)]] + ([[(boucle, 9.5, False, MUTED)]] if boucle else [])
            tx(s, x + 34, y + 4, W - 40, 46, paras, A.MIDDLE)


# ---------------------------------------------------------------- conviction 1 : culture produit
def culture(prs):
    ref = deck_draw.slide_par_titre(prs, T_INTRO)
    rang = list(prs.slides).index(ref)
    neuve = deck_base.core().copier_slide(prs, deck_draw.slide_par_titre(prs, T_MODELE), index=rang + 1)
    vider(neuve)
    deck_draw.set_run(neuve.shapes.title, T_CULTURE)
    signature_cyan(neuve)
    trois_cartes(neuve, [
        ("LE CHANGEMENT", "Une culture, pas un outil", ["Le mode produit change les décisions, les rôles et les habitudes.", "Les outils viennent ensuite."]),
        ("LE SPONSOR", "Une règle posée d'en haut", ["Chaque investissement est discuté en termes de valeur pour l'usager.", "Le sponsor porte ce message."]),
        ("LA PREUVE", "Des managers formés", ["Les managers apprennent à piloter par la valeur.", "La valeur est mesurée au T0, puis au T1."]),
    ])
    bas_de_conviction(neuve,
                      "Chaque équipe formule la valeur rendue à ses consommateurs et la revoit à chaque cycle ; les décisions citent un indicateur.",
                      "Atelier de cadrage de la valeur par capacité, message du sponsor, coaching des managers et des équipes sur le terrain.")


# ---------------------------------------------------------------- planning hypothetique sans cout
def planning(prs):
    modele = deck_draw.slide_par_titre(prs, T_DEUX_TEMPS)
    rang = list(prs.slides).index(modele)
    s = deck_base.core().copier_slide(prs, modele, index=rang + 1)
    vider(s)
    deck_draw.set_run(s.shapes.title, T_PLANNING)
    signature_cyan(s)
    capsule(s, 489, 94, 170, 20, BLANC, NAVY)
    tx(s, 489, 94, 170, 20, [[("HYPOTHÈSE · SANS CHIFFRAGE", 10, True, NAVY)]], A.MIDDLE, P.CENTER)
    LX, TX0, TW = X0, 165, 494
    q = TW / 4
    for k in range(4):
        capsule(s, TX0 + k * q + 2, 122, q - 4, 20, CLAIR)
        tx(s, TX0 + k * q + 2, 122, q - 4, 20, [[(f"Trimestre {k + 1}", 10, True, NAVY)]], A.MIDDLE, P.CENTER)
    barres = [("Assessment flash", TX0, 100, True), ("Pilote Produit + IA", TX0 + 108, 160, True), ("Ancrage", TX0 + 276, TW - 276, True)]
    for r, (nom, x, w, _) in enumerate(barres):
        y = 150 + r * 30
        tx(s, LX, y, 116, 24, [[(nom, 11, True, NAVY)]], A.MIDDLE)
        forme(s, S.RECTANGLE, TX0, y + 23, TW, 1, CLAIR, None)
        capsule(s, x, y, w, 24, NAVY)
        tx(s, x, y, w, 24, [[("durée à valider", 10, True, BLANC)]], A.MIDDLE, P.CENTER)
    y = 240
    tx(s, LX, y, 116, 24, [[("Pilotage", 11, True, NAVY)]], A.MIDDLE)
    for x, lab in ((TX0 + 86, "Choix du pilote"), (TX0 + 254, "Comité de décision"), (TX0 + TW - 14, "Passage de main")):
        forme(s, S.DIAMOND, x, y + 5, 14, 14, NAVY, None)
    tx(s, TX0 + 104, y, 90, 24, [[("choix du pilote", 10, False, MUTED)]], A.MIDDLE)
    tx(s, TX0 + 272, y, 120, 24, [[("comité de décision", 10, False, MUTED)]], A.MIDDLE)
    tx(s, TX0 + TW - 122, y, 104, 24, [[("passage de main", 10, False, MUTED)]], A.MIDDLE, P.RIGHT)
    goutte(s, X0, 284, 300, 76, NAVY, 20)
    etiquette(s, X0 + 14, 273, 130, "LIVRABLES CLÉS")
    tx(s, X0 + 16, 298, 270, 58, [[(t, 10.5, False, BLANC)] for t in ("Assessment flash : baseline, pilote choisi", "Pilote : produit Infra livré, indicateurs T0 → T1", "Ancrage : relais formés, passage de main")], A.MIDDLE, espace=2)
    goutte(s, X0 + 314, 284, 300, 76, CLAIR, 20)
    etiquette(s, X0 + 328, 273, 190, "HYPOTHÈSES DE CALENDRIER")
    tx(s, X0 + 330, 298, 270, 58, [[("• " + t, 10.5, False, NAVY)] for t in ("Démarrage après le choix du périmètre", "Durées à caler avec vous", "Le chiffrage est traité à part")], A.MIDDLE, espace=2)


# ---------------------------------------------------------------- exemples de livrables en encarts
LIVRABLES_1 = [
    ("Cartographie des capacités", "grille", "Voir les capacités, leurs consommateurs et leurs propriétaires.", "Format : tableau", "Qui : ensemble", "Quand : début"),
    ("Baseline des indicateurs", "barres", "Fixer l'état de départ (T0) pour mesurer ensuite l'apport.", "Format : tableau", "Qui : OCTO", "Quand : début"),
    ("Fiche produit Infra", "document", "Décrire l'offre, la valeur, les engagements et le propriétaire.", "Format : fiche", "Qui : ensemble", "Quand : pilote"),
    ("Grille de priorisation", "cible", "Arbitrer les demandes sur la valeur, la faisabilité et la charge.", "Format : grille", "Qui : ensemble", "Quand : pilote"),
]
LIVRABLES_2 = [
    ("Roadmap produit", "calendrier", "Planifier les évolutions et l'arrêt des capacités.", "Format : frise", "Qui : client", "Quand : pilote"),
    ("Tableau de bord", "jauge", "Suivre l'usage, la qualité et le coût, du T0 au T1.", "Format : écran", "Qui : client", "Quand : continu"),
    ("Charte IA et RACI", "personnes", "Cadrer qui prépare, qui valide et qui décide quand l'IA aide.", "Format : charte", "Qui : ensemble", "Quand : pilote"),
    ("Plan de passage de main", "fleche", "Transférer l'animation aux relais formés du client.", "Format : plan", "Qui : OCTO", "Quand : fin"),
]


def encarts_livrables(s, liste):
    for k, (nom, glyphe, sert, f, qui, quand) in enumerate(liste):
        x, y = X0 + (k % 2) * 314, 98 + (k // 2) * 132
        goutte(s, x, y, 300, 124, CLAIR, 22)
        disque_picto(s, x + 34, y + 36, glyphe)
        tx(s, x + 62, y + 12, 226, 22, [[(nom, 13, True, NAVY)]], A.MIDDLE)
        tx(s, x + 62, y + 38, 226, 40, [[(sert, 10.5, False, NAVY)]])
        for j, c in enumerate((f, qui, quand)):
            cx = x + 14 + j * 96
            capsule(s, cx, y + 90, 92, 22, BLANC, NAVY)
            tx(s, cx, y + 90, 92, 22, [[(c, 9.5, True, NAVY)]], A.MIDDLE, P.CENTER)
    tx(s, X0, 362, 614, 14, [[("Exemples illustratifs : formats et contenus à adapter en atelier avec vous.", 10, False, MUTED)]])


def livrables(prs):
    modele = deck_draw.slide_par_titre(prs, T_PILOTE)
    rang = list(prs.slides).index(modele)
    for k, (titre, liste) in enumerate(((T_LIV1, LIVRABLES_1), (T_LIV2, LIVRABLES_2))):
        s = deck_base.core().copier_slide(prs, modele, index=rang + 1 + k)
        vider(s)
        deck_draw.set_run(s.shapes.title, titre)
        signature_cyan(s)
        encarts_livrables(s, liste)


def retouche(prs):
    introduction(prs)
    culture(prs)
    renumeroter(prs)
    mini_cartes(prs)
    planning(prs)
    livrables(prs)
    bandeau(prs)
    sections(prs)


ETAPES = [{"declarer": _ids(), "f": retouche, "nouvelles_max": 4, "nom": "introduction des convictions, culture produit, planning, livrables"}]
