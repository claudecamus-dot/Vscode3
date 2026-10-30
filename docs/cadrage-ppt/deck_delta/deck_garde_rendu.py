"""Garde geometrique d'un deck : formes hors slide et textes hors boite (port adapte de Vscode8 `scripts/garde_rendu.py`).

Reutilise `verifier_geometrie` / `verifier_debordements_texte` de la skill pptx-deck du projet (canon, non modifie). Ne remplace PAS le rendu reel.
Adaptation : le chemin de la skill se deduit de deck_base.RACINE ; `compter_defauts` renvoie le nombre de defauts SANS rien imprimer, pour que deck.py compare
la base et le resultat (un defaut deja present dans la base n'est pas un defaut du delta). Renomme `deck_garde_rendu` (Vscode8 : `garde_rendu`) pour ne pas
masquer `pptx-verify/garde_rendu.py` du hub, un autre outil.
"""
import os
import sys

import deck_base
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

SKILL = os.path.join(deck_base.RACINE, ".claude", "skills", "pptx-deck", "scripts")


def _pptx_deck():
    if SKILL not in sys.path:
        sys.path.append(SKILL)
    import pptx_deck
    return pptx_deck


def _lignes_nulles(prs):
    """Cles « slide N: 'nom' dimension non positive » des connecteurs/lignes de largeur OU hauteur nulle (nulle par construction : faux positif)."""
    cles = set()
    for si, slide in enumerate(prs.slides, start=1):
        for shp in slide.shapes:
            balise = shp._element.tag.rsplit("}", 1)[-1]
            if balise == "cxnSp" or getattr(shp, "shape_type", None) == MSO_SHAPE_TYPE.LINE:
                try:
                    nul = shp.width == 0 or shp.height == 0
                except Exception:
                    nul = False
                if nul:
                    cles.add(f"slide {si}: '{shp.name or 'shape'}' dimension non positive")
    return cles


def detecter_defauts(prs):
    """Liste de (genre, detail) : formes hors slide et textes hors boite, hors connecteurs/lignes de dimension nulle."""
    D = _pptx_deck()
    nulles = _lignes_nulles(prs)
    geo = [p for p in D.verifier_geometrie(prs) if not any(p.startswith(c) for c in nulles)]
    return [("forme hors slide", p) for p in geo] + [("texte hors boite", p) for p in D.verifier_debordements_texte(prs)]


def compter_defauts(chemin_pptx):
    """Nombre de defauts d'un fichier (aucune sortie)."""
    return len(detecter_defauts(Presentation(str(chemin_pptx))))
