"""Delta V16 (base V15) : deux slides trop chargees scindees en deux, redessinees en cartes a en-tete et pastille (modele AGIMA 35, voir delta_V14).

- « Chaque conviction se verifie par un critere et devient un chantier » (1 410 caracteres) -> « Une conviction, un critere, un chantier : 1 a 3 » puis « : 4 et 5 » (+ regle de passage) ;
- « Dix problemes, dix apports, dix indicateurs a suivre » (1 104 caracteres, tableau de 10 lignes) -> « ... : 1 a 5 » puis « : 6 a 10 », une carte par probleme.
Le texte est repris mot pour mot, seul le decoupage change. Mesure : caracteres par slide relevés par python-pptx sur la V15 (11 : 1 410 ; 38 : 1 104 ; les plus chargees du deck).
Usage (depuis docs/cadrage-ppt/) : py deck_delta/deck.py delta_V16 --base V15
"""
import deck_base
import deck_draw
from delta_V9 import BLANC, CYAN, EMU, MUTED, NAVY, POLICE
from delta_V10 import sections
from delta_V11 import navigation_8
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt

BASE = "V15"
NOTE = "V16 : slides 11 et 38 scindees en deux, cartes a en-tete et pastille"
OPS = []

T11 = "Chaque conviction se vérifie par un critère et devient un chantier"
T38 = "Dix problèmes, dix apports, dix indicateurs à suivre"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _run(p, texte, taille, couleur, gras):
    r = p.add_run()
    r.text = texte
    r.font.size, r.font.bold, r.font.name = Pt(taille), gras, POLICE
    r.font.color.rgb = couleur


def _cadre(sh, ancre, marges):
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Pt(marges[0])
    tf.margin_top = tf.margin_bottom = Pt(marges[1])
    tf.vertical_anchor = ancre
    return tf


def carte(s, x, y, w, h, haut, titre, blocs, tag=None, taille_titre=10):
    """Carte a en-tete (AGIMA 35) : corps blanc au trait navy fin sous l'en-tete navy de meme largeur, pastille de rubrique blanche a cheval sur le bord haut droit."""
    corps = deck_draw.rect(s, x, y + haut, w, h - haut, BLANC, MSO_SHAPE.ROUND_2_SAME_RECTANGLE, ligne=NAVY)
    corps.line.width = Pt(0.75)
    geom = corps._element.spPr.find(f"{A}prstGeom")
    for gd in list(geom.find(f"{A}avLst")):
        geom.find(f"{A}avLst").remove(gd)
    corps.adjustments[0] = 0
    corps.adjustments[1] = min(0.5, 12 / min(w, h - haut))
    corps.name = "corps-carte"
    entete = deck_draw.rect(s, x, y, w, haut, NAVY, MSO_SHAPE.ROUND_2_SAME_RECTANGLE)
    entete.name = "entete-carte"
    entete.adjustments[0] = min(0.5, 12 / haut)
    entete.adjustments[1] = 0
    _run(_cadre(entete, MSO_ANCHOR.MIDDLE, (8, 3)).paragraphs[0], titre, taille_titre, BLANC, True)
    if tag:
        lp = 14 + len(tag) * 4.4
        pastille = deck_draw.rect(s, x + w - lp - 8, y - 7, lp, 14, BLANC, MSO_SHAPE.ROUNDED_RECTANGLE, ligne=NAVY)
        pastille.name = "pastille-rubrique"
        pastille.adjustments[0] = 0.5
        pastille.line.width = Pt(0.75)
        p = _cadre(pastille, MSO_ANCHOR.MIDDLE, (2, 0)).paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        _run(p, tag, 6.5, NAVY, True)
    zone = deck_draw.rect(s, x, y + haut, w, h - haut, BLANC, MSO_SHAPE.RECTANGLE)
    zone.fill.background()
    zone.name = "texte-carte"
    tf = _cadre(zone, MSO_ANCHOR.TOP, (8, 6))
    for k, (etiquette, texte) in enumerate(blocs):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        if k:
            p.space_before = Pt(8)
        _run(p, etiquette, 7, MUTED, True)
        q = tf.add_paragraph()
        _run(q, texte, 9.5, NAVY, False)


def _vider(s, garder):
    formes = list(s.shapes)
    for j, sh in enumerate(formes):
        if j not in garder and not sh.name.startswith("nav-"):
            sh._element.getparent().remove(sh._element)


def _scinder(prs, titre_modele, garder, titre1, titre2):
    """Vide la slide modele (hors titre, kicker, navigation et formes gardees), la copie juste apres, retitre les deux. Renvoie (slide1, slide2)."""
    s1 = deck_draw.slide_par_titre(prs, titre_modele)
    _vider(s1, garder)
    rang = list(prs.slides).index(s1)
    s2 = deck_base.core().copier_slide(prs, s1, index=rang + 1)
    deck_draw.set_run(next(sh for sh in s1.shapes if sh.is_placeholder), titre1)
    deck_draw.set_run(next(sh for sh in s2.shapes if sh.is_placeholder), titre2)
    return s1, s2


CONVICTIONS = [
    ("Une plateforme n'est adoptée que si elle résout une douleur mesurée chez ceux qui la consomment.",
     "Les douleurs de chaque consommateur sont documentées, avec une mesure.", "Interviews et situation de référence mesurée."),
    ("Sans porteur produit, l'infra retourne au guichet.",
     "Chaque capacité clé a un propriétaire nommé, une roadmap et un indicateur.", "Désigner les propriétaires, bâtir la roadmap, définir l'indicateur de valeur."),
    ("Les priorités se décident sur des faits mesurés, pas sur le rapport de force.",
     "Une grille de priorisation partagée alimente les arbitrages budgétaires.", "Choisir les indicateurs, scorer et prioriser, ritualiser la décision."),
    ("Le progrès se mesure de façon régulière et transparente, à chaque niveau.",
     "Des indicateurs de progrès sont suivis à rythme régulier, succès comme erreurs.", "Installer la mesure et un rituel de revue qui rend les équipes redevables."),
    ("L'IA gagne du temps là où il est perdu : mesuré puis vérifié, pas promis.",
     "Les gisements de temps perdu sont repérés et les données utilisables autorisées.", "Un cas d'usage borné, après contrôle de confidentialité."),
]
PROBLEMES = [
    ("L'infra est perçue comme un centre de coûts", "Clarifier la proposition de valeur", "Valeur délivrée par capacité"),
    ("Les utilisateurs ne comprennent pas l'offre", "Formaliser un catalogue de capacités", "Demandes couvertes par le catalogue"),
    ("Les demandes sont trop spécifiques", "Identifier les besoins récurrents", "Part de demandes standardisées"),
    ("Les outils sont peu adoptés", "Travailler l'UX et l'adoption", "Taux d'adoption"),
    ("Les équipes applicatives contournent l'infra", "Rendre l'offre simple et désirable", "Usages hors plateforme"),
    ("Les priorités sont politiques", "Prioriser par valeur, usage, risque, adoption", "Backlog priorisé sur critères partagés"),
    ("Les capacités livrées ne sont pas utilisées", "Mesurer l'usage réel", "Usage par capacité livrée"),
    ("L'automatisation manque de financement", "Construire un business case produit", "Capacité Run libérée, puis réallouée"),
    ("L'obsolescence est subie", "Piloter le cycle de vie des capacités", "Décommissionnements tenus"),
    ("La plateforme est trop technique", "Définir des parcours compréhensibles", "Délai de mise à disposition"),
]


def convictions(prs):
    s1, s2 = _scinder(prs, T11, {0, 1, 2, 6}, "Une conviction, un critère, un chantier : 1 à 3", "Une conviction, un critère, un chantier : 4 et 5")
    w, gap, y, h, haut = 195, 14.5, 118, 205, 60
    for s, lot in ((s1, range(0, 3)), (s2, range(3, 5))):
        for col, k in enumerate(lot):
            conviction, critere, chantier = CONVICTIONS[k]
            carte(s, 45 + col * (w + gap), y, w, h, haut, conviction, [("LE CRITÈRE OBSERVABLE", critere), ("LE CHANTIER À MENER", chantier)], tag=f"CONVICTION {k + 1}")
    x = 45 + 2 * (w + gap)
    panneau = deck_draw.rect(s2, x, y, w, h, NAVY, MSO_SHAPE.ROUNDED_RECTANGLE)
    panneau.adjustments[0] = 12 / min(w, h)
    panneau.name = "panneau-regle"
    etiquette = deck_draw.rect(s2, x + 10, y + 12, 112, 19, CYAN, MSO_SHAPE.ROUNDED_RECTANGLE)
    etiquette.adjustments[0] = 0.5
    p = _cadre(etiquette, MSO_ANCHOR.MIDDLE, (2, 0)).paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _run(p, "RÈGLE DE PASSAGE", 9, NAVY, True)
    texte = deck_draw.rect(s2, x + 10, y + 40, w - 20, h - 52, NAVY, MSO_SHAPE.RECTANGLE)
    texte.fill.background()
    tf = _cadre(texte, MSO_ANCHOR.TOP, (0, 0))
    _run(tf.paragraphs[0], "Une capacité entre dans la roadmap quand elle a un propriétaire, une valeur validée et un indicateur.", 10.5, BLANC, False)
    q = tf.add_paragraph()
    q.space_before = Pt(10)
    _run(q, "L'IA prépare, nos coachs challengent, le client décide.", 10.5, BLANC, True)


def problemes(prs):
    s1, s2 = _scinder(prs, T38, {0, 45}, "Dix problèmes, dix apports, dix indicateurs : 1 à 5", "Dix problèmes, dix apports, dix indicateurs : 6 à 10")
    w, gap, y, h, haut = 114, 11, 112, 225, 66
    for s, lot in ((s1, range(0, 5)), (s2, range(5, 10))):
        for col, k in enumerate(lot):
            probleme, apport, indicateur = PROBLEMES[k]
            carte(s, 45 + col * (w + gap), y, w, h, haut, probleme, [("L'APPORT", apport), ("L'INDICATEUR", indicateur)], tag=f"{k + 1} / 10", taille_titre=9.5)


def retouche(prs):
    convictions(prs)
    problemes(prs)
    navigation_8(prs)
    sections(prs)


def _ids():
    return [f"id:{s.slide_id}" for s in Presentation(deck_base.resoudre_base(BASE)).slides]


ETAPES = [{"declarer": _ids(), "f": retouche, "nouvelles_max": 2, "nom": "slides 11 et 38 scindees, cartes a en-tete"}]
