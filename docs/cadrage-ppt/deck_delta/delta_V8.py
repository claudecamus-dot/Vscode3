"""Delta V8 (base V7) : redessin de la slide des convictions (illisible), correction des pastilles de la boucle, pictogrammes.

Usage (depuis docs/cadrage-ppt/) : py deck_delta/deck.py delta_V8 --base V7 --dry-run   puis sans --dry-run.
Pictos : dessines en autoshapes (disque + glyphe simple), pas d'image : pas de dependance, couleur de charte uniquement.
"""
import deck_draw
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR

BASE = "V7"
NOTE = "V8 : slide Convictions redessinee (3 colonnes lisibles), pastilles de la boucle dans leur cadre, pictos sur Coach augmente et Agents"

NAVY = RGBColor(0x0E, 0x23, 0x56)
CYAN = RGBColor(0x00, 0xD2, 0xDD)
MUTED = RGBColor(0x58, 0x65, 0x86)
CLAIR = RGBColor(0xE7, 0xE9, 0xEE)
LIGNE = RGBColor(0xD5, 0xD8, 0xE0)
BLANC = RGBColor(0xFF, 0xFF, 0xFF)
EMU = 12700

T_CONV_ANCIEN = "Chaque conviction fonde un principe, un critère observable et un chantier"
T_CONV = "Chaque conviction se vérifie par un critère et devient un chantier"
T_BOUCLE = "Une boucle continue pour décider sur des faits, expérimenter et ancrer"
T_AGENTS = "Des agents augmentent le coach ; d'autres améliorent l'Infra du client"
T_COACH = "Le coach produit augmenté est le moteur de la transformation"
OPS = []

# (conviction, critere observable, chantier) : condensations du texte V7 (colonne « principe » fusionnee dans la conviction)
LIGNES = [
    ("Une plateforme n'est adoptée que si elle résout une douleur mesurée chez ceux qui la consomment.",
     "Les douleurs de chaque consommateur sont documentées, avec une mesure.",
     "Interviews et situation de référence mesurée."),
    ("Sans porteur produit, l'infra retourne au guichet.",
     "Chaque capacité clé a un propriétaire nommé, une roadmap et un indicateur.",
     "Désigner les propriétaires, bâtir la roadmap, définir l'indicateur de valeur."),
    ("Les priorités se décident sur des faits mesurés, pas sur le rapport de force.",
     "Une grille de priorisation partagée alimente les arbitrages budgétaires.",
     "Choisir les indicateurs, scorer et prioriser, ritualiser la décision."),
    ("Le progrès se mesure de façon régulière et transparente, à chaque niveau.",
     "Des indicateurs de progrès sont suivis à rythme régulier, succès comme erreurs.",
     "Installer la mesure et un rituel de revue qui rend les équipes redevables."),
    ("L'IA gagne du temps là où il est perdu : mesuré puis vérifié, pas promis.",
     "Les gisements de temps perdu sont repérés et les données utilisables autorisées.",
     "Un cas d'usage borné, après contrôle de confidentialité."),
]


def _texte(s, x, y, w, h, paras, ancre=MSO_ANCHOR.TOP):
    """deck_draw.texte + police du deck (Outfit) sur chaque run (sinon la police par defaut de PowerPoint s'applique)."""
    tb = deck_draw.texte(s, x, y, w, h, paras, ancre=ancre)
    for para in tb.text_frame.paragraphs:
        for run in para.runs:
            run.font.name = "Outfit"
    return tb


def _supprimer(sh):
    sh._element.getparent().remove(sh._element)


def redessiner_convictions(prs):
    s = deck_draw.slide_par_titre(prs, T_CONV_ANCIEN)
    deck_draw.set_run(s.shapes.title, T_CONV)
    garder = {0, 1, 2, 37, 38, 39, 40}  # titre, pastille « proposition », bandeau de la regle de passage, bandeau de chapitre
    for i, sh in enumerate(list(s.shapes)):
        if i not in garder:
            _supprimer(sh)
    x0, x_conv, x_crit, x_chant, fin = 45, 83, 330, 520, 659
    for x, w, libelle in ((x_conv, 235, "LA CONVICTION"), (x_crit, 180, "LE CRITÈRE OBSERVABLE"), (x_chant, 135, "LE CHANTIER À MENER")):
        _texte(s, x, 92, w, 12, [[(libelle, 8.5, True, MUTED)]])
    y, h, gap = 108, 40, 4.5
    for n, (conv, crit, chant) in enumerate(LIGNES, 1):
        accent = n == 1  # « un sur N en accent »
        carte = deck_draw.rect(s, x0, y, fin - x0, h, BLANC, MSO_SHAPE.ROUNDED_RECTANGLE, ligne=NAVY if accent else LIGNE)
        carte.adjustments[0] = 0.18
        _texte(s, x0 + 8, y, 28, h, [[(str(n), 22, True, RGBColor(0xC5, 0xCA, 0xD8))]], ancre=MSO_ANCHOR.MIDDLE)
        _texte(s, x_conv, y, 235, h, [[(conv, 9.5, True, NAVY)]], ancre=MSO_ANCHOR.MIDDLE)
        _texte(s, x_crit, y, 180, h, [[(crit, 9, False, NAVY)]], ancre=MSO_ANCHOR.MIDDLE)
        _texte(s, x_chant, y, 129, h, [[(chant, 9, False, NAVY)]], ancre=MSO_ANCHOR.MIDDLE)
        y += h + gap


def _forme_texte(s, debut):
    return next(sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.startswith(debut))


def pastilles_boucle(prs):
    """Les 3 pastilles « Actions possibles » dépassaient du cadre (Formation coupée) : trois largeurs qui tiennent dans le cadre."""
    s = deck_draw.slide_par_titre(prs, T_BOUCLE)
    cadre = _forme_texte(s, "Actions possibles")
    droite = cadre.left + cadre.width
    plan = (("Assainissement", 74), ("Évolution technique", 86), ("Formation", 54))
    x = cadre.left + 6 * EMU
    for nom, w in plan:
        txt = _forme_texte(s, nom)
        cx, cy = txt.left + txt.width // 2, txt.top + txt.height // 2
        fonds = [sh for sh in s.shapes if sh is not txt and sh.shape_type == 1 and sh.left <= cx <= sh.left + sh.width and sh.top <= cy <= sh.top + sh.height and sh.width < 120 * EMU]
        for sh in [txt, *fonds]:
            sh.left, sh.width = int(x), int(w * EMU)
        x += (w + 4) * EMU
    assert x - 4 * EMU <= droite, "les pastilles dépassent encore du cadre"


def picto(s, genre, cx, cy, d, fond, trait):
    """Disque de diametre d (pt) centre en (cx, cy) + glyphe simple en autoshapes ; genre : engrenage | bouclier | personne | cube | eclair."""
    disque = deck_draw.rect(s, cx - d / 2, cy - d / 2, d, d, fond, MSO_SHAPE.OVAL)
    disque.name = f"picto-{genre}"
    k = d / 24
    if genre == "engrenage":
        deck_draw.rect(s, cx - 8 * k, cy - 8 * k, 16 * k, 16 * k, trait, MSO_SHAPE.GEAR_6)
        deck_draw.rect(s, cx - 3 * k, cy - 3 * k, 6 * k, 6 * k, fond, MSO_SHAPE.OVAL)
    elif genre == "bouclier":
        deck_draw.rect(s, cx - 6 * k, cy - 7 * k, 12 * k, 14 * k, trait, MSO_SHAPE.FLOWCHART_OFFPAGE_CONNECTOR)
        deck_draw.rect(s, cx - 2.5 * k, cy - 2 * k, 5 * k, 3 * k, fond, MSO_SHAPE.RECTANGLE)
    elif genre == "personne":
        deck_draw.rect(s, cx - 3 * k, cy - 8 * k, 6 * k, 6 * k, trait, MSO_SHAPE.OVAL)
        deck_draw.rect(s, cx - 6 * k, cy - 1 * k, 12 * k, 8 * k, trait, MSO_SHAPE.ROUND_2_SAME_RECTANGLE)
    elif genre == "cube":
        deck_draw.rect(s, cx - 7 * k, cy - 7 * k, 14 * k, 14 * k, trait, MSO_SHAPE.CUBE)
    elif genre == "eclair":
        deck_draw.rect(s, cx - 6 * k, cy - 8 * k, 12 * k, 16 * k, trait, MSO_SHAPE.LIGHTNING_BOLT)
    return disque


def pictos_agents(prs):
    s = deck_draw.slide_par_titre(prs, T_AGENTS)
    for debut, genre in (("Le coach humain valide", "bouclier"), ("Les équipes opérationnelles", "personne")):
        txt = _forme_texte(s, debut)
        cy = (txt.top + txt.height / 2) / EMU
        picto(s, genre, txt.left / EMU - 16, cy, 22, NAVY, BLANC)


def pictos_coach(prs):
    s = deck_draw.slide_par_titre(prs, T_COACH)
    for debut, genre, fond, trait in (("COACH", "personne", CLAIR, NAVY), ("AGENTS DU COACH", "engrenage", NAVY, BLANC), ("CLIENT", "cube", NAVY, BLANC)):
        pastille = next(sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip() == debut)
        cy = (pastille.top + pastille.height / 2) / EMU
        picto(s, genre, pastille.left / EMU - 5, cy, 18, fond, trait)


ETAPES = [
    {"declarer": [T_CONV_ANCIEN], "f": redessiner_convictions, "nouvelles_max": 0, "nom": "convictions redessinees"},
    {"declarer": [T_BOUCLE], "f": pastilles_boucle, "nouvelles_max": 0, "nom": "pastilles de la boucle"},
    {"declarer": [T_AGENTS], "f": pictos_agents, "nouvelles_max": 0, "nom": "pictos agents"},
    {"declarer": [T_COACH], "f": pictos_coach, "nouvelles_max": 0, "nom": "pictos coach"},
]
