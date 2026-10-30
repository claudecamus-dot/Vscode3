"""Delta V20 (base V19) : premiere version du plan du comite (T3_CONSOLIDATION, 2026-10-30) : ouverture resserree, bandeau numerote, sections sans intercalaires.

- ouverture : les trois slides qui disaient trois fois la meme chose (lettre « quatre atouts », « cinq points », « votre demande, notre reponse en trois temps, nos engagements »)
  deviennent UNE slide : la troisieme, la plus complete (demande, trois temps, engagements), retitree en assertion. Reprise : les quatre atouts de la lettre sont deja
  dans ses engagements (aucun ROI promis, l'IA prepare / le client decide, relais) ; le constat, la cible et le pourquoi maintenant sont repris par « Contexte et enjeux ».
- intercalaires : les sept intercalaires de chapitre disparaissent (decision D1 du comite, defaut « non » : le bandeau de navigation, le sommaire et les sections portent le
  chapitre) ; celui des Annexes reste. Le sommaire garde ses sept pastilles.
- bandeau : le surtitre de chaque slide de chapitre porte le numero du chapitre du sommaire (« 03 · NOTRE APPROCHE · 2/4 ») ; le chapitre se lit sur la pastille navy de la navigation.
- sections PowerPoint : une par chapitre, deduites de la pastille navy (plus des intercalaires, qui n'existent plus) ; ouverture = couverture, slide d'ouverture, sommaire.
Les fonctions de chapitre de delta_V9 / delta_V11 (`_chapitres`, `navigation_8`) lisent les intercalaires : elles ne s'appliquent plus a partir de la V20 (ce delta n'importe plus delta_V9 : son import resout la base V8, archivee depuis).
Usage (depuis docs/cadrage-ppt/) : py deck_delta/deck.py delta_V20 --base V19
"""
import re
import uuid

import deck_base
import deck_draw
from lxml import etree
from pptx import Presentation

BASE = "V19"
NAVY = "0E2356"
EMU = 12700
NOTE = "V20 : ouverture en une slide, intercalaires retires, bandeau numerote, sections par chapitre"

T_OUVERTURE = "Votre demande, notre réponse en trois temps, nos engagements"
T_OUVERTURE_NEUF = "En trois temps et trois engagements, votre Infra passe du guichet au produit"
NUMEROS = {"Contexte et enjeux": "01", "Nos convictions": "02", "Notre approche": "03", "Démarche générale": "04",
           "Démarche détaillée": "05", "Dispositif et prix": "06", "Prochain pas": "07", "Annexes": "A"}
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
EXT_SECTIONS = "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"


def _base():
    prs = Presentation(deck_base.resoudre_base(BASE))
    sorties = []
    for s in prs.slides:
        if _est_intercalaire(s) and _nom_chapitre(s) != "Annexes":
            sorties.append(s.slide_id)
        elif s.shapes.title is not None and s.shapes.title.text_frame.text in ("Notre engagement : quatre atouts pour passer du guichet au produit",
                                                                               "Cinq points pour passer d'un guichet saturé à une Infra créatrice de valeur"):
            sorties.append(s.slide_id)
    return [f"id:{s.slide_id}" for s in prs.slides if s.slide_id not in sorties], sorties


def _est_intercalaire(s):
    return s.slide_layout.name.startswith(("50 -", "51 -"))


def _nom_chapitre(s):
    return next((sh.text_frame.text for sh in s.shapes if sh.is_placeholder and sh.has_text_frame and sh.top > 100 * EMU
                 and not re.fullmatch(r"\d{2}|[A-Z]", sh.text_frame.text)), None)


def _kicker(s):
    return next((sh for sh in s.shapes if sh.has_text_frame and not sh.name.startswith("nav-")
                 and re.fullmatch(r"[^\d]+ · \d+/\d+", sh.text_frame.text)), None)


_RESTANTES, _SUPPRIMEES = _base()
OPS = [{"slide": f"id:{i}", "op": "supprimer"} for i in _SUPPRIMEES]


def chapitre_de(s):
    """Chapitre d'une slide = nom de la pastille de navigation en aplat navy (None : couverture, ouverture, sommaire, intercalaire)."""
    for sh in s.shapes:
        if sh.name.startswith("nav-") and sh.fill.type == 1 and str(sh.fill.fore_color.rgb) == NAVY:
            return sh.name[4:]
    return None


def ouverture(prs):
    s = deck_draw.slide_par_titre(prs, T_OUVERTURE)
    deck_draw.set_run(s.shapes.title, T_OUVERTURE_NEUF)
    k = next(sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.startswith("OUVERTURE"))
    deck_draw.set_run(k, "OUVERTURE")


def bandeau(prs):
    """Surtitre « NN · CHAPITRE · n/N » recompte sur les slides de chaque chapitre (ordre du deck)."""
    par_chapitre = {}
    for s in prs.slides:
        ch = chapitre_de(s)
        if ch:
            par_chapitre.setdefault(ch, []).append(s)
    for ch, slides in par_chapitre.items():
        for i, s in enumerate(slides, 1):
            k = next((sh for sh in s.shapes if sh.has_text_frame and not sh.name.startswith("nav-")
                      and re.fullmatch(r"(\d{2}|A · )?[^\d]+ · \d+/\d+|\d{2} · [^\d]+ · \d+/\d+", sh.text_frame.text)), None)
            if k is None:
                k = _kicker(s)
            nouveau = f"{NUMEROS[ch]} · {ch.upper()} · {i}/{len(slides)}"
            if k is not None and k.text_frame.text != nouveau:
                deck_draw.set_run(k, nouveau)


def sections(prs):
    """Une section par chapitre : ouverture (jusqu'au sommaire), un chapitre par pastille navy, annexes (intercalaire compris)."""
    groupes = []
    for s in prs.slides:
        nom = chapitre_de(s) or ("Annexes" if _est_intercalaire(s) else None)
        if nom is None:
            nom = groupes[-1][0] if groupes and groupes[-1][0] == "Ouverture" else ("Ouverture" if not groupes else groupes[-1][0])
        if groupes and groupes[-1][0] == nom:
            groupes[-1][1].append(deck_draw.id_de(prs, s))
        else:
            groupes.append((nom, [deck_draw.id_de(prs, s)]))
    pres = prs.part._element
    ext_lst = pres.find(f"{P}extLst")
    if ext_lst is None:
        ext_lst = etree.SubElement(pres, f"{P}extLst")
    for ext in ext_lst.findall(f"{P}ext"):
        if ext.get("uri") == EXT_SECTIONS.strip("{}") or ext.get("uri") == EXT_SECTIONS:
            ext_lst.remove(ext)
    ext = etree.SubElement(ext_lst, f"{P}ext", uri=EXT_SECTIONS)
    liste = etree.SubElement(ext, f"{{{deck_draw.P14}}}sectionLst", nsmap={"p14": deck_draw.P14})
    for nom, ids in groupes:
        sec = etree.SubElement(liste, f"{{{deck_draw.P14}}}section", name=nom, id="{" + str(uuid.uuid4()).upper() + "}")
        sl = etree.SubElement(sec, f"{{{deck_draw.P14}}}sldIdLst")
        for i in ids:
            etree.SubElement(sl, f"{{{deck_draw.P14}}}sldId", id=i)


def retouche(prs):
    ouverture(prs)
    bandeau(prs)
    sections(prs)


ETAPES = [{"declarer": _RESTANTES, "f": retouche, "nouvelles_max": 0, "nom": "ouverture, bandeau numerote, sections"}]
