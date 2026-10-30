"""Delta V26 (base V25) : demande utilisateur du 2026-10-30 — « ok pour les 6 convictions, mais dans 1 seul chapitre, lance la refonte de forme »
(styles choisis : A cartes goutte, D processus en chevrons, E tableaux en bandes).

Fond : six convictions ajoutees au chapitre « Nos convictions » (organisation, equipe, pilotage, demande, RUN : service vecu, RUN pret), issues du comite de lecture de
Telechargements (comite_conv/C1..C4, comite_pdf/P1..P9). Aucun chiffre : les criteres sont observables, les valeurs de reference sont a mesurer chez le client.
Forme : slides 10 et 11 (convictions 1 a 5) en cartes goutte ; « Deux temps d'engagement » en processus ; « L'IA est gouvernee » en cartes + capsules ; « Chiffrage » en tableau a bande.
Usage (depuis docs/cadrage-ppt/) : py deck_delta/deck.py delta_V26 --base V25
"""
import deck_base
import deck_draw
from delta_V20 import bandeau
from delta_V21 import sections
from forme_v26 import *  # noqa: F401,F403
from pptx import Presentation

BASE = "V25"
NOTE = "V26 : six convictions (organisation, equipe, pilotage, demande, RUN) dans le chapitre Nos convictions ; refonte de forme des slides cles"
OPS = []

T_MODELE = "Deux convictions de plus, vérifiées de la même façon"
T10 = "Trois convictions, chacune vérifiée par un critère et un chantier"
T18 = "Deux temps d'engagement : diagnostiquer, puis expérimenter et déployer"
T27 = "L'IA est gouvernée : données, droits, validations et traçabilité"
T31 = "Chiffrage : la forme du tableau jours × rôles, toutes valeurs à valider"


def _ids():
    return [f"id:{s.slide_id}" for s in Presentation(deck_base.resoudre_base(BASE)).slides]


def T(sh):
    return sh.text_frame.text


def paras(sh):
    return [p.text for p in sh.text_frame.paragraphs]


def par_nom(s, nom):
    return [sh for sh in s.shapes if sh.name == nom]


# ---------------------------------------------------------------- convictions 1 a 5 (cartes goutte)
def convictions_cartes(prs):
    s = deck_draw.slide_par_titre(prs, T10)
    heads, pills, bodies = (par_nom(s, n) for n in ("entete-carte", "pastille-rubrique", "texte-carte"))
    data = [(T(h), T(p), paras(b)) for h, p, b in zip(heads, pills, bodies)]
    vider(s)
    signature_cyan(s)
    cartes_conviction(s, data, [X0 + k * 212 for k in range(3)])

    s = deck_draw.slide_par_titre(prs, "Deux convictions de plus, vérifiées de la même façon")
    heads, pills, bodies = (par_nom(s, n) for n in ("entete-carte", "pastille-rubrique", "texte-carte"))
    data = [(T(h), T(p), paras(b)) for h, p, b in zip(heads, pills, bodies)]
    titre_regle, texte_regle = T(par_nom(s, "Rounded Rectangle 87")[0]), T(par_nom(s, "Rectangle 88")[0])
    vider(s)
    signature_cyan(s)
    cartes_conviction(s, data, [X0, X0 + 212])
    x = X0 + 424
    goutte(s, x, 98, 190, 266, NAVY, 30)
    etiquette(s, x + 16, 98 + 22, 120, titre_regle)
    tx(s, x + 16, 98 + 62, 158, 190, [[(texte_regle, 12, False, BLANC)]])


def cartes_conviction(s, data, xs):
    W, Y = 190, 98
    for (head, pill, body), x in zip(data, xs):
        goutte(s, x, Y, W, 266, CLAIR, 30)
        goutte(s, x, Y, W, 96, NAVY, 30)
        tx(s, x + 16, Y + 8, W - 30, 80, [[(head, 13, True, BLANC)]], A.MIDDLE)
        etiquette(s, x + 16, Y + 96 - 11, 104, pill)
        tx(s, x + 16, Y + 120, W - 32, 14, [[(body[0], 10.5, True, MUTED)]])
        tx(s, x + 16, Y + 137, W - 32, 56, [[(body[1], 11.5, False, NAVY)]])
        tx(s, x + 16, Y + 198, W - 32, 14, [[(body[2], 10.5, True, MUTED)]])
        tx(s, x + 16, Y + 215, W - 32, 46, [[(body[3], 11.5, False, NAVY)]])


# ---------------------------------------------------------------- six convictions nouvelles
NOUVELLES = [
    ("Conviction 6 : financer des produits et des équipes stables",
     "Part du budget infra rattachée à des équipes ou produits nommés, sur plusieurs trimestres.",
     "Définir le produit infra, nommer son responsable, passer d'un budget de projet à une capacité annuelle."),
    ("Conviction 7 : une petite équipe produit qui construit et exploite",
     "Temps de construction protégé ; demandes récurrentes passées en self-service.",
     "Cadrer la composition de l'équipe, le binôme produit et la répartition entre construire et exploiter."),
    ("Conviction 8 : trois instances, chacune avec sa décision",
     "Chaque instance a une carte d'identité écrite ; chaque revue trimestrielle produit une décision tracée.",
     "Atelier des cartes d'identité, seuils de délégation, calendrier annuel, premier cycle sur 2 ou 3 capacités."),
    ("Conviction 9 : un flux d'entrée unique qui qualifie avant de servir",
     "Toute demande a une fiche et un statut visibles ; le délai de qualification est mesuré.",
     "Cartographier les flux entrants, fiche de demande standard, grille de priorisation partagée."),
    ("Conviction 10 : piloter le service vécu, pas seulement le SLA",
     "Écart entre SLA et expérience mesuré au T0 puis au T1 ; part de RUN subi suivie.",
     "Objectifs de service par capacité, revue trimestrielle des écarts, charge de RUN rendue visible."),
    ("Conviction 11 : aucun produit livré sans RUN prêt",
     "Contrôle « RUN prêt » validé avant chaque mise en service ; deux personnes capables par module.",
     "Contrôle de préparation au RUN sur un pilote, post-mortems reliés au backlog, IA d'aide au diagnostic."),
]


def f_organisation(s):
    trois_cartes(s, [
        ("AUJOURD'HUI", "Un budget par projet", ["L'équipe est dissoute à la fin du projet.", "Le RUN est transféré à d'autres équipes."]),
        ("CIBLE", "Une capacité par produit", ["Une équipe stable, qui construit et exploite.", "Un budget annuel rattaché au produit."]),
        ("LE TEST", "Produit ou composant ?", ["Produits : Kubernetes, CI/CD, IAM.", "Composants : VM, serveur, SAN.", "Puis-je décrire l'offre, la valeur, les engagements ?"]),
    ])


def f_equipe(s):
    trois_cartes(s, [
        ("LE PRINCIPE", "Construire et exploiter", ["Une même équipe conçoit, livre et assure le service.", "Du temps de construction est protégé."]),
        ("LE RÔLE", "Un binôme porte le produit", ["Un responsable produit et un référent technique.", "Pas un PO posé dans l'Infra."]),
        ("LE RÉSULTAT", "Moins de tickets", ["Les demandes récurrentes sont automatisées.", "Elles sont rendues en libre-service."]),
    ])


def f_pilotage(s):
    cx, cw = [113 + k * 182 for k in range(3)], 172
    noms = ["Comité d'investissement", "Revue produit", "Suivi d'exécution"]
    for k, nom in enumerate(noms):
        chevron(s, cx[k] - 6, 100, 188, 34, NAVY, premier=(k == 0))
        tx(s, cx[k] + 6, 100, 160, 34, [[(nom, 11.5, True, BLANC)]], A.MIDDLE, P.CENTER)
    lignes = [
        ("Décide", ["arrêter, étendre ou pivoter une capacité", "backlog, roadmap et budget du produit", "ajustements, risques, escalades"]),
        ("Avec qui", ["sponsor, DSI, responsables produit", "responsable produit, métiers, équipe", "équipe, coach, relais client"]),
        ("Cadence", ["trimestrielle, à valider", "à fixer avec vous", "à fixer avec vous"]),
    ]
    for r, (lab, cells) in enumerate(lignes):
        y = 142 + r * 36
        if r % 2 == 0:
            forme(s, S.RECTANGLE, X0, y, 614, 34, CLAIR, None)
        tx(s, X0 + 8, y, 58, 34, [[(lab, 10.5, True, MUTED)]], A.MIDDLE)
        for k, c in enumerate(cells):
            tx(s, cx[k], y, cw, 34, [[(c, 10.5, k == 0 and r == 0, NAVY)]], A.MIDDLE)
    tx(s, X0, 254, 614, 14, [[("Les deux comités du pilote, suivi et décision, en sont la première étape.", 10.5, False, MUTED)]])


def f_demande(s):
    W, G, y, h = 140, 18, 104, 94
    cartes = [("Recevoir", "Un seul point d'entrée ; une fiche : problème, valeur, indicateur."),
              ("Qualifier", "Valeur, faisabilité et charge estimées avant d'engager l'équipe."),
              ("Prioriser", "Une grille explicite ; au-delà d'un seuil de charge, un ROI est demandé."),
              ("Servir", "Statut visible ; les demandes récurrentes passent en self-service.")]
    for k, (t, c) in enumerate(cartes):
        x = X0 + k * (W + G)
        sombre = k == 3
        col = BLANC if sombre else NAVY
        goutte(s, x, y, W, h, NAVY if sombre else CLAIR, 22)
        pastille(s, x + 22, y, 26, k + 1, sombre=not sombre)
        tx(s, x + 38, y + 8, W - 46, 18, [[(t, 12.5, True, col)]], A.MIDDLE)
        tx(s, x + 12, y + 32, W - 22, h - 36, [[(c, 10.5, False, col)]])
    tx(s, X0, 205, 614, 14, [[("LA CAPACITÉ SE RÉPARTIT PAR DÉCISION, PAS PAR ORDRE D'ARRIVÉE", 10.5, True, MUTED)]])
    for lab, x, w, sombre in (("RUN", X0, 250, True), ("Nouveaux sujets", X0 + 254, 220, False), ("Amélioration", X0 + 478, 136, False)):
        capsule(s, x, 222, w, 24, NAVY if sombre else CLAIR)
        tx(s, x, 222, w, 24, [[(lab, 10.5, True, BLANC if sombre else NAVY)]], A.MIDDLE, P.CENTER)
    tx(s, X0, 250, 614, 12, [[("Répartition illustrative, à calibrer produit par produit et à revoir chaque trimestre.", 10, False, MUTED)]])


def f_service(s):
    CX, CW = [57, 252, 459], 185
    capsule(s, X0, 100, 614, 24, NAVY)
    for j, h_ in enumerate(("Ce que mesure le SLA", "Ce que vit le consommateur", "Ce que nous suivons, du T0 au T1")):
        tx(s, CX[j], 100, CW, 24, [[(h_, 10.5, True, BLANC)]], A.MIDDLE)
    lignes = [("Disponibilité du service", "Une demande qui attend, même quand le service est « vert »", "Objectif de service publié par capacité"),
              ("Incidents clos dans le délai", "Le même incident qui revient", "Taux d'incidents récurrents"),
              ("Tickets traités", "Du temps passé à subir plutôt qu'à améliorer", "Part de RUN subi et maîtrisé")]
    for r, ligne in enumerate(lignes):
        y = 130 + r * 44
        if r % 2 == 0:
            forme(s, S.RECTANGLE, X0, y, 614, 42, CLAIR, None)
        for j, c in enumerate(ligne):
            tx(s, CX[j], y, CW, 42, [[(c, 11, j == 2, NAVY)]], A.MIDDLE)


def f_run_pret(s):
    xs = [X0 + k * 150 for k in range(4)]
    for k, nom in enumerate(["Construire", "Porte « RUN prêt »", "Exploiter", "Apprendre"]):
        chevron(s, xs[k], 100, 164, 36, NAVY if k == 1 else CLAIR, premier=(k == 0))
        tx(s, xs[k] + (6 if k == 0 else 20), 100, 128, 36, [[(nom, 11, True, BLANC if k == 1 else NAVY)]], A.MIDDLE, P.CENTER)
    goutte(s, X0, 152, 300, 112, NAVY, 24)
    tx(s, X0 + 16, 162, 270, 16, [[("LA PORTE « RUN PRÊT »", 10.5, True, BLANC)]])
    tx(s, X0 + 16, 182, 270, 78, [[("• " + t, 10.5, False, BLANC)] for t in (
        "Deux personnes capables par module", "Supervision et astreinte définies", "Budget RUN dans le business case", "Procédure d'incident documentée")], espace=3)
    goutte(s, X0 + 314, 152, 300, 112, CLAIR, 24)
    tx(s, X0 + 330, 162, 270, 16, [[("LA BOUCLE D'APPRENTISSAGE", 10.5, True, MUTED)]])
    tx(s, X0 + 330, 182, 270, 78, [[("Chaque incident majeur : un post-mortem, puis une action au backlog.", 11, False, NAVY)],
                                  [("L'IA aide à diagnostiquer ; la décision et la mise en œuvre restent humaines et tracées.", 11, False, NAVY)]], espace=6)


FORMES = [f_organisation, f_equipe, f_pilotage, f_demande, f_service, f_run_pret]


def nouvelles_convictions(prs):
    ref = deck_draw.slide_par_titre(prs, T_MODELE)
    modele_rang = list(prs.slides).index(ref)
    for k, ((titre, critere, chantier), f) in enumerate(zip(NOUVELLES, FORMES)):
        neuve = deck_base.core().copier_slide(prs, ref, index=modele_rang + 1 + k)
        vider(neuve)
        deck_draw.set_run(neuve.shapes.title, titre)
        signature_cyan(neuve)
        f(neuve)
        bas_de_conviction(neuve, critere, chantier)


# ---------------------------------------------------------------- slides 18, 27, 31 : refonte
def deux_temps(prs):
    s = deck_draw.slide_par_titre(prs, T18)
    sh = list(s.shapes)
    bandes = [T(sh[2]), T(sh[4])]
    cols = [(T(sh[7 + 10 * k]), T(sh[9 + 10 * k]), T(sh[10 + 10 * k]), T(sh[11 + 10 * k]), T(sh[14 + 10 * k] if k < 2 else sh[33]), T(sh[35 + 2 * k])) for k in range(3)]
    note = T(sh[40])
    vider(s)
    signature_cyan(s)
    W, G = 190, 22
    capsule(s, X0, 96, W, 22, NAVY)
    tx(s, X0, 96, W, 22, [[(bandes[0], 10.5, True, BLANC)]], A.MIDDLE, P.CENTER)
    capsule(s, X0 + W + G, 96, 2 * W + G, 22, CLAIR)
    tx(s, X0 + W + G, 96, 2 * W + G, 22, [[(bandes[1], 10.5, True, NAVY)]], A.MIDDLE, P.CENTER)
    for k, (ti, nu, co, cd, du, mo) in enumerate(cols):
        x, y, sombre = X0 + k * (W + G), 140, k == 0
        fond, enc = (NAVY, BLANC) if sombre else (CLAIR, NAVY)
        goutte(s, x, y, W, 214, fond, 30)
        pastille(s, x + 38, y, 40, nu, sombre=not sombre)
        tx(s, x + 14, y + 25, W - 28, 22, [[(ti, 15, True, enc)]])
        tx(s, x + 14, y + 46, W - 28, 86, [[(co, 10.5, False, enc)]])
        tx(s, x + 14, y + 133, W - 28, 27, [[(cd, 10.5, True, enc)]])
        capsule(s, x + 14, y + 161, W - 28, 20, BLANC if sombre else NAVY)
        tx(s, x + 14, y + 161, W - 28, 20, [[(du, 10.5, True, NAVY if sombre else BLANC)]], A.MIDDLE, P.CENTER)
        forme(s, S.ROUNDED_RECTANGLE, x + 14, y + 184, W - 28, 26, None, BLANC if sombre else NAVY, adj=0.3)
        tx(s, x + 20, y + 184, W - 40, 26, [[(mo, 10.5, True, enc)]], A.MIDDLE, P.CENTER)
    tx(s, X0, 358, 614, 14, [[(note, 11, False, MUTED)]])


def chiffrage(prs):
    s = deck_draw.slide_par_titre(prs, T31)
    sh = list(s.shapes)
    hdr = [T(sh[i]) for i in (4, 5, 6, 7)]
    rows = [[T(sh[b + j]) for j in range(4)] for b in (10, 15, 20, 25)]
    note, cond_t = T(sh[30]), T(sh[32])
    cond = [[T(sh[b + j]) for j in range(3)] for b in (34, 38, 42, 46)]
    vider(s)
    signature_cyan(s)
    CX, CW = [62, 296, 396, 512], [220, 90, 100, 140]
    capsule(s, X0, 98, 614, 24, NAVY)
    for j in range(4):
        tx(s, CX[j], 98, CW[j], 24, [[(hdr[j], 10.5, True, BLANC)]], A.MIDDLE)
    for r in range(3):
        y = 128 + r * 30
        for j in range(4):
            tx(s, CX[j], y, CW[j], 28, [[(rows[r][j], 12 if j == 0 else 14, True, NAVY)]], A.MIDDLE)
        forme(s, S.RECTANGLE, X0 + 6, y + 29, 602, 1, CLAIR, None)
    goutte(s, X0, 222, 614, 32, NAVY, 14)
    for j in range(4):
        tx(s, CX[j], 222, CW[j], 32, [[(rows[3][j], 12, True, BLANC)]], A.MIDDLE)
    tx(s, X0, 261, 614, 15, [[(note, 11, False, MUTED)]])
    goutte(s, X0, 290, 614, 74, CLAIR, 24)
    etiquette(s, X0 + 16, 290 - 11, 196, cond_t)
    for k, (lab, val, st) in enumerate(cond):
        x = X0 + 18 + k * 148
        tx(s, x, 309, 138, 14, [[(lab, 11.5, True, MUTED)]])
        tx(s, x, 324, 138, 20, [[(val, 16, True, NAVY)]])
        tx(s, x, 343, 138, 14, [[(st, 11, False, MUTED)]])


def runs_of(sh):
    return [(r.text, bool(r.font.bold)) for r in sh.text_frame.paragraphs[0].runs]


def ia_gouvernee(prs):
    s = deck_draw.slide_par_titre(prs, T27)
    sh = list(s.shapes)
    tagl, gard = T(sh[2]), T(sh[4])
    steps = [(T(sh[7 + 5 * k]), paras(sh[8 + 5 * k])) for k in range(4)]
    face2, sub, leg, ag, hu = T(sh[24]), T(sh[25]), T(sh[26]), T(sh[27]), T(sh[28])
    raci = []
    for b in (30, 34, 38, 42):
        raci.append((T(sh[b]), [str(o.fill.fore_color.rgb) == "0E2356" for o in (sh[b + 1], sh[b + 2])]))
    items = [runs_of(sh[i]) for i in (47, 49, 51)]
    auto_l, chips = T(sh[52]), [T(sh[i]) for i in (54, 56, 58, 60)]
    vider(s)
    forme(s, S.ROUNDED_RECTANGLE, X0, 92, 5, 17, CYAN, adj=0.5)
    tx(s, X0 + 14, 92, 600, 17, [[(tagl, 13, True, NAVY)]], A.MIDDLE)
    capsule(s, X0, 113, 614, 22, NAVY)
    tx(s, X0 + 14, 113, 590, 22, [[(gard, 10.5, True, BLANC)]], A.MIDDLE)
    W, G = 146, 10
    for k, (ti, bl) in enumerate(steps):
        x = X0 + k * (W + G)
        goutte(s, x, 148, W, 74, CLAIR, 20)
        pastille(s, x + 22, 148, 26, k + 1, True)
        tx(s, x + 40, 152, W - 46, 16, [[(bl[0], 12.5, True, NAVY)]])
        tx(s, x + 14, 171, W - 22, 50, [[(b, 10.5, False, NAVY)] for b in bl[1:]])
    tx(s, X0, 228, 614, 13, [[(face2, 10.5, True, MUTED)]])
    tx(s, X0, 243, 292, 30, [[(sub, 10.5, True, NAVY)]])
    tx(s, X0, 273, 150, 13, [[(leg, 10.5, False, MUTED)]])
    tx(s, 192, 274, 58, 13, [[(ag, 10.5, True, NAVY)]], align=P.CENTER)
    tx(s, 261, 274, 58, 13, [[(hu, 10.5, True, NAVY)]], align=P.CENTER)
    for r, (lab, fl) in enumerate(raci):
        y = 288 + r * 14
        forme(s, S.RECTANGLE, X0, y, 288, 1, CLAIR, None)
        tx(s, X0, y + 1, 158, 14, [[(lab, 11, False, NAVY)]], A.MIDDLE)
        for cx_, f in zip((221, 290), fl):
            rond(s, cx_, y + 8, 10, NAVY if f else BLANC, NAVY, 1.25)
    for k, it in enumerate(items):
        y = 244 + k * 33
        goutte(s, 351, y, 308, 30, CLAIR, 12)
        tx(s, 363, y, 288, 30, [[(t, 10.5, b, NAVY) for t, b in it]], A.MIDDLE)
    tx(s, X0, 349, 116, 18, [[(auto_l, 10.5, True, MUTED)]], A.MIDDLE)
    for k, c in enumerate(chips):
        x = 170 + k * 122
        capsule(s, x, 348, 118, 19, NAVY if k == 3 else CLAIR)
        tx(s, x, 348, 118, 19, [[(c, 10.5, True, BLANC if k == 3 else NAVY)]], A.MIDDLE, P.CENTER)


def retouche(prs):
    convictions_cartes(prs)
    deux_temps(prs)
    ia_gouvernee(prs)
    chiffrage(prs)
    nouvelles_convictions(prs)
    bandeau(prs)
    sections(prs)


ETAPES = [{"declarer": _ids(), "f": retouche, "nouvelles_max": 6, "nom": "six convictions et refonte de forme (A, D, E)"}]
