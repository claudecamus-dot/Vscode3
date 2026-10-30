"""Constructeurs de formes et de slides pour les ETAPES de dessin d'un delta (port reduit de Vscode8 `scripts/deck_draw.py`).

Portes : id_de, deplacer_apres, nouvelle_apres, clone_forme, set_run, rect, texte (tels quels, sections p14 devenues OPTIONNELLES : le deck V4 n'en a pas,
la source levait AssertionError). Nouveaux : `slide_par_titre` (titre EXACT, meme definition que deck_rapport.titre_slide) et `ajouter_sur_gabarit`
(copie d'une slide modele via `copier_slide` du core, puis textes remplaces).
Non portes : la palette violette Accenture, `puce_ronde`, `liste_puces`, `cadre_chapitre` (formes 2 a 6 et pied « BPM017346 » propres au deck Accenture).
Ce module ne connait ni versions, ni gardes : il dessine. Les gardes sont appliquees autour par deck_garde.appliquer_etape.
"""
import copy
import re

from deck_lib import plat
from deck_rapport import titre_slide
from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.packuri import PackURI
from pptx.util import Emu, Pt

RID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
pt = lambda v: Emu(int(round(v * 12700)))  # noqa: E731


def _sections(prs):
    return prs.part._element.find(f".//{{{P14}}}sectionLst")


def id_de(prs, slide):
    """sldId (chaine) de la slide."""
    return [e for e in prs.slides._sldIdLst if prs.part.related_part(e.get(RID)) is slide.part][0].get("id")


def _retirer_des_sections(ext, sid):
    for sec in ext:
        for x in list(sec.find(f"{{{P14}}}sldIdLst")):
            if x.get("id") == sid:
                x.getparent().remove(x)


def _placer_dans_section(ext, ref_id, sid):
    for sec in ext:
        for x in sec.find(f"{{{P14}}}sldIdLst"):
            if x.get("id") == ref_id:
                x.addnext(etree.Element(f"{{{P14}}}sldId", id=sid))
                return
    raise AssertionError("section de la reference introuvable")


def deplacer_apres(prs, slide_a_deplacer, reference):
    """Deplace une slide juste apres la reference : plan de slides ET, s'il y en a, liste de sa section."""
    lst, ext = prs.slides._sldIdLst, _sections(prs)
    sid, rid = id_de(prs, slide_a_deplacer), id_de(prs, reference)
    e = [x for x in lst if x.get("id") == sid][0]
    lst.remove(e)
    [x for x in lst if x.get("id") == rid][0].addnext(e)
    if ext is not None:
        _retirer_des_sections(ext, sid)
        _placer_dans_section(ext, rid, sid)


def nouvelle_apres(prs, reference, layout, masquee=False):
    """Slide vide (placeholders du layout retires) inseree juste apres `reference`, dans le plan ET, s'il y en a, dans la section de la reference."""
    lst, ext = prs.slides._sldIdLst, _sections(prs)
    new = prs.slides.add_slide(layout)
    n = max(int(re.search(r"(\d+)", str(s.part.partname)).group(1)) for s in prs.slides) + 100
    new.part.partname = PackURI(f"/ppt/slides/slide{n}.xml")
    if masquee:
        new._element.set("show", "0")
    arbre = new.shapes._spTree
    for el in list(arbre)[2:]:
        arbre.remove(el)
    e = list(lst)[-1]
    lst.remove(e)
    ref_id = id_de(prs, reference)
    [x for x in lst if x.get("id") == ref_id][0].addnext(e)
    if ext is not None:
        _placer_dans_section(ext, ref_id, e.get("id"))
    return new


def clone_forme(src_shape, dst_slide):
    """Copie profonde d'une forme sans relation (une image ou un media exige `copier_slide` du core)."""
    el = copy.deepcopy(src_shape._element)
    dst_slide.shapes._spTree.append(el)
    return dst_slide.shapes[-1]


def set_run(shape, texte_):
    """Remplace le texte d'une forme par `texte_` en gardant la mise en forme du premier run ; supprime les autres runs et paragraphes."""
    p = shape.text_frame.paragraphs[0]
    runs = p.runs
    runs[0].text = texte_
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    for extra in shape.text_frame.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def rect(s, x, y, w, h, fill, forme=MSO_SHAPE.RECTANGLE, ligne=None):
    """Forme pleine ; x, y, w, h en points ; `fill` un RGBColor fourni par l'appelant (aucune palette ici)."""
    r = s.shapes.add_shape(forme, pt(x), pt(y), pt(w), pt(h))
    r.fill.solid()
    r.fill.fore_color.rgb = fill
    if ligne is None:
        r.line.fill.background()
    else:
        r.line.color.rgb = ligne
        r.line.width = Pt(0.75)
    r.shadow.inherit = False
    return r


def texte(s, x, y, w, h, paras_, ancre=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
    """Zone de texte sans marges ; `paras_` : [[(texte, taille, gras, RGBColor), ...], ...] (un paragraphe par liste de runs)."""
    tb = s.shapes.add_textbox(pt(x), pt(y), pt(w), pt(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = ancre
    for i, runs in enumerate(paras_):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for t, taille, gras, coul in runs:
            r = p.add_run()
            r.text = t
            r.font.size, r.font.bold = Pt(taille), gras
            r.font.color.rgb = coul
    return tb


def slide_par_titre(prs, titre):
    """LA slide dont le titre EXACT (casse et espaces normalises, meme definition que deck_rapport.titre_slide : placeholder titre, sinon premiere forme texte) est `titre`.
    ValueError si zero ou plusieurs : aucun prefixe n'est accepte."""
    cible = plat(titre).casefold()
    trouvees = [sl for sl in prs.slides if titre_slide(sl).casefold() == cible]
    if len(trouvees) != 1:
        raise ValueError(f"titre exact {titre!r} : {len(trouvees)} slide(s) ; 1 exigee")
    return trouvees[0]


def ajouter_sur_gabarit(core, prs, titre_modele, remplacements):
    """Ajoute, juste apres la slide modele (designee par TITRE EXACT), une copie de celle-ci (layout, images, fond, relations : `core.copier_slide`),
    puis remplace le texte des formes : `remplacements` = {texte actuel exact d'une forme: nouveau texte} (une forme par cle, sinon ValueError ;
    la mise en forme du premier run est gardee). Renvoie la nouvelle slide. Les sections p14 ne sont pas mises a jour par `copier_slide` : un deck qui en a
    verra son controle d'integrite refuser une slide sans section (visible, jamais silencieux)."""
    modele = slide_par_titre(prs, titre_modele)
    rang = list(prs.slides).index(modele)
    neuve = core.copier_slide(prs, modele, index=rang + 1)
    for ancien, nouveau in remplacements.items():
        formes = [sh for sh in neuve.shapes if sh.has_text_frame and plat(sh.text_frame.text) == plat(ancien)]
        if len(formes) != 1:
            raise ValueError(f"gabarit {titre_modele!r} : forme de texte {ancien!r} -> {len(formes)} trouvee(s) ; 1 exigee")
        set_run(formes[0], nouveau)
    return neuve


__all__ = ["RGBColor", "id_de", "deplacer_apres", "nouvelle_apres", "clone_forme", "set_run", "rect", "texte", "slide_par_titre", "ajouter_sur_gabarit"]
