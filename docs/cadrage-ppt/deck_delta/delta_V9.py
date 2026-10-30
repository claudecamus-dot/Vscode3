"""Delta V9 (base V8) : apprentissages Vscode8 (navigation en haut) + charte du deck Naval Group (formes plus arrondies, pictos dessines) + police Outfit explicite.

Source des mecanismes (jamais du contenu) :
- Vscode8 : une ligne de navigation en haut de chaque slide de contenu, le chapitre courant en aplat sombre, les autres en clair ;
- deck « Revue V2 … Naval Group … » (OCTO 2024) : cartes aux grands rayons, pictos en disque blanc cerne de navy, glyphe navy + detail cyan, pose a cheval sur le bord d'une carte ;
  police Outfit portee explicitement par chaque run (le theme dit Arial dans les deux decks : sans nom de police explicite, le texte retombe sur Arial).
Usage (depuis docs/cadrage-ppt/) : py deck_delta/deck.py delta_V9 --base V8 --dry-run   puis sans --dry-run.
"""
import re

import deck_base
import deck_draw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt

BASE = "V8"
NOTE = "V9 : navigation en haut, pictos dessines, formes plus arrondies, police Outfit explicite sur tous les textes"

NAVY = RGBColor(0x0E, 0x23, 0x56)
CYAN = RGBColor(0x00, 0xD2, 0xDD)
MUTED = RGBColor(0x58, 0x65, 0x86)
CLAIR = RGBColor(0xE7, 0xE9, 0xEE)
BLANC = RGBColor(0xFF, 0xFF, 0xFF)
EMU = 12700
POLICE = "Outfit"
OPS = []

CHAPITRES_NAV = ["Contexte et enjeux", "Nos convictions", "Notre approche", "Démarche générale", "Démarche détaillée", "Preuve", "Dispositif et prix", "Prochain pas", "Annexes"]
T_CONV = "Cinq convictions : une cible à construire et une démarche pour y conduire"
T_REUSSIR = "Réussir : cinq conditions que nous posons avant de commencer"
T_IA = "L'IA apporte au coach du temps et de la matière, pas la décision"


def _ids_base():
    """Selecteurs `id:N` de TOUTES les slides de la base (resolus avant les etapes : l'ordre ne change pas pendant ce delta)."""
    prs = Presentation(deck_base.resoudre_base(BASE))
    return [f"id:{s.slide_id}" for s in prs.slides]


def _est_intercalaire(s):
    return s.slide_layout.name.startswith(("50 -", "51 -"))


def _nom_chapitre(s):
    return next((sh.text_frame.text for sh in s.shapes if sh.is_placeholder and sh.has_text_frame and sh.top > 100 * EMU
                 and not re.fullmatch(r"\d{2}|[A-Z]", sh.text_frame.text)), None)


def _chapitres(prs):
    """[(slide, chapitre courant ou None)] : None avant le premier intercalaire (ouverture) et pour couverture, sommaire, intercalaires."""
    courant, sortie = None, []
    for s in prs.slides:
        if _est_intercalaire(s):
            courant = _nom_chapitre(s)
            sortie.append((s, None))
        elif s.slide_layout.name.startswith(("95 -", "01 -", "02 -", "03 -")) or courant is None:
            sortie.append((s, None))
        else:
            sortie.append((s, courant))
    return sortie


def navigation(prs):
    """Une ligne de pastilles en haut de chaque slide de contenu : le chapitre courant en aplat navy, les autres en clair."""
    largeurs = [len(n) * 3.3 + 16 for n in CHAPITRES_NAV]
    total_dispo = 632 - 4 * (len(CHAPITRES_NAV) - 1)
    echelle = total_dispo / sum(largeurs)
    for s, chapitre in _chapitres(prs):
        if chapitre is None:
            continue
        x = 44.0
        for nom, w in zip(CHAPITRES_NAV, largeurs, strict=True):
            w *= echelle
            actif = nom == chapitre
            chip = deck_draw.rect(s, x, 7, w, 13, NAVY if actif else CLAIR, MSO_SHAPE.ROUNDED_RECTANGLE)
            chip.name = "nav-" + nom
            chip.adjustments[0] = 0.5
            tf = chip.text_frame
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.word_wrap = False
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = nom
            r.font.size, r.font.bold, r.font.name = Pt(6), actif, POLICE
            r.font.color.rgb = BLANC if actif else MUTED
            x += w + 4


def arrondir(prs):
    """Cartes aux grands rayons (charte Naval) : tout rectangle arrondi de plus de 40 pt de haut passe a un rayon d'environ 16 pt."""
    for s in prs.slides:
        for sh in s.shapes:
            if sh.shape_type == 1 and sh.name.startswith("Rounded Rectangle") and sh.height > 40 * EMU and not sh.name.startswith("nav-"):
                cote = min(sh.width, sh.height)
                sh.adjustments[0] = min(0.5, 16 * EMU / cote)


def police(prs):
    """Police Outfit explicite sur chaque run sans nom de police (sinon : Arial, la police du theme)."""
    for s in prs.slides:
        for sh in s.shapes:
            cadres = []
            if sh.has_text_frame:
                cadres.append(sh.text_frame)
            if sh.shape_type == 19:  # TABLE
                cadres += [c.text_frame for ligne in sh.table.rows for c in ligne.cells]
            for tf in cadres:
                for para in tf.paragraphs:
                    for r in para.runs:
                        if not r.font.name:
                            r.font.name = POLICE


def picto(s, genre, cx, cy, d):
    """Disque blanc cerne de navy, glyphe navy et detail cyan (style des pictos du deck Naval)."""
    disque = deck_draw.rect(s, cx - d / 2, cy - d / 2, d, d, BLANC, MSO_SHAPE.OVAL, ligne=NAVY)
    disque.name = f"picto-{genre}"
    disque.line.width = Pt(1)
    k = d / 24

    def forme(x, y, w, h, coul, kind):
        sh = deck_draw.rect(s, cx + x * k, cy + y * k, w * k, h * k, coul, kind)
        sh.name = f"picto-{genre}-glyphe"
        return sh
    if genre == "personne":
        forme(-3, -8, 6, 6, NAVY, MSO_SHAPE.OVAL)
        forme(-6, -1, 12, 8, NAVY, MSO_SHAPE.ROUND_2_SAME_RECTANGLE)
        forme(2, 3, 4, 4, CYAN, MSO_SHAPE.OVAL)
    elif genre == "cible":
        forme(-8, -8, 16, 16, NAVY, MSO_SHAPE.DONUT)
        forme(-2.5, -2.5, 5, 5, CYAN, MSO_SHAPE.OVAL)
    elif genre == "barres":
        forme(-7, 0, 4, 7, NAVY, MSO_SHAPE.RECTANGLE)
        forme(-2, -4, 4, 11, CYAN, MSO_SHAPE.RECTANGLE)
        forme(3, -7, 4, 14, NAVY, MSO_SHAPE.RECTANGLE)
    elif genre == "engrenage":
        forme(-8, -8, 16, 16, NAVY, MSO_SHAPE.GEAR_6)
        forme(-2.5, -2.5, 5, 5, CYAN, MSO_SHAPE.OVAL)
    elif genre == "bouclier":
        forme(-6, -7, 12, 14, NAVY, MSO_SHAPE.FLOWCHART_OFFPAGE_CONNECTOR)
        forme(-2.5, -3, 5, 4, CYAN, MSO_SHAPE.RECTANGLE)
    elif genre == "loupe":
        forme(-7, -7, 11, 11, NAVY, MSO_SHAPE.DONUT)
        poignee = forme(2, 3, 7, 2.5, CYAN, MSO_SHAPE.RECTANGLE)
        poignee.rotation = 45
    elif genre == "eclair":
        forme(-5, -8, 10, 16, CYAN, MSO_SHAPE.LIGHTNING_BOLT)
    elif genre == "cube":
        forme(-7, -7, 14, 14, NAVY, MSO_SHAPE.CUBE)
        forme(1, -7, 5, 5, CYAN, MSO_SHAPE.OVAL)
    else:
        raise ValueError(genre)


CARTES = [3, 7, 10, 13, 16]  # rang des 5 cartes dans les slides a cartes numerotees (modele : « Cinq convictions »)
PICTOS_PAR_SLIDE = {
    T_CONV: ("personne", "cible", "barres", "engrenage", "eclair"),
    T_REUSSIR: ("personne", "cible", "barres", "cube", "bouclier"),
    T_IA: ("loupe", "barres", "engrenage", "cible", "eclair"),
}


def pictos_cartes(prs):
    """Un picto a cheval sur le coin haut droit de chacune des 5 cartes des slides de convictions."""
    for titre, genres in PICTOS_PAR_SLIDE.items():
        s = deck_draw.slide_par_titre(prs, titre)
        formes = list(s.shapes)
        for rang, genre in zip(CARTES, genres, strict=True):
            carte = formes[rang]
            picto(s, genre, (carte.left + carte.width) / EMU - 18, carte.top / EMU, 24)


def retouche_globale(prs):
    """Une seule etape pour les 3 retouches sur toutes les slides : la simulation (garde + rendu) se paie une fois, pas trois (28 min mesurees en 4 etapes)."""
    navigation(prs)
    arrondir(prs)
    police(prs)


TOUTES = _ids_base()
ETAPES = [
    {"declarer": [T_CONV, T_REUSSIR, T_IA], "f": pictos_cartes, "nouvelles_max": 0, "nom": "pictos sur les cartes de convictions"},
    {"declarer": TOUTES, "f": retouche_globale, "nouvelles_max": 0, "nom": "navigation, formes arrondies, police Outfit"},
]
