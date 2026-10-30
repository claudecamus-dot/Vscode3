"""Delta V17 (base V16) : finitions des slides scindees (revue du rendu V16) : texte aligne a gauche (le premier paragraphe heritait du centrage), cartes moins hautes
(la moitie basse etait vide), corps a 10,5 pt."""
import deck_draw
from delta_V9 import EMU
from delta_V16 import T11, T38  # noqa: F401
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt

BASE = "V16"
NOTE = "V17 : slides scindees, texte a gauche, cartes moins hautes"
OPS = []
TITRES = ["Une conviction, un critère, un chantier : 1 à 3", "Une conviction, un critère, un chantier : 4 et 5",
          "Dix problèmes, dix apports, dix indicateurs : 1 à 5", "Dix problèmes, dix apports, dix indicateurs : 6 à 10"]


def finitions(prs):
    for titre in TITRES:
        s = deck_draw.slide_par_titre(prs, titre)
        total = 168 if "chantier" in titre else 185
        for sh in s.shapes:
            if sh.name in ("corps-carte", "texte-carte"):
                sh.height = (total - (60 if "chantier" in titre else 66)) * EMU
            elif sh.name == "panneau-regle":
                sh.height = total * EMU
            if sh.has_text_frame and sh.name in ("texte-carte", "entete-carte") or (sh.has_text_frame and sh.text_frame.text.startswith(("Une capacité entre", "RÈGLE"))
                                                                                    and sh.name.startswith("Rectangle")):
                for p in sh.text_frame.paragraphs:
                    p.alignment = PP_ALIGN.LEFT
            if sh.name == "texte-carte":
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.size == Pt(9.5):
                            r.font.size = Pt(10.5)
                        elif r.font.size == Pt(7):
                            r.font.size = Pt(7.5)


ETAPES = [{"declarer": TITRES, "f": finitions, "nouvelles_max": 0, "nom": "finitions des slides scindees"}]
